import ast
import csv
import io
import json
import math
import os
import random
import re
import secrets
import unicodedata
import zipfile
from datetime import datetime
from functools import wraps
from collections import defaultdict
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from flask import (
    Flask, abort, flash, g, has_request_context, jsonify, make_response, redirect, render_template,
    request, send_file, session, url_for
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from psycopg import DatabaseError, Error as DatabaseConnectionError, IntegrityError
from psycopg_pool import PoolTimeout
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from database import close_pool, connection as database_connection
from policy_defaults import DEFAULT_RULES_VERSION, DEFAULT_STUDENT_RULES_EN, DEFAULT_STUDENT_RULES_ES

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
AUDIO_DIR = os.getenv("AUDIO_DIR", os.path.join(BASE_DIR, "static", "audio"))
ALLOWED_AUDIO_EXTENSIONS = {"wav", "mp3", "m4a", "ogg", "aac"}


def env_bool(name, default):
    raw = os.getenv(name)
    if raw is None:
        return default
    normalized = raw.strip().lower()
    if normalized in {"1", "true", "yes", "on"}:
        return True
    if normalized in {"0", "false", "no", "off"}:
        return False
    raise RuntimeError(f"{name} must be one of: 1/0, true/false, yes/no, on/off")


app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "dev-change-this-secret-key")
app.config["TEACHER_ADMIN_NAME"] = os.getenv("TEACHER_ADMIN_NAME", "Administrator")
app.config["TEACHER_ADMIN_EMAIL"] = os.getenv("TEACHER_ADMIN_EMAIL", "admin@assessment.local")
app.config["TEACHER_ADMIN_PASSWORD"] = os.getenv("TEACHER_ADMIN_PASSWORD", "ChangeMe123")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SECURE"] = env_bool("SESSION_COOKIE_SECURE", True)
app.config["SESSION_COOKIE_SAMESITE"] = os.getenv("SESSION_COOKIE_SAMESITE", "Lax").strip().capitalize()
if app.config["SESSION_COOKIE_SAMESITE"] not in {"Lax", "Strict", "None"}:
    raise RuntimeError("SESSION_COOKIE_SAMESITE must be Lax, Strict, or None")
if app.config["SESSION_COOKIE_SAMESITE"] == "None" and not app.config["SESSION_COOKIE_SECURE"]:
    raise RuntimeError("SESSION_COOKIE_SAMESITE=None requires SESSION_COOKIE_SECURE=1")

TYPE_LABELS = {
    "multiple_choice": "Multiple Choice",
    "true_false": "True / False",
    "short_answer": "Short Answer",
    "numeric": "Numeric Answer",
    "order": "Put in Order",
    "matching": "Matching",
    "listening": "Audio / Listening",
}

SUPPORTED_UI_LANGUAGES = {"es": "Español", "en": "English"}

TRANSLATIONS_ES = {
    # Shared navigation / actions
    "Menu": "Menú", "Results": "Resultados", "Students": "Estudiantes", "Question bank": "Banco de preguntas",
    "Bank": "Banco", "Catalog": "Catálogo", "Test setup": "Configuración", "Setup": "Config.", "Log out": "Cerrar sesión",
    "Edit": "Editar", "Duplicate": "Duplicar", "Cancel": "Cancelar", "Save changes": "Guardar cambios",
    "Create question": "Crear pregunta", "Search": "Buscar", "Section": "Sección", "Status": "Estado",
    "Active": "Activo", "Inactive": "Inactivo", "All": "Todos", "Apply filters": "Aplicar filtros",
    "Selected": "Seleccionada", "Bank only": "Solo banco", "No audio": "Sin audio",
    "Remove from test": "Quitar del examen", "Select for test": "Seleccionar para examen",
    "No questions found": "No se encontraron preguntas", "Change the filters or create a new question.": "Cambia los filtros o crea una nueva pregunta.",
    # Question types
    "Multiple Choice": "Opción múltiple", "True / False": "Verdadero / Falso", "Short Answer": "Respuesta corta",
    "Numeric Answer": "Respuesta numérica", "Put in Order": "Ordenar", "Matching": "Relacionar", "Audio / Listening": "Audio / Comprensión auditiva",
    # Student access / exam
    "questions": "preguntas", "categories": "categorías", "Auto-scored": "Calificación automática",
    "Full name": "Nombre completo", "Write your full name": "Escribe tu nombre completo", "Section / group": "Sección / grupo",
    "Select your section": "Selecciona tu sección", "Example: 1° A": "Ejemplo: 1° A", "Start assessment": "Iniciar evaluación",
    "Before you start": "Antes de comenzar", "Teacher access": "Acceso docente",
    "Secure assessment": "Evaluación segura", "answered": "respondidas", "question": "pregunta", "questions_lower": "preguntas",
    "Exam monitoring": "Monitoreo del examen", "window/app focus changes recorded": "cambios de ventana/app registrados",
    "Enter focus mode": "Activar modo enfoque", "Integrity notice": "Aviso de integridad",
    "Play audio": "Reproducir audio", "plays available": "reproducciones disponibles", "play available": "reproducción disponible",
    "Listening source": "Fuente de audio", "Upload audio file": "Subir archivo de audio", "Browser voice (script)": "Voz del navegador (guion)",
    "Use a prerecorded audio file.": "Usar un archivo de audio pregrabado.", "The browser will read the script aloud when the student presses play.": "El navegador leerá el guion en voz alta cuando el estudiante pulse reproducir.",
    "Speech language": "Idioma de la voz", "Automatic / browser default": "Automático / predeterminado del navegador",
    "English (United States)": "Inglés (Estados Unidos)", "Spanish (Latin America)": "Español (Latinoamérica)", "Spanish (Spain)": "Español (España)",
    "Listening script": "Guion de listening", "Required when browser voice is selected.": "Obligatorio cuando se selecciona la voz del navegador.",
    "Preview the audio before saving.": "Preview the audio before saving.",
    "Select an audio file to test playback.": "Select an audio file to test playback.",
    "Browser voice is not supported on this device.": "La voz del navegador no es compatible con este dispositivo.",
    "Could not load the listening script.": "No se pudo cargar el guion de listening.",
    "Type your answer in the box.": "Escribe tu respuesta en el campo.", "Your answer": "Tu respuesta",
    "Enter a numeric value. Use a decimal point or comma if needed.": "Ingresa un valor numérico. Usa punto o coma decimal si es necesario.",
    "Numeric answer": "Respuesta numérica", "Drag the cards or use the arrow buttons to arrange the items in the correct order.": "Arrastra las tarjetas o usa las flechas para colocarlas en el orden correcto.",
    "Move up": "Mover arriba", "Move down": "Mover abajo", "Choose the best match for each item.": "Elige la relación correcta para cada elemento.",
    "Select...": "Seleccionar...", "Previous": "Anterior", "Next": "Siguiente", "Review & submit": "Revisar y enviar",
    "Ready to submit?": "¿Listo para enviar?", "Review your progress before sending your answers.": "Revisa tu progreso antes de enviar tus respuestas.",
    "Answered": "Respondidas", "Missing": "Pendientes", "Total": "Total", "Keep reviewing": "Seguir revisando", "Submit assessment": "Enviar evaluación",
    # Result
    "Your Result": "Tu resultado", "Assessment complete": "Evaluación finalizada", "Correct": "Correctas", "Percentage": "Porcentaje",
    "Submitted": "Enviado", "Performance by category": "Rendimiento por categoría", "Performance by question type": "Rendimiento por tipo de pregunta",
    "Exam integrity": "Integridad del examen", "No incidents": "Sin incidencias", "Review": "Revisar",
    "Estimated window/app changes": "Cambios estimados de ventana/app", "Detected time away": "Tiempo detectado fuera", "Blocked actions": "Acciones bloqueadas",
    "Download report PDF": "Descargar informe PDF", "Print result": "Imprimir resultado",
    "Download exam PDF": "Descargar PDF del examen",
    "Back to roster": "Volver al padrón", "Back to dashboard": "Volver al panel",
    "Teacher": "Docente", "Student name": "Nombre del estudiante", "Student ID": "NIE", "Section": "Sección", "Date": "Fecha",
    "Instructions": "Instrucciones", "Questions": "Preguntas", "Use blue or black ink.": "Usa tinta azul o negra.",
    "Write clearly and keep your answers legible.": "Escribe con claridad y mantén tus respuestas legibles.",
    "Do not write on the answer key; this copy is for student use only.": "No escribas sobre la clave de respuestas; esta copia es solo para uso del estudiante.",
    "Order line": "Línea de orden", "Answer line": "Línea de respuesta", "Left items": "Elementos de la izquierda", "Right items": "Elementos de la derecha",
    # Teacher dashboard
    "Teacher dashboard": "Panel docente", "submitted results": "resultados enviados", "active questions": "preguntas activas",
    "questions in the full bank": "preguntas en el banco completo", "Export Excel": "Exportar Excel", "Export CSV": "Exportar CSV",
    "Manage students": "Gestionar estudiantes", "Active test": "Examen activo", "Question bank stat": "Banco de preguntas",
    "Student roster": "Padrón de estudiantes", "student attempts": "intentos de estudiantes", "across all subjects": "en todas las asignaturas",
    "across": "en", "section(s)": "sección(es)", "Student": "Estudiante", "Assessment": "Evaluación", "Subject": "Asignatura",
    "Grade": "Nota", "Integrity": "Integridad", "Actions": "Acciones", "Report": "Informe", "No submitted results yet.": "Aún no hay resultados enviados.",
    # Students
    "Student roster eyebrow": "Padrón de estudiantes", "Students & sections": "Estudiantes y secciones", "Access settings": "Configuración de acceso",
    "Create a section first, then manage the students inside it.": "Crea primero una sección y luego gestiona los estudiantes dentro de ella.",
    "Use the section filter only to narrow by NIE, name, or section.": "Usa el filtro de sección solo para acotar por NIE, nombre o sección.",
    "Open section": "Abrir sección",
    "Active students": "Estudiantes activos", "available in the roster": "disponibles en el padrón", "Sections": "Secciones", "active groups": "grupos activos",
    "Inactive students": "Estudiantes inactivos", "kept for historical records": "conservados para historial", "Manual enrollment": "Registro manual",
    "Add student": "Agregar estudiante", "Roster": "Padrón", "Student code": "Código del estudiante", "optional": "opcional",
    "ID, code or carnet": "ID, código o carnet", "Notes": "Notas", "Short teacher note": "Nota breve del docente",
    "+ Add student": "+ Agregar estudiante", "Organization": "Organización", "Create section": "Crear sección", "Group": "Grupo",
    "Section name": "Nombre de sección", "Description": "Descripción", "Grade, shift or other reference": "Grado, turno u otra referencia",
    "+ Create section": "+ Crear sección", "Groups": "Grupos", "Save section": "Guardar sección", "Directory": "Directorio",
    "shown": "mostrados", "students": "estudiantes", "active": "activos", "selected": "seleccionados",
    "Student name or code...": "Nombre o código del estudiante...", "All sections": "Todas las secciones",
    "Edit student": "Editar estudiante", "Active student": "Estudiante activo", "Save changes student": "Guardar cambios",
    "No students found": "No se encontraron estudiantes", "Add a student manually or change the current filters.": "Agrega un estudiante manualmente o cambia los filtros actuales.",
    "No description yet.": "Aún sin descripción.", "Select all": "Seleccionar todo", "Bulk action": "Acción masiva",
    "Apply to selected": "Aplicar a seleccionados", "Move to another section": "Mover a otra sección",
    "Target section": "Sección destino", "0 selected": "0 seleccionados", "Add a student to this section or adjust the current filters.": "Agrega un estudiante a esta sección o ajusta los filtros actuales.",
    "Choose a section first.": "Primero elegí una sección.",
    "Archive selected students": "Archivar estudiantes seleccionados",
    # Question bank/editor
    "Reusable content": "Contenido reutilizable", "+ New question": "+ Nueva pregunta", "Type": "Tipo", "All types": "Todos los tipos",
    "All subjects": "Todas las asignaturas", "All categories": "Todas las categorías", "All status": "Todos los estados",
    "Bank only filter": "Solo banco", "Apply": "Aplicar", "Bulk action": "Acción masiva", "Add to test": "Agregar al examen",
    "Remove from test bulk": "Quitar del examen", "Archive": "Archivar", "Apply to selected": "Aplicar a seleccionadas",
    "Question editor": "Editor de preguntas", "New question": "Nueva pregunta", "Edit question": "Editar pregunta",
    "1. Classification": "1. Clasificación", "Subject and category": "Asignatura y categoría", "Category": "Categoría", "Question type": "Tipo de pregunta",
    "2. Question": "2. Pregunta", "Question prompt": "Enunciado", "Write the question the student will see...": "Escribe la pregunta que verá el estudiante...",
    "Include in current test": "Incluir en el examen actual", "If disabled, it stays in the bank for later.": "Si está desactivada, queda guardada en el banco para usar después.",
    "3. Answer options": "3. Opciones de respuesta", "+ Add option": "+ Agregar opción", "Correct": "Correcta", "Option": "Opción",
    "Audio": "Audio", "Current audio": "Audio actual", "Replace audio (optional)": "Reemplazar audio (opcional)", "Audio file": "Archivo de audio",
    "Teacher script / notes — never shown to students": "Guion / notas del docente — nunca se muestran a estudiantes",
    "Optional transcript, explanation or notes...": "Transcripción, explicación o notas opcionales...", "3. Correct answer": "3. Respuesta correcta", "Answer": "Respuesta",
    "True": "Verdadero", "False": "Falso", "3. Accepted answers": "3. Respuestas aceptadas", "Accepted answers": "Respuestas aceptadas",
    "Case-sensitive grading": "Distinguir mayúsculas y minúsculas", "3. Numeric grading": "3. Calificación numérica", "Correct value": "Valor correcto",
    "Tolerance ±": "Tolerancia ±", "3. Correct order": "3. Orden correcto", "Items": "Elementos", "3. Matching pairs": "3. Pares para relacionar",
    "+ Add pair": "+ Agregar par", "Left item": "Elemento izquierdo", "Matching answer": "Respuesta relacionada",
    # Catalog
    "Subjects & categories": "Asignaturas y categorías", "Subjects": "Asignaturas", "Create subject": "Crear asignatura", "Subject name": "Nombre de asignatura",
    "Categories": "Categorías", "Create category": "Crear categoría", "Category name": "Nombre de categoría", "Sort order": "Orden",
    "Save subject": "Guardar asignatura", "Save category": "Guardar categoría",
    # Setup
    "Current assessment": "Evaluación actual", "Test setup heading": "Configuración del examen", "Preview student page ↗": "Vista del estudiante ↗",
    "Institution": "Institución", "Assessment title": "Título de la evaluación", "Current subject": "Asignatura actual",
    "Student-page description": "Descripción para estudiantes", "Audio play limit": "Límite de reproducciones de audio", "Student access": "Acceso de estudiantes",
    "Open access — type name and section": "Acceso abierto — escribir nombre y sección", "Roster validation — registered students only": "Validación de padrón — solo estudiantes registrados",
    "Interface language": "Idioma de la interfaz", "Spanish": "Español", "English": "English", "Save test settings": "Guardar configuración",
    "Live composition": "Composición activa", "By category": "Por categoría", "No categories yet.": "Aún no hay categorías.",
    "Edit active questions": "Editar preguntas activas", "Safe editing": "Edición segura",
    # Login
    "Teacher access heading": "Acceso docente", "Enter dashboard": "Entrar al panel", "Back to assessment": "Volver a la evaluación",
    # Flash / validation
    "Enter your full name and section to begin.": "Ingresa tu nombre completo y sección para comenzar.",
    "This assessment is not available yet. Please contact your teacher.": "Esta evaluación aún no está disponible. Contacta a tu docente.",
    "Your name and section were not found in the active student roster. Check the information or contact your teacher.": "Tu nombre y sección no aparecen en el padrón activo. Verifica los datos o contacta a tu docente.",
    "Incorrect teacher PIN.": "PIN docente incorrecto.", "Write a section or group name.": "Escribe un nombre de sección o grupo.",
    "Section created.": "Sección creada.", "That section already exists.": "Esa sección ya existe.", "Section name cannot be empty.": "El nombre de la sección no puede estar vacío.",
    "Section updated.": "Sección actualizada.", "Another section already uses that name.": "Otra sección ya usa ese nombre.",
    "Student name and section are required.": "El nombre del estudiante y la sección son obligatorios.", "Select a valid section.": "Selecciona una sección válida.",
    "Student added to the roster.": "Estudiante agregado al padrón.", "That student already exists in the section, or the student code is already in use.": "Ese estudiante ya existe en la sección o el código ya está en uso.",
    "Student updated.": "Estudiante actualizado.", "Student status updated.": "Estado del estudiante actualizado.",
    "Question added to the bank.": "Pregunta agregada al banco.", "Question updated. Existing student attempts keep their original snapshot.": "Pregunta actualizada. Los intentos ya iniciados conservan su versión original.",
    "Question availability updated.": "Disponibilidad de la pregunta actualizada.", "Question duplicated as bank-only so you can edit it safely.": "Pregunta duplicada como 'solo banco' para que puedas editarla con seguridad.",
    "Select at least one question.": "Selecciona al menos una pregunta.", "Choose a valid bulk action.": "Elige una acción masiva válida.",
    "Write a subject name.": "Escribe un nombre de asignatura.", "Subject created.": "Asignatura creada.", "A subject with that name already exists.": "Ya existe una asignatura con ese nombre.",
    "Subject updated.": "Asignatura actualizada.", "Choose a valid subject.": "Elige una asignatura válida.", "Write a category name.": "Escribe un nombre de categoría.",
    "Selected subject does not exist.": "La asignatura seleccionada no existe.", "Category created.": "Categoría creada.", "That category already exists in this subject.": "Esa categoría ya existe en esta asignatura.",
    "Category updated.": "Categoría actualizada.", "Institution, assessment title, and a valid subject are required.": "La institución, el título de la evaluación y una asignatura válida son obligatorios.",
    "Assessment settings saved.": "Configuración de la evaluación guardada.",
    "Add at least two options and select the correct one.": "Agrega al menos dos opciones y selecciona la correcta.",
    "Add students manually, organize them by section, and optionally require roster validation before an assessment begins.": "Agrega estudiantes manualmente, organízalos por sección y, si lo deseas, exige validación contra el padrón antes de iniciar una evaluación.",
    "Assessment identity": "Identidad de la evaluación", "Back to bank": "Volver al banco", "Bank organization": "Organización del banco",
    "Choose the subject and configure the assessment students will see now.": "Elige la asignatura y configura la evaluación que verán los estudiantes.",
    "Create a section first before adding students.": "Crea primero una sección antes de agregar estudiantes.",
    "Create as many subjects and categories as you need. Categories are the topics, units, competencies or blocks used to organize and report questions.": "Crea todas las asignaturas y categorías que necesites. Las categorías pueden representar temas, unidades, competencias o bloques para organizar y reportar las preguntas.",
    "Create reusable questions and classify them by subject and category.": "Crea preguntas reutilizables y clasifícalas por asignatura y categoría.",
    "During the assessment, tab/app changes, focus losses and blocked actions may be recorded for academic-integrity review. Audio questions, when present, can be played up to {count} time(s).": "Durante la evaluación se pueden registrar cambios de pestaña/app, pérdidas de foco y acciones bloqueadas para revisión de integridad académica. Las preguntas con audio pueden reproducirse hasta {count} vez/veces.",
    "Each left item is paired with the answer on its right.": "Cada elemento de la izquierda se relaciona con la respuesta de la derecha.",
    "Enter PIN": "Ingresa el PIN", "Example: 3.14": "Ejemplo: 3.14",
    "Example: correct value 10 with tolerance 0.5 accepts answers from 9.5 to 10.5.": "Ejemplo: un valor correcto de 10 con tolerancia 0.5 acepta respuestas entre 9.5 y 10.5.",
    "Filter": "Filtrar", "First step": "Primer paso", "Second step": "Segundo paso", "Third step": "Tercer paso",
    "Focus changes": "Cambios de foco", "Inactive students cannot pass roster validation.": "Los estudiantes inactivos no pueden superar la validación del padrón.",
    "Leaving this exam, switching tabs/apps, right-clicking, copying, pasting, or blocked shortcuts may be recorded for teacher review.": "Salir del examen, cambiar de pestaña/app, hacer clic derecho, copiar, pegar o usar atajos bloqueados puede registrarse para revisión docente.",
    "Open results, question banks, subjects, categories and assessment settings.": "Accede a resultados, bancos de preguntas, asignaturas, categorías y configuración de evaluaciones.",
    "Optional description": "Descripción opcional",
    "Organize reusable questions by subject and category. Activate only the questions you want in each subject's current test.": "Organiza preguntas reutilizables por asignatura y categoría. Activa únicamente las preguntas que quieras usar en el examen actual de cada asignatura.",
    "Question, subject or category...": "Pregunta, asignatura o categoría...", "Remove from subject test": "Quitar del examen de la asignatura",
    "Roster validation matches the student's full name and section against Students.": "La validación de padrón compara el nombre completo y la sección con los estudiantes registrados.",
    "Score": "Puntaje", "Select all": "Seleccionar todo", "Select for subject test": "Seleccionar para el examen de la asignatura",
    "Select section": "Seleccionar sección", "Selected for test": "Seleccionada para examen", "Student full name": "Nombre completo del estudiante",
    "Student submissions will appear here automatically.": "Los envíos de los estudiantes aparecerán aquí automáticamente.",
    "Students already taking a test keep the exact question snapshot they started with. New changes affect only future attempts.": "Los estudiantes que ya iniciaron conservan exactamente la versión de preguntas con la que comenzaron. Los cambios nuevos solo afectan intentos futuros.",
    "These are technical indicators for teacher review; they do not automatically prove misconduct.": "Estos son indicadores técnicos para revisión docente; no prueban automáticamente una conducta indebida.",
    "This assessment has no active questions yet. The teacher can add questions to the bank and activate them from the dashboard.": "Esta evaluación aún no tiene preguntas activas. El docente puede agregar preguntas al banco y activarlas desde el panel.",
    "This changes the application interface, not the text of questions you created.": "Esto cambia la interfaz de la aplicación, no el texto de las preguntas que hayas creado.",
    "Use only when capitalization must matter.": "Úsalo solo cuando las mayúsculas y minúsculas deban importar.", "View": "Ver",
    "Write one accepted answer per line. By default, capitalization is ignored.": "Escribe una respuesta aceptada por línea. Por defecto se ignoran mayúsculas y minúsculas.",
    "Write one item per line in the correct order. Students receive them shuffled.": "Escribe un elemento por línea en el orden correcto. Los estudiantes los recibirán desordenados.",
    "Your full name and section must match the teacher roster.": "Tu nombre completo y sección deben coincidir con el padrón del docente.",
    "Your report contains the score summary only; it does not reveal the answer key.": "Tu informe contiene únicamente el resumen de calificación; no muestra la clave de respuestas.",
    "category(s)": "categoría(s)", "selected": "seleccionada(s)", "subject(s)": "asignatura(s)",
    "Question type": "Tipo de pregunta", "Exam integrity summary": "Resumen de integridad del examen",
    "Total detected time away": "Tiempo total detectado fuera",
    "Integrity events are technical indicators for teacher review and are not, by themselves, proof of misconduct.": "Los eventos de integridad son indicadores técnicos para revisión docente y, por sí solos, no constituyen prueba de conducta indebida.",
    "This report records the student's submitted result. It does not display the answer key.": "Este informe registra el resultado enviado por el estudiante. No muestra la clave de respuestas.",
    "Attempt ID": "ID del intento", "Started": "Inicio", "Integrity Status": "Estado de integridad",
    "Tab/App Visibility Changes": "Cambios de visibilidad pestaña/app", "Total Time Away (s)": "Tiempo total fuera (s)",
    "Longest Time Away (s)": "Mayor tiempo fuera (s)", "Window Blur Events": "Pérdidas de foco de ventana",
    "Page Hide Events": "Eventos de página oculta", "Right-click Attempts": "Intentos de clic derecho",
    "Copy Attempts": "Intentos de copiar", "Cut Attempts": "Intentos de cortar", "Paste Attempts": "Intentos de pegar",
    "Blocked Shortcuts": "Atajos bloqueados", "Fullscreen Exits": "Salidas de pantalla completa",
    "Event Type": "Tipo de evento", "Server Time": "Hora del servidor", "Details": "Detalles",
    "Integrity Events": "Eventos de integridad", "No subject selected": "No hay asignatura seleccionada", "No subject": "Sin asignatura",
    "Select a valid question type.": "Selecciona un tipo de pregunta válido.", "Select a valid subject and category.": "Selecciona una asignatura y categoría válidas.",
    "The selected category does not belong to that subject.": "La categoría seleccionada no pertenece a esa asignatura.", "Write the question prompt.": "Escribe el enunciado de la pregunta.",
    "This question type needs at least two answer options.": "Este tipo de pregunta necesita al menos dos opciones de respuesta.", "Select a non-empty correct answer.": "Selecciona una respuesta correcta que no esté vacía.",
    "Audio must be WAV, MP3, M4A, OGG, or AAC.": "El audio debe ser WAV, MP3, M4A, OGG o AAC.", "Upload an audio file for the audio/listening question.": "Sube un archivo de audio para la pregunta de audio/comprensión auditiva.",
    "Choose True or False as the correct answer.": "Elige Verdadero o Falso como respuesta correcta.", "Add at least one accepted answer, one per line.": "Agrega al menos una respuesta aceptada, una por línea.",
    "Numeric answer and tolerance must be valid numbers.": "La respuesta numérica y la tolerancia deben ser números válidos.", "Ordering questions need at least two items, one per line.": "Las preguntas de ordenar necesitan al menos dos elementos, uno por línea.",
    "Matching questions need at least two complete pairs.": "Las preguntas de relacionar necesitan al menos dos pares completos.",
}

TRANSLATIONS_ES.update({
    "Teachers": "Docentes", "Teacher users": "Usuarios docentes", "Manage teacher accounts": "Administrar cuentas docentes",
    "Teacher name": "Nombre del docente", "Teacher email": "Correo del docente", "Role": "Rol",
    "Administrator": "Administrador", "Teacher": "Docente", "Create teacher": "Crear docente",
    "New teacher": "Nuevo docente", "Teacher account created.": "Cuenta docente creada.",
    "Teacher account updated.": "Cuenta docente actualizada.", "Teacher account disabled.": "Cuenta docente desactivada.",
    "Teacher account enabled.": "Cuenta docente activada.", "Reset teacher password": "Restablecer contraseña docente",
    "Teacher password reset.": "Contraseña docente restablecida.", "Only administrators can manage teacher accounts.": "Solo los administradores pueden gestionar cuentas docentes.",
    "Invalid teacher email or password.": "Correo o contraseña docente incorrectos.",
    "Your teacher account is inactive.": "Tu cuenta docente está inactiva.",
    "Signed in as": "Sesión de", "Shared institutional catalog": "Catálogo institucional compartido",
    "Teacher workspace": "Espacio del docente", "My content": "Mi contenido",
    "Create separate teacher accounts. Each teacher has an isolated question bank, exams and results.": "Crea cuentas docentes separadas. Cada docente tiene un banco de preguntas, exámenes y resultados aislados.",
    "Last login": "Último acceso", "Deactivate": "Desactivar", "Activate": "Activar",
    "Reset password": "Restablecer contraseña", "Teacher name and a valid email are required.": "Se requiere el nombre del docente y un correo válido.",
    "That email address is already assigned to another teacher.": "Ese correo ya está asignado a otro docente.",
    "At least one active administrator is required.": "Debe existir al menos un administrador activo.",
    "You cannot disable your own active session.": "No puedes desactivar tu propia sesión activa.",
    "Your question counts are shown for your teacher workspace.": "Los conteos de preguntas corresponden a tu espacio docente.",
})


def _ui_translation_keys():
    keys = set()
    tree = ast.parse(Path(__file__).read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call) or not node.args or not isinstance(node.args[0], ast.Constant):
            continue
        name = getattr(node.func, "id", getattr(node.func, "attr", ""))
        if name in {"tr", "rt", "flash_ui"} and isinstance(node.args[0].value, str):
            keys.add(node.args[0].value)
    pattern = re.compile(r"(?<![A-Za-z_])t\(\s*['\"]([^'\"]+)['\"]")
    for template in (Path(__file__).parent / "templates").glob("*.html"):
        keys.update(pattern.findall(template.read_text(encoding="utf-8")))
    return keys


for _translation_key in _ui_translation_keys():
    TRANSLATIONS_ES.setdefault(_translation_key, _translation_key)

TRANSLATIONS_ES.update({
    "A version with that name already exists.": "Ya existe una version con ese nombre.",
    "Account security": "Seguridad de la cuenta", "Accounts and access": "Cuentas y acceso",
    "Active exams": "Examenes activos", "Add questions to at least one active version before publishing.": "Agrega preguntas al menos a una version activa antes de publicar.",
    "Administration": "Administracion", "Assessment statistics": "Estadisticas de evaluaciones", "Assigned": "Asignadas",
    "Average grade": "Nota promedio", "Average percentage": "Porcentaje promedio",
    "CSV and Excel downloads follow the filters below.": "Las descargas CSV y Excel respetan los filtros siguientes.",
    "Choose a valid fixed version.": "Elige una version fija valida.", "Choose a valid section.": "Elige una seccion valida.",
    "Clean review rate": "Tasa sin incidencias", "Collapse navigation": "Contraer navegacion", "Content": "Contenido", "Copied": "Copiado",
    "Create exams by subject, build multiple versions, and assign them to sections.": "Crea examenes por asignatura, prepara varias versiones y asignalas a secciones.",
    "Create your first exam and then build one or more versions from the question bank.": "Crea tu primer examen y luego prepara una o mas versiones desde el banco de preguntas.",
    "Developing": "En desarrollo", "Each assessment keeps its assigned version and can be submitted only once.": "Cada evaluacion conserva su version asignada y solo puede enviarse una vez.",
    "Exam archived.": "Examen archivado.", "Exam title and a valid subject are required.": "Se requieren el titulo del examen y una asignatura valida.",
    "Exam title is required.": "El titulo del examen es obligatorio.", "Expand navigation": "Expandir navegacion", "Exports": "Exportaciones",
    "Institution and exam policy": "Institucion y politica de examenes", "Manage exams": "Gestionar examenes",
    "Monitor outcomes, review integrity signals, and move directly into assessment work.": "Supervisa resultados, revisa indicadores de integridad y gestiona las evaluaciones.",
    "More": "Mas", "Needs review": "Requiere revision", "Needs support": "Requiere apoyo", "Next step": "Siguiente paso",
    "No submitted grades": "Sin notas enviadas", "Nothing pending": "Nada pendiente", "Operations": "Operaciones",
    "Performance distribution": "Distribucion del rendimiento", "Questions": "Preguntas", "Results overview": "Resumen de resultados",
    "Reusable assessment items": "Preguntas reutilizables", "Review queue": "Cola de revision", "Roster and sections": "Padron y secciones",
    "Secondary destinations": "Destinos secundarios", "See what is ready, resume active work, and review submitted results.": "Consulta lo disponible, continua el trabajo activo y revisa los resultados enviados.",
    "Skip to main content": "Saltar al contenido principal", "Strong": "Solido", "Student dashboard": "Panel del estudiante",
    "Subjects and categories": "Asignaturas y categorias", "Submission explorer": "Explorador de envios",
    "Submissions and integrity": "Envios e integridad", "Submitted grade profile": "Perfil de notas enviadas", "Submitted results": "Resultados enviados",
    "Take the current view with you": "Descarga la vista actual", "Teaching resources": "Recursos docentes",
    "Technical signals support review; they are not proof of misconduct.": "Los indicadores tecnicos apoyan la revision; no constituyen prueba de conducta indebida.",
    "This assessment is assigned to your section.": "Esta evaluacion esta asignada a tu seccion.", "Version name is required.": "El nombre de la version es obligatorio.",
    "Versions and assignments": "Versiones y asignaciones", "Workspace": "Espacio de trabajo",
    "You have submitted every assigned assessment.": "Has enviado todas las evaluaciones asignadas.", "Your assigned work": "Tus evaluaciones asignadas",
    "average grade": "nota promedio", "complete": "completado", "flagged submissions": "envios marcados", "focus changes": "cambios de foco",
    "no incidents": "sin incidencias", "out of 10": "sobre 10", "pending": "pendientes", "published exams": "examenes publicados",
    "ready to resume": "listas para continuar", "ready to start": "listas para iniciar", "submitted attempts": "intentos enviados",
    "total": "total", "total exams": "examenes totales",
})

TRANSLATIONS_ES.update({
    "Correct answers": "Respuestas correctas",
    "Answer every question with a valid response before submitting.": "Responde todas las preguntas con una respuesta valida antes de enviar.",
    "You must accept the current usage rules before continuing.": "Debes aceptar las reglas de uso vigentes antes de continuar.",
    "You must accept the current usage rules before starting or resuming an exam.": "Debes aceptar las reglas de uso vigentes antes de iniciar o continuar un examen.",
    "Usage rules": "Reglas de uso",
    "Read and accept these rules before continuing.": "Lee y acepta estas reglas antes de continuar.",
    "I have read and accept the usage rules.": "He leido y acepto las reglas de uso.",
    "Accept and continue": "Aceptar y continuar",
    "Rules version": "Version de reglas",
    "Require acknowledgment each session": "Exigir aceptacion en cada sesion",
    "Spanish usage rules": "Reglas de uso en espanol",
    "English usage rules": "Reglas de uso en ingles",
    "Institution and both usage-rule translations are required.": "La institucion y ambas traducciones de las reglas de uso son obligatorias.",
    "Complete the current question before continuing.": "Completa correctamente la pregunta actual antes de continuar.",
    "Enter a finite numeric value.": "Ingresa un valor numerico finito.",
    "Arrange every item before continuing.": "Ordena todos los elementos antes de continuar.",
    "Choose a different match for every item.": "Elige una relacion diferente para cada elemento.",
    "Raw grade": "Nota academica",
    "Penalty points": "Puntos de penalizacion",
    "Adjusted final grade": "Nota final ajustada",
    "Manual grade penalties": "Penalizaciones manuales de nota",
    "Penalty reason": "Motivo de la penalizacion",
    "Apply penalty": "Aplicar penalizacion",
    "Penalty history": "Historial de penalizaciones",
    "Penalty active": "Penalizacion activa",
    "Revoked": "Revocada",
    "Revoke penalty": "Revocar penalizacion",
    "Active penalty reasons": "Motivos de penalizacion activos",
    "A reason is required for every penalty.": "Toda penalizacion requiere un motivo.",
    "Penalty points must be greater than zero and may not exceed the current adjusted grade.": "Los puntos deben ser mayores que cero y no pueden superar la nota ajustada actual.",
    "Penalty applied.": "Penalizacion aplicada.",
    "Penalty revoked.": "Penalizacion revocada.",
    "This deduction was applied manually after teacher review. Integrity events never deduct points automatically.": "Esta deduccion fue aplicada manualmente despues de la revision docente. Los eventos de integridad nunca descuentan puntos automaticamente.",
    "No penalties have been applied.": "No se han aplicado penalizaciones.",
    "CSV columns: Subject, Category, Type, Question, Options, Answer, Pairs, Order, Tolerance, CaseSensitive, Script, and Audio.": "Columnas CSV: Asignatura, Categoria, Tipo, Pregunta, Opciones, Respuesta, Pares, Orden, Tolerancia, SensibleMayusculas, Script y Audio.",
    "The CSV must include student ID, first name, and last name columns.": "El CSV debe incluir las columnas NIE, Nombre y Apellido.",
    "The CSV must include type and question columns.": "El CSV debe incluir las columnas Tipo y Pregunta.",
    "Revoke this penalty?": "Revocar esta penalizacion?",
    "The raw academic score will remain unchanged.": "La calificacion academica original permanecera sin cambios.",
})

TRANSLATIONS_ES.update({
    "Integrity event review": "Revision de eventos de integridad",
    "Review the recorded browser signals in context before making a decision.": "Revisa las senales del navegador registradas y su contexto antes de tomar una decision.",
    "Telemetry supports teacher judgment; it is not proof of misconduct by itself.": "La telemetria apoya el criterio docente; por si sola no prueba una conducta indebida.",
    "Away periods": "Periodos fuera del examen", "Total away time": "Tiempo total fuera",
    "Fullscreen exits": "Salidas de pantalla completa", "No integrity events were recorded.": "No se registraron eventos de integridad.",
    "Technical details": "Detalles tecnicos", "Question {number}": "Pregunta {number}",
    "Question reference unavailable": "Referencia de pregunta no disponible", "Time unavailable": "Hora no disponible",
    "Left the exam for {duration}": "Salio del examen durante {duration}",
    "The exam tab became hidden and a return was recorded after {duration}.": "La pestana del examen quedo oculta y se registro el regreso despues de {duration}.",
    "The exam tab became hidden and a return was recorded.": "La pestana del examen quedo oculta y se registro el regreso.",
    "Left or switched away from the exam tab": "Salio o cambio de la pestana del examen",
    "The exam tab became hidden; no matching return was recorded.": "La pestana del examen quedo oculta; no se registro un regreso correspondiente.",
    "Returned to the exam": "Regreso al examen",
    "A return to the exam was recorded without a matching tab departure.": "Se registro un regreso al examen sin una salida de pestana correspondiente.",
    "The exam window lost focus": "La ventana del examen perdio el foco",
    "The browser recorded that the exam window was no longer active.": "El navegador registro que la ventana del examen dejo de estar activa.",
    "Left or reloaded the exam page": "Salio o recargo la pagina del examen",
    "The browser recorded that the exam page was left, closed, or reloaded.": "El navegador registro que la pagina del examen se abandono, cerro o recargo.",
    "Blocked context menu attempt": "Intento bloqueado de abrir el menu contextual",
    "The exam blocked an attempt to open the context menu.": "El examen bloqueo un intento de abrir el menu contextual.",
    "Blocked copy attempt": "Intento bloqueado de copiar", "The exam blocked an attempt to copy content.": "El examen bloqueo un intento de copiar contenido.",
    "Blocked cut attempt": "Intento bloqueado de cortar", "The exam blocked an attempt to cut content.": "El examen bloqueo un intento de cortar contenido.",
    "Blocked paste attempt": "Intento bloqueado de pegar", "The exam blocked an attempt to paste content.": "El examen bloqueo un intento de pegar contenido.",
    "Blocked attempt to {action}": "Intento bloqueado de {action}",
    "The exam blocked a restricted keyboard shortcut associated with {action}.": "El examen bloqueo un atajo de teclado restringido asociado con {action}.",
    "Attempted a restricted shortcut": "Intento usar un atajo restringido",
    "The exam blocked a restricted keyboard shortcut.": "El examen bloqueo un atajo de teclado restringido.",
    "Left fullscreen": "Salio de pantalla completa",
    "The browser recorded that fullscreen mode ended during the exam.": "El navegador registro que el modo de pantalla completa termino durante el examen.",
    "Recorded integrity event": "Evento de integridad registrado",
    "The browser recorded an event that is available for technical review.": "El navegador registro un evento disponible para revision tecnica.",
    "copy": "copiar", "cut": "cortar", "paste": "pegar", "select all": "seleccionar todo", "save": "guardar",
    "print": "imprimir", "view page source": "ver el codigo fuente", "open developer tools": "abrir las herramientas de desarrollo",
    "use the address bar": "usar la barra de direcciones", "open a new tab": "abrir una pestana nueva",
    "open a new window": "abrir una ventana nueva", "close the tab": "cerrar la pestana",
})

TRANSLATIONS_ES.update({
    "Email": "Correo electrónico",
    "Email address": "Correo electrónico",
    "Student email": "Correo del estudiante",
    "student@example.com": "estudiante@ejemplo.com",
    "Password": "Contraseña",
    "Current password": "Contraseña actual",
    "New password": "Nueva contraseña",
    "Confirm new password": "Confirmar nueva contraseña",
    "Sign in": "Iniciar sesión",
    "Student account": "Cuenta del estudiante",
    "Student accounts": "Cuentas de estudiantes",
    "Account access": "Acceso por cuenta",
    "Students sign in with their registered email and password.": "Los estudiantes inician sesión con su correo registrado y contraseña.",
    "Initial password": "Contraseña inicial",
    "Password setup": "Configuración de contraseña",
    "Generate secure temporary password": "Generar contraseña temporal segura",
    "Use a teacher-defined password": "Usar una contraseña definida por el docente",
    "Teacher-defined password": "Contraseña definida por el docente",
    "At least 8 characters, including a letter and a number.": "Mínimo 8 caracteres, incluyendo una letra y un número.",
    "Require password change at next sign-in": "Exigir cambio de contraseña en el próximo inicio",
    "Recommended for generated or reset passwords.": "Recomendado para contraseñas generadas o restablecidas.",
    "Temporary credentials": "Credenciales temporales",
    "Copy credentials": "Copiar credenciales",
    "Password shown only once": "La contraseña se muestra solo una vez",
    "Save or deliver these credentials now. The stored password cannot be viewed later.": "Guarda o entrega estas credenciales ahora. La contraseña almacenada no puede consultarse después.",
    "Account ready": "Cuenta lista",
    "Account pending": "Cuenta pendiente",
    "No password set": "Sin contraseña configurada",
    "Reset password": "Restablecer contraseña",
    "Generate new temporary password": "Generar nueva contraseña temporal",
    "Set a new password": "Definir una nueva contraseña",
    "New teacher-defined password": "Nueva contraseña definida por el docente",
    "Update password": "Actualizar contraseña",
    "Change password": "Cambiar contraseña",
    "Password updated successfully.": "Contraseña actualizada correctamente.",
    "Your password must be changed before you start the assessment.": "Debes cambiar tu contraseña antes de iniciar la evaluación.",
    "For security, choose a new password that only you know.": "Por seguridad, elige una nueva contraseña que solo tú conozcas.",
    "Your current password is incorrect.": "Tu contraseña actual es incorrecta.",
    "The new passwords do not match.": "Las nuevas contraseñas no coinciden.",
    "Password must be at least 8 characters and include at least one letter and one number.": "La contraseña debe tener al menos 8 caracteres e incluir al menos una letra y un número.",
    "Invalid email or password, or the account is inactive.": "Correo o contraseña incorrectos, o la cuenta está inactiva.",
    "Enter your email and password to begin.": "Ingresa tu correo y contraseña para comenzar.",
    "A valid student email is required.": "Se requiere un correo electrónico válido para el estudiante.",
    "That email address is already assigned to another student.": "Ese correo electrónico ya está asignado a otro estudiante.",
    "Choose an initial password method.": "Elige un método de contraseña inicial.",
    "Student added. Temporary credentials are shown below.": "Estudiante agregado. Las credenciales temporales se muestran abajo.",
    "Student account updated.": "Cuenta del estudiante actualizada.",
    "Password reset. New credentials are shown below.": "Contraseña restablecida. Las nuevas credenciales se muestran abajo.",
    "You need an active student account with email and password to start an assessment.": "Necesitas una cuenta activa de estudiante con correo y contraseña para iniciar una evaluación.",
    "Forgot your password? Ask your teacher to reset it.": "¿Olvidaste tu contraseña? Solicita al docente que la restablezca.",
    "Sign out": "Cerrar sesión",
    "Password changed. You can continue.": "Contraseña cambiada. Puedes continuar.",
    "Student email or code...": "Correo, nombre o código del estudiante...",
    "Email and password are required for every student account.": "Cada cuenta de estudiante requiere correo y contraseña.",
    "Legacy roster students without credentials are marked Account pending until a teacher sets an email and password.": "Los estudiantes migrados sin credenciales aparecen como Cuenta pendiente hasta que el docente configure correo y contraseña.",
    "Account": "Cuenta",
    "Last sign-in": "Último acceso",
    "Never": "Nunca",
    "Email is used as the student username and must be unique.": "El correo se usa como nombre de usuario del estudiante y debe ser único.",
    "Student sign-in": "Inicio de sesión del estudiante",
    "Registered accounts only — email + password": "Solo cuentas registradas — correo + contraseña",
    "Student authentication": "Autenticación de estudiantes",
    "Passwords are stored securely as hashes. Teachers can reset them, but cannot view an existing password.": "Las contraseñas se almacenan de forma segura como hashes. El docente puede restablecerlas, pero no puede ver una contraseña existente.",
    "Only required when you choose a teacher-defined password.": "Solo se requiere cuando eliges una contraseña definida por el docente.",
})


TRANSLATIONS_ES.update({
    "Filters": "Filtros",
    "Filter results": "Filtrar resultados",
    "Clear filters": "Limpiar filtros",
    "All subjects": "Todas las asignaturas",
    "All assessments": "Todas las evaluaciones",
    "All integrity states": "Todos los estados de integridad",
    "From date": "Desde",
    "To date": "Hasta",
    "Search student or email...": "Buscar estudiante o correo...",
    "Filtered results": "Resultados filtrados",
    "of total": "de un total de",
    "Download filtered Excel": "Descargar Excel filtrado",
    "Download filtered CSV": "Descargar CSV filtrado",
    "Current filters are applied to exports.": "Los filtros actuales también se aplican a las exportaciones.",
    "Edit section": "Editar sección",
    "Edit subject": "Editar asignatura",
    "Edit category": "Editar categoría",
    "Reset student password": "Restablecer contraseña del estudiante",
    "Archive student": "Archivar estudiante",
    "Archive section": "Archivar sección",
    "Archive subject": "Archivar asignatura",
    "Archive category": "Archivar categoría",
    "Archive question": "Archivar pregunta",
    "Archive": "Archivar",
    "Delete / archive": "Eliminar / archivar",
    "This action requires confirmation.": "Esta acción requiere confirmación.",
    "Are you sure?": "¿Estás seguro?",
    "This action cannot be undone from this screen.": "Esta acción no se puede deshacer desde esta pantalla.",
    "Yes, continue": "Sí, continuar",
    "No, cancel": "No, cancelar",
    "Student archived.": "Estudiante archivado.",
    "Section archived.": "Sección archivada.",
    "Subject archived.": "Asignatura archivada.",
    "Category archived.": "Categoría archivada.",
    "Question archived.": "Pregunta archivada.",
    "Move or archive the students in this section before archiving it.": "Mueve o archiva los estudiantes de esta sección antes de archivarla.",
    "At least one active subject must remain.": "Debe quedar al menos una asignatura activa.",
    "Archived items remain in historical results but are hidden from normal management views.": "Los elementos archivados se conservan en resultados históricos, pero se ocultan de las vistas normales de gestión.",
    "Edit details": "Editar datos",
    "Close": "Cerrar",
    "Save section": "Guardar sección",
    "Save subject": "Guardar asignatura",
    "Save category": "Guardar categoría",
    "Danger zone": "Zona de riesgo",
    "Results shown": "Resultados mostrados",
})

TRANSLATIONS_ES.update({
    "Assessments": "Exámenes", "Exams": "Exámenes", "Exam manager": "Gestor de exámenes",
    "Create exam": "Crear examen", "+ Create exam": "+ Crear examen", "Manage exam": "Gestionar examen",
    "Exam title": "Título del examen", "Exam description": "Descripción del examen", "Published": "Publicado",
    "Draft": "Borrador", "Versions": "Versiones", "Version": "Versión", "Version name": "Nombre de versión",
    "Add version": "Agregar versión", "+ Add version": "+ Agregar versión", "Duplicate version": "Duplicar versión",
    "Question selection": "Selección de preguntas", "Edit questions": "Editar preguntas", "Save question selection": "Guardar selección de preguntas",
    "Assignments": "Asignaciones", "Assign to section": "Asignar a sección", "Assignment mode": "Modo de asignación",
    "Random version": "Versión aleatoria", "Fixed version": "Versión fija", "Choose fixed version": "Elegir versión fija",
    "Assign / update": "Asignar / actualizar", "Unassign": "Desasignar", "Available exams": "Exámenes disponibles",
    "My assessments": "Mis evaluaciones", "Welcome": "Bienvenido", "Available": "Disponible", "Completed": "Completado",
    "In progress": "En progreso", "Resume": "Continuar", "Start exam": "Iniciar examen", "No exams available": "No hay exámenes disponibles",
    "Your teacher has not assigned an active exam to your section yet.": "Tu docente aún no ha asignado un examen activo a tu sección.",
    "You can complete each assigned exam only once.": "Puedes resolver cada examen asignado una sola vez.",
    "A version will be assigned automatically when you start.": "Se te asignará una versión automáticamente al iniciar.",
    "This exam uses a fixed version for your section.": "Este examen usa una versión fija para tu sección.",
    "Publish exam": "Publicar examen", "Unpublish exam": "Ocultar examen", "Archive exam": "Archivar examen",
    "Exam created.": "Examen creado.", "Exam updated.": "Examen actualizado.", "Version created.": "Versión creada.",
    "Version duplicated.": "Versión duplicada.", "Version archived.": "Versión archivada.", "Version restored.": "Versión restaurada.",
    "Archive version": "Archivar versión", "Restore version": "Restaurar versión", "Archived versions": "Versiones archivadas",
    "This version is used by an active fixed assignment. Change or remove that assignment before archiving it.": "Esta versión está siendo usada por una asignación fija activa. Cambia o elimina esa asignación antes de archivarla.",
    "A published exam with active assignments must keep at least one usable version.": "Un examen publicado con asignaciones activas debe conservar al menos una versión utilizable.",
    "Archived versions stay linked to historical attempts and can be restored later.": "Las versiones archivadas permanecen vinculadas a los intentos históricos y pueden restaurarse después.",
    "Question selection saved.": "Selección de preguntas guardada.",
    "Assignment saved.": "Asignación guardada.", "Assignment removed.": "Asignación eliminada.",
    "You already completed this exam.": "Ya completaste este examen.", "This exam is not available for your section.": "Este examen no está disponible para tu sección.",
    "No valid exam version is available.": "No hay una versión válida disponible para este examen.",
    "This version has no questions yet.": "Esta versión todavía no tiene preguntas.",
    "Questions are reusable content. Add them to exam versions from the Exams area.": "Las preguntas son contenido reutilizable. Agrégalas a versiones de examen desde el área Exámenes.",
    "Global settings": "Configuración global", "Configure the application interface and institution defaults.": "Configura la interfaz y los valores generales de la institución.",
    "Student portal": "Portal del estudiante", "Sign in to see the exams assigned to your section.": "Inicia sesión para ver los exámenes asignados a tu sección.",
    "Assessment Studio": "Assessment Studio", "Back to exams": "Volver a exámenes", "Question count": "Cantidad de preguntas",
    "Assigned sections": "Secciones asignadas", "No versions yet.": "Aún no hay versiones.", "No assignments yet.": "Aún no hay asignaciones.",
    "Select questions from the bank": "Selecciona preguntas del banco", "Selected questions": "Preguntas seleccionadas",
    "Save exam": "Guardar examen", "Exam status": "Estado del examen", "Published exams": "Exámenes publicados",
    "Results remain linked to the exact exam version used by the student.": "Los resultados quedan vinculados a la versión exacta usada por el estudiante.",
    "Version assigned": "Versión asignada", "One attempt only": "Un solo intento", "Section assignment": "Asignación por sección",
})


TRANSLATIONS_ES.update({
    "Import students": "Importar estudiantes", "Import CSV": "Importar CSV", "Download CSV template": "Descargar plantilla CSV",
    "Bulk student import": "Importación masiva de estudiantes", "CSV file": "Archivo CSV", "Choose CSV file": "Seleccionar archivo CSV",
    "Required columns": "Columnas requeridas", "NIE, First name and Last name are required. Email is optional when automatic email generation is enabled.": "NIE, Nombre y Apellido son obligatorios. Correo es opcional cuando está activada la generación automática.",
    "Generate email from NIE": "Generar correo desde NIE", "Email domain": "Dominio de correo", "Use the Correo column from the CSV": "Usar la columna Correo del CSV",
    "Batch password": "Contraseña del lote", "Generate a temporary password for this batch": "Generar una contraseña temporal para este lote",
    "Use a teacher-defined temporary password": "Usar una contraseña temporal definida por el docente", "Import students now": "Importar estudiantes ahora",
    "Bulk import complete": "Importación masiva completada", "Imported": "Importados", "Skipped": "Omitidos", "Temporary batch password": "Contraseña temporal del lote",
    "This password is shown only once. Students should change it at first sign-in.": "Esta contraseña se muestra una sola vez. Los estudiantes deben cambiarla en su primer inicio de sesión.",
    "No valid students were imported.": "No se importaron estudiantes válidos.", "The CSV must include NIE, Nombre and Apellido columns.": "El CSV debe incluir las columnas NIE, Nombre y Apellido.",
    "The CSV must include a Correo column when automatic email generation is disabled.": "El CSV debe incluir una columna Correo cuando la generación automática de correo está desactivada.",
    "Select a valid section for the import.": "Selecciona una sección válida para la importación.", "Enter a valid email domain.": "Ingresa un dominio de correo válido.",
    "Question CSV import": "Importar preguntas CSV", "Import question bank": "Importar banco de preguntas", "Bulk question import": "Importación masiva de preguntas",
    "Default subject": "Asignatura predeterminada", "Default category": "Categoría predeterminada", "Use CSV value when provided": "Usar el valor del CSV cuando exista",
    "Import questions now": "Importar preguntas ahora", "Question import complete": "Importación de preguntas completada", "Questions imported": "Preguntas importadas",
    "Rows skipped": "Filas omitidas", "Listening questions needing a source": "Preguntas de listening que requieren una fuente de audio", "Listening questions needing audio": "Preguntas de listening que requieren una fuente de audio", "Requires audio": "Requiere audio o guion",
    "Requires listening source": "Requiere audio o guion", "Listening questions without an audio file or browser-voice script cannot be added to an exam until you complete one of those sources.": "Las preguntas de listening sin archivo de audio ni guion de voz no pueden agregarse a un examen hasta completar una de esas fuentes.",
    "CSV columns": "Columnas CSV", "Use | to separate options or accepted answers, => for matching pairs, and internal type names such as multiple_choice, true_false, short_answer, numeric, order, matching or listening.": "Usa | para separar opciones o respuestas aceptadas, => para pares de relación y nombres internos de tipo como multiple_choice, true_false, short_answer, numeric, order, matching o listening.",
    "The CSV must include Tipo and Pregunta columns.": "El CSV debe incluir las columnas Tipo y Pregunta.", "No valid questions were imported.": "No se importaron preguntas válidas.",
    "Back to results": "Volver a resultados", "Back to previous page": "Volver a la página anterior", "Selected now": "Seleccionadas ahora",
    "+ Create subject": "+ Crear asignatura", "+ Create category": "+ Crear categoría", "Organization actions": "Acciones de organización",
    "Create subjects and categories only when you need them; keeping these forms in modals makes the catalog easier to scan.": "Crea asignaturas y categorías solo cuando las necesites; mantener estos formularios en modales hace el catálogo más fácil de revisar.",
    "CSV import notes": "Notas de importación CSV", "Automatic emails use the NIE as the username portion.": "Los correos automáticos usan el NIE como parte de usuario.",
    "Existing NIE or email records are skipped instead of overwritten.": "Los NIE o correos ya existentes se omiten en lugar de sobrescribirse.",
    "No categories found": "No se encontraron categorías", "No teachers found": "No se encontraron docentes", "Preview student page": "Vista del estudiante",
    "Subject and category names referenced by the CSV must already exist. If those cells are blank, choose default values below.": "Los nombres de asignatura y categoría indicados en el CSV deben existir previamente. Si esas celdas están vacías, elige valores predeterminados abajo.",
    "First name": "Nombre", "Last name": "Apellido", "Personalized email": "Correo personalizado",
    "Only .csv files are accepted.": "Solo se aceptan archivos .csv.", "The CSV file is empty.": "El archivo CSV está vacío.",
    "The CSV encoding could not be read.": "No se pudo leer la codificación del archivo CSV.", "The CSV does not contain headers.": "El CSV no contiene encabezados.",
})

DEFAULT_SETTINGS = {
    "institution_name": os.getenv("INSTITUTION_NAME", "Your Institution"),
    "assessment_title": "General Assessment",
    "assessment_subtitle": "Secure assessment with automatic grading and reusable question banks.",
    "listening_max_plays": "2",
    "current_subject_id": "1",
    "student_access_mode": "accounts",
    "ui_language": "es",
    "student_rules_es": DEFAULT_STUDENT_RULES_ES,
    "student_rules_en": DEFAULT_STUDENT_RULES_EN,
    "rules_version": DEFAULT_RULES_VERSION,
    "require_rules_acknowledgment": "1",
    "rules_updated_at": "",
}


@app.after_request
def add_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "no-referrer")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=(), display-capture=()")
    response.headers.setdefault(
        "Content-Security-Policy",
        "default-src 'self'; script-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net; style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data:; media-src 'self'; connect-src 'self' https://cdn.jsdelivr.net; frame-ancestors 'none'; "
        "object-src 'none'; base-uri 'self'; form-action 'self'"
    )
    if request.path.startswith(("/exam", "/student", "/api/integrity-event", "/api/listening-script", "/result/")):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
    return response


def get_db():
    return database_connection()


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def setting(conn, key):
    row = conn.execute("SELECT value FROM app_settings WHERE key = %s", (key,)).fetchone()
    return row["value"] if row else DEFAULT_SETTINGS.get(key, "")


def current_settings():
    with get_db() as conn:
        return {key: setting(conn, key) for key in DEFAULT_SETTINGS}


def rules_required(conn):
    return setting(conn, "require_rules_acknowledgment") == "1"


def current_rules(conn, language=None):
    language = language if language in SUPPORTED_UI_LANGUAGES else (setting(conn, "ui_language") or "es")
    return {
        "version": setting(conn, "rules_version") or "1",
        "text": setting(conn, f"student_rules_{language}"),
        "language": language,
        "updated_at": setting(conn, "rules_updated_at"),
    }


def bump_rules_version(version):
    version = str(version or DEFAULT_RULES_VERSION).strip()
    match = re.match(r"^(.*?)(\d+)$", version)
    if match:
        return f"{match.group(1)}{int(match.group(2)) + 1}"
    return f"{version}.1"


def rules_acknowledged(conn):
    if not rules_required(conn):
        return True
    rules = current_rules(conn)
    return bool(
        session.get("rules_acknowledged_version") == rules["version"]
        and session.get("rules_acknowledged_at")
    )


def get_ui_language():
    if has_request_context() and hasattr(g, "_assessment_ui_language"):
        return g._assessment_ui_language
    language = None
    if has_request_context():
        attempt_id = session.get("attempt_id") if session else None
        if attempt_id and not session.get("teacher_authenticated"):
            try:
                with get_db() as conn:
                    cols = table_columns(conn, "attempts")
                    if "ui_language" in cols:
                        row = conn.execute("SELECT ui_language FROM attempts WHERE id=%s", (attempt_id,)).fetchone()
                        language = row["ui_language"] if row and row["ui_language"] else None
            except DatabaseConnectionError:
                language = None
    if language not in SUPPORTED_UI_LANGUAGES:
        try:
            with get_db() as conn:
                language = setting(conn, "ui_language")
        except DatabaseConnectionError:
            language = DEFAULT_SETTINGS.get("ui_language", "es")
    if language not in SUPPORTED_UI_LANGUAGES:
        language = "es"
    if has_request_context():
        g._assessment_ui_language = language
    return language


def tr(text, language=None, **kwargs):
    language = language or get_ui_language()
    translated = TRANSLATIONS_ES.get(text, text) if language == "es" else text
    if kwargs:
        try:
            translated = translated.format(**kwargs)
        except (KeyError, ValueError):
            pass
    return translated


def localized_type_labels(language=None):
    language = language or get_ui_language()
    return {key: tr(label, language=language) for key, label in TYPE_LABELS.items()}


def exam_js_strings(language=None):
    language = language or get_ui_language()
    if language == "es":
        return {
            "rightClick": "El clic derecho está deshabilitado durante esta evaluación. El intento fue registrado.",
            "copy": "Copiar está deshabilitado durante esta evaluación.",
            "cut": "Cortar contenido está deshabilitado durante esta evaluación.",
            "paste": "Pegar está deshabilitado durante esta evaluación.",
            "shortcut": "Ese atajo del navegador está deshabilitado durante la evaluación.",
            "focusRecorded": "Cambio de foco #{count} registrado.",
            "focusUnavailable": "Modo enfoque no disponible",
            "focusActive": "Modo enfoque activo",
            "focusExited": "Se salió del modo enfoque. El evento fue registrado.",
            "focusReenter": "Volver al modo enfoque",
            "playAvailable": "1 reproducción disponible",
            "playsAvailable": "{count} reproducciones disponibles",
            "ttsUnsupported": "La voz del navegador no es compatible con este dispositivo.",
            "ttsLoadError": "No se pudo cargar el guion de listening.",
            "question": "Pregunta {count}",
            "allAnswered": "Las {count} preguntas tienen una respuesta. Ya puedes enviar la evaluación.",
            "missing": "Aún tienes {count} ejercicio(s) sin una respuesta completa.",
            "submitting": "Enviando…",
            "invalidCurrent": "Completa correctamente la pregunta actual antes de continuar.",
            "invalidNumeric": "Ingresa un valor numerico finito.",
            "invalidOrder": "Ordena todos los elementos antes de continuar.",
            "invalidMatching": "Completa todas las relaciones sin repetir opciones.",
        }
    return {
        "rightClick": "Right-click is disabled during this assessment. The attempt was recorded.",
        "copy": "Copying is disabled during this assessment.",
        "cut": "Cutting content is disabled during this assessment.",
        "paste": "Pasting is disabled during this assessment.",
        "shortcut": "That browser shortcut is disabled during the assessment.",
        "focusRecorded": "Focus change #{count} recorded.",
        "focusUnavailable": "Focus mode unavailable",
        "focusActive": "Focus mode active",
        "focusExited": "Focus mode was exited. The event was recorded.",
        "focusReenter": "Re-enter focus mode",
        "playAvailable": "1 play available",
        "playsAvailable": "{count} plays available",
        "ttsUnsupported": "Browser voice is not supported on this device.",
        "ttsLoadError": "Could not load the listening script.",
        "question": "Question {count}",
        "allAnswered": "All {count} questions have a response. You can submit now.",
        "missing": "You still have {count} exercise(s) without a complete response.",
        "submitting": "Submitting…",
        "invalidCurrent": "Complete the current question before continuing.",
        "invalidNumeric": "Enter a finite numeric value.",
        "invalidOrder": "Arrange every item before continuing.",
        "invalidMatching": "Complete every match without duplicate selections.",
    }


def flash_ui(message, category="message", **kwargs):
    flash(tr(message, **kwargs), category)


@app.context_processor
def inject_ui_helpers():
    language = get_ui_language()
    teacher = None
    if session.get("teacher_authenticated") and session.get("teacher_id"):
        try:
            with get_db() as conn:
                teacher = conn.execute("SELECT id,full_name,email,role FROM teachers WHERE id=%s AND is_active=1", (session.get("teacher_id"),)).fetchone()
        except DatabaseConnectionError:
            teacher = None
    return {
        "t": lambda text, **kwargs: tr(text, language=language, **kwargs),
        "ui_language": language,
        "supported_ui_languages": SUPPORTED_UI_LANGUAGES,
        "current_teacher_user": teacher,
    }


EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
PASSWORD_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789"


def normalize_email(value):
    return str(value or "").strip().casefold()[:254]


def valid_email(value):
    email = normalize_email(value)
    return bool(email and EMAIL_RE.fullmatch(email))


def valid_student_password(password):
    value = str(password or "")
    return 8 <= len(value) <= 128 and any(ch.isalpha() for ch in value) and any(ch.isdigit() for ch in value)


def generate_temp_password(length=12):
    length = max(10, min(int(length), 24))
    chars = [secrets.choice("ABCDEFGHJKLMNPQRSTUVWXYZ"), secrets.choice("abcdefghijkmnopqrstuvwxyz"), secrets.choice("23456789")]
    chars.extend(secrets.choice(PASSWORD_ALPHABET) for _ in range(length - len(chars)))
    random.SystemRandom().shuffle(chars)
    return "".join(chars)



def table_columns(conn, table):
    rows = conn.execute(
        """SELECT column_name FROM information_schema.columns
           WHERE table_schema=current_schema() AND table_name=%s""",
        (table,),
    ).fetchall()
    return {row["column_name"] for row in rows}


def bootstrap_runtime_data():
    """Create required runtime defaults after Alembic has migrated the schema."""
    os.makedirs(AUDIO_DIR, exist_ok=True)
    with get_db() as conn:
        conn.execute("SELECT pg_advisory_xact_lock(%s)", (724190315,))
        timestamp = now_iso()
        bootstrap = conn.execute("SELECT id FROM teachers ORDER BY id LIMIT 1").fetchone()
        if not bootstrap:
            email = normalize_email(app.config["TEACHER_ADMIN_EMAIL"])
            password = str(app.config["TEACHER_ADMIN_PASSWORD"] or "")
            if not valid_email(email) or not valid_student_password(password):
                raise RuntimeError("Valid TEACHER_ADMIN_EMAIL and TEACHER_ADMIN_PASSWORD are required for initial bootstrap")
            conn.execute(
                """INSERT INTO teachers(full_name,email,password_hash,role,is_active,must_change_password,created_at,updated_at)
                   VALUES(%s,%s,%s,'admin',1,0,%s,%s)""",
                (app.config["TEACHER_ADMIN_NAME"], email, generate_password_hash(password), timestamp, timestamp),
                )
            bootstrap = conn.execute("SELECT id FROM teachers ORDER BY id LIMIT 1").fetchone()
        test_teacher = None
        if os.getenv("TEST_DATABASE_URL"):
            test_teacher = conn.execute("SELECT id FROM teachers WHERE role='teacher' AND is_active=1 ORDER BY id LIMIT 1").fetchone()
            if not test_teacher:
                timestamp = now_iso()
                conn.execute(
                    """INSERT INTO teachers(full_name,email,password_hash,role,is_active,must_change_password,created_at,updated_at)
                       VALUES(%s,%s,%s,'teacher',1,0,%s,%s)""",
                    ("Test Teacher", "test.teacher@assessment.local", generate_password_hash("Teacher12345"), timestamp, timestamp),
                )
                test_teacher = conn.execute("SELECT id FROM teachers WHERE role='teacher' AND is_active=1 ORDER BY id LIMIT 1").fetchone()
        subject = conn.execute("SELECT id FROM subjects WHERE name=%s AND teacher_id=%s", ("General", bootstrap["id"])).fetchone()
        if not subject:
            subject = conn.execute(
                """INSERT INTO subjects(teacher_id,name,description,is_archived,created_at,updated_at)
                   VALUES(%s,%s,%s,0,%s,%s) RETURNING id""",
                (bootstrap["id"], "General", "Default subject. Rename it from Subjects & Categories.", timestamp, timestamp),
            ).fetchone()
        conn.execute(
            """INSERT INTO categories(teacher_id,subject_id,name,description,sort_order,is_archived,created_at,updated_at)
               VALUES(%s,%s,%s,%s,0,0,%s,%s) ON CONFLICT(subject_id,name) DO NOTHING""",
            (bootstrap["id"], subject["id"], "General", "Default category. Rename it or create additional categories.", timestamp, timestamp),
        )
        if test_teacher:
            test_subject = conn.execute("SELECT id FROM subjects WHERE teacher_id=%s AND name=%s", (test_teacher["id"], "General")).fetchone()
            if not test_subject:
                test_subject = conn.execute(
                    """INSERT INTO subjects(teacher_id,name,description,is_archived,created_at,updated_at)
                       VALUES(%s,%s,%s,0,%s,%s) RETURNING id""",
                    (test_teacher["id"], "General", "Default subject. Rename it from Subjects & Categories.", timestamp, timestamp),
                ).fetchone()
            conn.execute(
                """INSERT INTO categories(teacher_id,subject_id,name,description,sort_order,is_archived,created_at,updated_at)
                   VALUES(%s,%s,%s,%s,0,0,%s,%s) ON CONFLICT(subject_id,name) DO NOTHING""",
                (test_teacher["id"], test_subject["id"], "General", "Default category. Rename it or create additional categories.", timestamp, timestamp),
            )
        for key, value in DEFAULT_SETTINGS.items():
            conn.execute(
                "INSERT INTO app_settings(key,value) VALUES(%s,%s) ON CONFLICT(key) DO NOTHING",
                (key, value),
            )

def current_teacher_id():
    try:
        return int(session.get("teacher_id") or 0)
    except (TypeError, ValueError):
        return 0


def current_teacher(conn=None):
    teacher_id = current_teacher_id()
    if not teacher_id:
        return None
    owns = conn is None
    if owns:
        with get_db() as owned_conn:
            return owned_conn.execute("SELECT * FROM teachers WHERE id=%s AND is_active=1", (teacher_id,)).fetchone()
    return conn.execute("SELECT * FROM teachers WHERE id=%s AND is_active=1", (teacher_id,)).fetchone()


def is_teacher_admin(conn=None):
    row = current_teacher(conn)
    return bool(row and row["role"] == "admin")


def teacher_owned_question(conn, qid):
    return conn.execute("SELECT * FROM question_bank WHERE id=%s AND teacher_id=%s", (qid, current_teacher_id())).fetchone()


def current_subject(conn):
    teacher_id = current_teacher_id()
    try:
        subject_id = int(setting(conn, "current_subject_id") or 1)
    except ValueError:
        subject_id = 1
    row = conn.execute("SELECT * FROM subjects WHERE id=%s AND teacher_id=%s AND is_archived=0", (subject_id, teacher_id)).fetchone()
    if row:
        return row
    row = conn.execute("SELECT * FROM subjects WHERE teacher_id=%s AND is_archived=0 ORDER BY name LIMIT 1", (teacher_id,)).fetchone()
    return row


def row_to_question(row):
    data = json.loads(row["data_json"] or "{}")
    q = {
        "id": row["id"],
        "subject_id": int(row["subject_id"]),
        "subject_name": row["subject_name"],
        "category_id": int(row["category_id"]),
        "category_name": row["category_name"],
        "type": row["type"],
        "prompt": row["prompt"],
        "answer": json.loads(row["answer_json"]),
    }
    q.update(data)
    if row["audio"]:
        q["audio"] = row["audio"]
    if row["script"]:
        q["script"] = row["script"]
    if q["type"] == "listening":
        source = q.get("listening_source")
        if source not in {"audio", "tts"}:
            source = "audio" if row["audio"] else ("tts" if row["script"] else "audio")
        q["listening_source"] = source
        q["tts_lang"] = q.get("tts_lang") or "auto"
    return q


def listening_row_usable(row):
    if row["type"] != "listening":
        return True
    try:
        data = json.loads(row["data_json"] or "{}")
    except (TypeError, json.JSONDecodeError):
        data = {}
    source = data.get("listening_source")
    if source == "audio":
        return bool((row["audio"] or "").strip())
    if source == "tts":
        return bool((row["script"] or "").strip())
    return bool((row["audio"] or "").strip() or (row["script"] or "").strip())


def question_select_sql(extra_where=""):
    return f"""
        SELECT q.*, s.name AS subject_name, c.name AS category_name
        FROM question_bank q
        JOIN subjects s ON s.id=q.subject_id
        JOIN categories c ON c.id=q.category_id
        WHERE q.is_archived=0 AND s.is_archived=0 AND c.is_archived=0 {extra_where}
    """


def active_questions(conn=None):
    owns = conn is None
    if owns:
        with get_db() as owned_conn:
            return active_questions(owned_conn)
    subject = current_subject(conn)
    if not subject:
        return []
    rows = conn.execute(
        question_select_sql("AND q.subject_id=%s AND q.is_active=1") + " ORDER BY c.sort_order, c.name, q.type, q.created_at, q.id",
        (subject["id"],),
    ).fetchall()
    questions = [row_to_question(row) for row in rows]
    return questions


def attempt_questions(attempt):
    if attempt["questions_json"]:
        try:
            return json.loads(attempt["questions_json"])
        except (json.JSONDecodeError, TypeError):
            pass
    return []


def category_labels_for(questions):
    labels = {}
    for q in questions:
        labels[str(q["category_id"])] = q.get("category_name") or f"Category {q['category_id']}"
    return labels


def clean_question(q, rng):
    item = {
        "id": q["id"], "subject_id": q["subject_id"], "subject_name": q.get("subject_name", ""),
        "category_id": q["category_id"], "category_name": q.get("category_name", ""),
        "type": q["type"], "prompt": q["prompt"],
    }
    if q["type"] in {"multiple_choice", "listening", "true_false"}:
        choices = [list(choice) for choice in q.get("choices", [])]
        rng.shuffle(choices)
        item["choices"] = choices
    if q["type"] == "listening":
        source = q.get("listening_source") or ("audio" if q.get("audio") else "tts")
        item["listening_source"] = source
        item["tts_lang"] = q.get("tts_lang") or "auto"
        if source == "audio":
            item["audio"] = q.get("audio", "")
    if q["type"] == "order":
        items = [list(value) for value in q.get("items", [])]
        rng.shuffle(items)
        item["items"] = items
    if q["type"] == "matching":
        left = [list(value) for value in q.get("left", [])]
        right = [list(value) for value in q.get("right", [])]
        rng.shuffle(left)
        rng.shuffle(right)
        item["left"] = left
        item["right"] = right
    return item


def build_exam(seed, questions):
    rng = random.Random(seed)
    groups = {key: [] for key in TYPE_LABELS}
    for q in questions:
        if q["type"] in groups:
            groups[q["type"]].append(clean_question(q, rng))
    for items in groups.values():
        rng.shuffle(items)
    return {key: items for key, items in groups.items() if items}


SECURITY_EVENT_COLUMNS = {
    "visibility_hidden": "focus_departures",
    "focus_return": "focus_returns",
    "window_blur": "blur_events",
    "pagehide": "pagehide_events",
    "contextmenu": "context_menu_attempts",
    "copy": "copy_attempts",
    "cut": "cut_attempts",
    "paste": "paste_attempts",
    "shortcut": "shortcut_attempts",
    "fullscreen_exit": "fullscreen_exits",
}


def blocked_action_total(row):
    return sum(int(row[key] or 0) for key in (
        "context_menu_attempts", "copy_attempts", "cut_attempts", "paste_attempts", "shortcut_attempts"
    ))


def focus_change_total(row):
    return int(row["focus_departures"] or 0) + int(row["blur_events"] or 0)


def integrity_label(row):
    if focus_change_total(row) == 0 and blocked_action_total(row) == 0 and int(row["fullscreen_exits"] or 0) == 0:
        return "No incidents"
    return "Review"


def _record_value(record, key, default=None):
    try:
        value = record[key]
    except (KeyError, TypeError, IndexError):
        return default
    return default if value is None else value


def _integrity_detail(raw_detail):
    raw = str(raw_detail or "").strip()
    if not raw:
        return {}, "{}"
    try:
        parsed = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {}, raw
    return (parsed if isinstance(parsed, dict) else {}), json.dumps(parsed, ensure_ascii=False, indent=2)


def _duration_ms(value):
    try:
        duration = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(duration) or duration < 0:
        return None
    return min(duration, 86400000.0)


def format_integrity_duration(duration_ms, language="en"):
    duration = _duration_ms(duration_ms)
    if duration is None:
        return None
    total_seconds = duration / 1000.0
    decimal = f"{total_seconds:.1f}".rstrip("0").rstrip(".")
    if language == "es":
        decimal = decimal.replace(".", ",")
    if total_seconds < 60:
        if language == "es":
            unit = "segundo" if decimal == "1" else "segundos"
        else:
            unit = "second" if decimal == "1" else "seconds"
        return f"{decimal} {unit}"
    rounded_seconds = int(round(total_seconds))
    hours, remainder = divmod(rounded_seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    units = {
        "en": (("hour", "hours"), ("minute", "minutes"), ("second", "seconds")),
        "es": (("hora", "horas"), ("minuto", "minutos"), ("segundo", "segundos")),
    }["es" if language == "es" else "en"]
    parts = []
    for value, names in ((hours, units[0]), (minutes, units[1]), (seconds, units[2])):
        if value or not parts and value == seconds:
            parts.append(f"{value} {names[0] if value == 1 else names[1]}")
    return " ".join(parts)


def format_integrity_timestamp(value, language="en"):
    try:
        parsed = datetime.fromisoformat(str(value or "").replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return tr("Time unavailable", language=language)
    if language == "es":
        return parsed.strftime("%d/%m/%Y, %H:%M:%S")
    return f"{parsed.strftime('%b')} {parsed.day}, {parsed.year}, {parsed.strftime('%I:%M:%S %p').lstrip('0')}"


def _shortcut_action(shortcut, language):
    normalized = re.sub(r"\s+", "", str(shortcut or "")).casefold().replace("command", "cmd")
    key = normalized.split("+")[-1] if normalized else ""
    command = "ctrl+" in normalized or "cmd+" in normalized
    action = None
    if command:
        action = {
            "c": "copy", "x": "cut", "v": "paste", "a": "select all", "s": "save", "p": "print",
            "u": "view page source", "l": "use the address bar", "t": "open a new tab", "n": "open a new window", "w": "close the tab",
        }.get(key)
    if key == "f12" or (command and "shift+" in normalized and key in {"i", "j", "c", "k"}):
        action = "open developer tools"
    return tr(action, language=language) if action else None


def _question_presentation(attempt, language):
    questions = attempt_questions(attempt)
    try:
        groups = build_exam(_record_value(attempt, "seed", 0), questions)
        ordered = [question for items in groups.values() for question in items]
    except (KeyError, TypeError, ValueError):
        ordered = questions
    presentation = {}
    for number, question in enumerate(ordered, start=1):
        prompt = re.sub(r"\s+", " ", str(question.get("prompt", ""))).strip()
        if len(prompt) > 96:
            prompt = prompt[:93].rstrip() + "..."
        presentation[str(question.get("id", ""))] = {
            "label": tr("Question {number}", language=language, number=number),
            "excerpt": prompt,
        }
    return presentation


def integrity_event_presentation(raw_events, attempt, language="en"):
    question_labels = _question_presentation(attempt, language)
    items = []
    pending_departures = []
    summary = {"away_periods": 0, "total_away_ms": 0.0, "blocked_actions": 0, "fullscreen_exits": 0}
    blocked_types = {"contextmenu", "copy", "cut", "paste", "shortcut"}

    def base_item(event, detail, technical):
        question = question_labels.get(str(detail.get("question", "")), {})
        return {
            "raw_type": str(_record_value(event, "event_type", "")),
            "occurred_at": format_integrity_timestamp(_record_value(event, "occurred_at", ""), language),
            "question_label": question.get("label", tr("Question reference unavailable", language=language)),
            "question_excerpt": question.get("excerpt", ""),
            "technical_details": technical,
            "duration": None,
            "shortcut": None,
        }

    for event in raw_events:
        event_type = str(_record_value(event, "event_type", ""))
        detail, technical = _integrity_detail(_record_value(event, "detail_json", ""))
        item = base_item(event, detail, technical)
        duration = _duration_ms(detail.get("duration_ms"))
        if event_type == "visibility_hidden":
            summary["away_periods"] += 1
            item.update(tone="attention", icon="eye", title=tr("Left or switched away from the exam tab", language=language),
                        explanation=tr("The exam tab became hidden; no matching return was recorded.", language=language))
            pending_departures.append(len(items))
            items.append(item)
            continue
        if event_type == "focus_return" and duration is not None:
            summary["total_away_ms"] += duration
        if event_type == "focus_return" and pending_departures:
            departure = items[pending_departures.pop()]
            formatted_duration = format_integrity_duration(duration, language) if duration is not None else None
            departure.update(
                title=tr("Left the exam for {duration}", language=language, duration=formatted_duration) if formatted_duration else departure["title"],
                explanation=tr("The exam tab became hidden and a return was recorded after {duration}.", language=language, duration=formatted_duration) if formatted_duration else tr("The exam tab became hidden and a return was recorded.", language=language),
                duration=formatted_duration,
                technical_details=f"visibility_hidden\n{departure['technical_details']}\n\nfocus_return\n{technical}",
            )
            if departure["question_label"] == tr("Question reference unavailable", language=language):
                departure["question_label"] = item["question_label"]
                departure["question_excerpt"] = item["question_excerpt"]
            continue
        if event_type in blocked_types:
            summary["blocked_actions"] += 1
        if event_type == "fullscreen_exit":
            summary["fullscreen_exits"] += 1

        if event_type == "focus_return":
            item.update(tone="neutral", icon="arrow-left", title=tr("Returned to the exam", language=language),
                        explanation=tr("A return to the exam was recorded without a matching tab departure.", language=language),
                        duration=format_integrity_duration(duration, language) if duration is not None else None)
        elif event_type == "window_blur":
            item.update(tone="attention", icon="warning-circle", title=tr("The exam window lost focus", language=language), explanation=tr("The browser recorded that the exam window was no longer active.", language=language))
        elif event_type == "pagehide":
            item.update(tone="attention", icon="arrow-right", title=tr("Left or reloaded the exam page", language=language), explanation=tr("The browser recorded that the exam page was left, closed, or reloaded.", language=language))
        elif event_type == "contextmenu":
            item.update(tone="blocked", icon="warning-circle", title=tr("Blocked context menu attempt", language=language), explanation=tr("The exam blocked an attempt to open the context menu.", language=language))
        elif event_type in {"copy", "cut", "paste"}:
            titles = {"copy": "Blocked copy attempt", "cut": "Blocked cut attempt", "paste": "Blocked paste attempt"}
            explanations = {"copy": "The exam blocked an attempt to copy content.", "cut": "The exam blocked an attempt to cut content.", "paste": "The exam blocked an attempt to paste content."}
            item.update(tone="blocked", icon={"copy": "copy", "cut": "scissors", "paste": "clipboard"}[event_type], title=tr(titles[event_type], language=language), explanation=tr(explanations[event_type], language=language))
        elif event_type == "shortcut":
            shortcut = str(detail.get("shortcut", "")).strip()[:80]
            action = _shortcut_action(shortcut, language)
            item.update(tone="blocked", icon="keyboard", shortcut=shortcut or None,
                        title=tr("Blocked attempt to {action}", language=language, action=action) if action else tr("Attempted a restricted shortcut", language=language),
                        explanation=tr("The exam blocked a restricted keyboard shortcut associated with {action}.", language=language, action=action) if action else tr("The exam blocked a restricted keyboard shortcut.", language=language))
        elif event_type == "fullscreen_exit":
            item.update(tone="attention", icon="arrows-out", title=tr("Left fullscreen", language=language), explanation=tr("The browser recorded that fullscreen mode ended during the exam.", language=language))
        else:
            item.update(tone="neutral", icon="fingerprint", title=tr("Recorded integrity event", language=language), explanation=tr("The browser recorded an event that is available for technical review.", language=language))
        items.append(item)

    summary["total_away_time"] = format_integrity_duration(summary["total_away_ms"], language)
    return {"items": items, "summary": summary}


def merge_security_snapshot(conn, attempt_id, form):
    fields = {
        "focus_departures": "security_focus_departures",
        "blur_events": "security_blur_events",
        "pagehide_events": "security_pagehide_events",
        "context_menu_attempts": "security_contextmenu_attempts",
        "copy_attempts": "security_copy_attempts",
        "cut_attempts": "security_cut_attempts",
        "paste_attempts": "security_paste_attempts",
        "shortcut_attempts": "security_shortcut_attempts",
        "fullscreen_exits": "security_fullscreen_exits",
    }
    values = {}
    for column, form_name in fields.items():
        try:
            values[column] = max(0, int(form.get(form_name, 0)))
        except (TypeError, ValueError):
            values[column] = 0
    try:
        away = max(0.0, min(float(form.get("security_away_seconds", 0)), 86400.0))
    except (TypeError, ValueError):
        away = 0.0
    try:
        longest = max(0.0, min(float(form.get("security_longest_away_seconds", 0)), 86400.0))
    except (TypeError, ValueError):
        longest = 0.0
    assignments = [f"{column} = GREATEST({column}, %s)" for column in values]
    params = list(values.values())
    assignments += ["away_seconds = GREATEST(away_seconds, %s)", "longest_away_seconds = GREATEST(longest_away_seconds, %s)"]
    params += [away, longest, attempt_id]
    conn.execute(f"UPDATE attempts SET {', '.join(assignments)} WHERE id = %s", params)


def teacher_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        teacher_id = current_teacher_id()
        if not session.get("teacher_authenticated") or not teacher_id:
            return redirect(url_for("teacher"))
        with get_db() as conn:
            row = conn.execute("SELECT id,full_name,email,role,is_active FROM teachers WHERE id=%s", (teacher_id,)).fetchone()
        if not row or not row["is_active"]:
            session.clear()
            return redirect(url_for("teacher"))
        session["teacher_role"] = row["role"]
        session["teacher_name"] = row["full_name"]
        return view(*args, **kwargs)
    return wrapped


def admin_required(view):
    @wraps(view)
    @teacher_required
    def wrapped(*args, **kwargs):
        if session.get("teacher_role") != "admin":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def csrf_token():
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_urlsafe(24)
        session["csrf_token"] = token
    return token


app.jinja_env.globals["csrf_token"] = csrf_token


def verify_csrf():
    supplied = request.form.get("csrf_token", "")
    expected = session.get("csrf_token", "")
    if not supplied or not expected or not secrets.compare_digest(supplied, expected):
        abort(403)


def normalized_text(value, case_sensitive=False):
    clean = " ".join(str(value or "").strip().split())
    return clean if case_sensitive else clean.casefold()


def score_attempt(form, questions):
    score = 0.0
    category_scores = {}
    type_scores = {t: {"earned": 0.0, "total": 0.0} for t in TYPE_LABELS}
    captured = {}

    for q in questions:
        earned = 0.0
        category_key = str(q["category_id"])
        category_scores.setdefault(category_key, {"earned": 0.0, "total": 0.0})
        category_scores[category_key]["total"] += 1
        type_scores.setdefault(q["type"], {"earned": 0.0, "total": 0.0})
        type_scores[q["type"]]["total"] += 1

        if q["type"] in {"multiple_choice", "listening", "true_false"}:
            value = form.get(f"q_{q['id']}", "")
            captured[q["id"]] = value
            if value == q["answer"]:
                earned = 1.0
        elif q["type"] == "short_answer":
            value = form.get(f"q_{q['id']}", "")
            captured[q["id"]] = value
            case_sensitive = bool(q.get("case_sensitive", False))
            accepted = q.get("answer") if isinstance(q.get("answer"), list) else [q.get("answer")]
            target = normalized_text(value, case_sensitive)
            if target and target in {normalized_text(ans, case_sensitive) for ans in accepted}:
                earned = 1.0
        elif q["type"] == "numeric":
            raw = form.get(f"q_{q['id']}", "").strip()
            captured[q["id"]] = raw
            try:
                value = float(raw.replace(",", "."))
                answer = q.get("answer") or {}
                correct = float(answer.get("value"))
                tolerance = max(0.0, float(answer.get("tolerance", 0)))
                if abs(value - correct) <= tolerance:
                    earned = 1.0
            except (TypeError, ValueError):
                pass
        elif q["type"] == "order":
            raw = form.get(f"q_{q['id']}", "[]")
            try:
                value = json.loads(raw)
            except json.JSONDecodeError:
                value = []
            captured[q["id"]] = value
            if value == q["answer"]:
                earned = 1.0
        elif q["type"] == "matching":
            response = {}
            for left_key, _ in q.get("left", []):
                response[left_key] = form.get(f"q_{q['id']}__{left_key}", "")
            captured[q["id"]] = response
            if response == q["answer"]:
                earned = 1.0

        score += earned
        category_scores[category_key]["earned"] += earned
        type_scores[q["type"]]["earned"] += earned

    total = float(len(questions))
    percentage = round((score / total) * 100, 1) if total else 0.0
    grade10 = round((score / total) * 10, 1) if total else 0.0
    return round(score, 1), total, percentage, grade10, category_scores, type_scores, captured


def validate_required_answers(form, questions):
    """Validate and capture answers using the attempt snapshot as the only authority."""
    captured = {}
    invalid = []
    for q in questions:
        qid = str(q["id"])
        qtype = q.get("type")
        field = f"q_{qid}"
        valid = False
        if qtype in {"multiple_choice", "listening", "true_false"}:
            value = str(form.get(field, ""))
            allowed = {str(item[0]) for item in q.get("choices", [])}
            captured[qid] = value
            valid = bool(value) and value in allowed
        elif qtype == "short_answer":
            value = str(form.get(field, ""))
            captured[qid] = value
            valid = bool(value.strip())
        elif qtype == "numeric":
            value = str(form.get(field, ""))
            captured[qid] = value
            try:
                valid = math.isfinite(float(value.strip().replace(",", ".")))
            except (TypeError, ValueError):
                valid = False
        elif qtype == "order":
            raw = str(form.get(field, ""))
            touched = form.get(f"{field}__touched") == "1"
            try:
                value = json.loads(raw)
            except (TypeError, json.JSONDecodeError):
                value = []
            expected = [str(item[0]) for item in q.get("items", [])]
            normalized = [str(item) for item in value] if isinstance(value, list) else []
            captured[qid] = normalized
            captured[f"{qid}__touched"] = "1" if touched else "0"
            valid = touched and len(normalized) == len(expected) and len(set(normalized)) == len(normalized) and set(normalized) == set(expected)
        elif qtype == "matching":
            expected_left = [str(item[0]) for item in q.get("left", [])]
            allowed = {str(item[0]) for item in q.get("right", [])}
            response = {left: str(form.get(f"{field}__{left}", "")) for left in expected_left}
            captured[qid] = response
            values = list(response.values())
            valid = bool(expected_left) and all(value in allowed and value for value in values) and len(set(values)) == len(values)
        if not valid:
            invalid.append(qid)
    return invalid, captured


def active_penalty_total(conn, attempt_id):
    row = conn.execute(
        "SELECT COALESCE(SUM(points),0) AS total FROM attempt_penalties WHERE attempt_id=%s AND is_active=1",
        (attempt_id,),
    ).fetchone()
    return round(float(row["total"] or 0), 2)


def adjusted_grade10(raw_grade, penalty_total):
    return round(max(0.0, float(raw_grade or 0) - float(penalty_total or 0)), 2)


def penalties_for_attempt(conn, attempt_id):
    return conn.execute(
        """SELECT p.*,t.full_name AS teacher_name,rv.full_name AS revoked_by_name
           FROM attempt_penalties p JOIN teachers t ON t.id=p.teacher_id
           LEFT JOIN teachers rv ON rv.id=p.revoked_by
           WHERE p.attempt_id=%s ORDER BY p.created_at DESC,p.id DESC""",
        (attempt_id,),
    ).fetchall()


def version_questions(conn, version_id):
    rows = conn.execute(
        """SELECT q.*, s.name AS subject_name, c.name AS category_name
           FROM exam_version_questions evq
           JOIN question_bank q ON q.id=evq.question_id
           JOIN subjects s ON s.id=q.subject_id
           JOIN categories c ON c.id=q.category_id
           WHERE evq.version_id=%s
           ORDER BY evq.position, q.created_at, q.id""",
        (version_id,),
    ).fetchall()
    return [row_to_question(row) for row in rows]


def student_exam_rows(conn, student_id):
    student = conn.execute(
        """SELECT st.*, sec.name AS section_name FROM students st
           JOIN sections sec ON sec.id=st.section_id
           WHERE st.id=%s AND st.is_active=1 AND st.is_archived=0 AND sec.is_archived=0""",
        (student_id,),
    ).fetchone()
    if not student:
        return None, []
    rows = conn.execute(
        """SELECT ea.*, e.title, e.description, e.is_published, e.subject_id, s.name AS subject_name,
                  e.teacher_id AS teacher_id,
                  ev.name AS fixed_version_name,
                  (SELECT COUNT(*) FROM exam_versions v WHERE v.exam_id=e.id AND v.is_active=1
                     AND EXISTS(SELECT 1 FROM exam_version_questions qx WHERE qx.version_id=v.id)) AS version_count,
                   a.id AS attempt_id, a.status AS attempt_status, a.grade10, a.percentage,
                   COALESCE((SELECT SUM(p.points) FROM attempt_penalties p WHERE p.attempt_id=a.id AND p.is_active=1),0) AS penalty_points,
                   GREATEST(0,COALESCE(a.grade10,0)-COALESCE((SELECT SUM(p.points) FROM attempt_penalties p WHERE p.attempt_id=a.id AND p.is_active=1),0)) AS adjusted_grade10,
                  av.name AS assigned_version_name
           FROM exam_assignments ea
           JOIN exams e ON e.id=ea.exam_id
           JOIN subjects s ON s.id=e.subject_id
           LEFT JOIN exam_versions ev ON ev.id=ea.fixed_version_id
           LEFT JOIN attempts a ON a.assignment_id=ea.id AND a.student_id=%s
           LEFT JOIN exam_versions av ON av.id=a.exam_version_id
           WHERE ea.section_id=%s AND ea.is_active=1 AND e.is_archived=0
             AND (e.is_published=1 OR a.id IS NOT NULL)
           ORDER BY e.title , ea.id""",
        (student_id, student["section_id"]),
    ).fetchall()
    return student, rows


def teacher_report_scope(conn=None):
    teacher = current_teacher(conn)
    return teacher, bool(teacher and teacher["role"] == "admin")


def teacher_reports_section_rows(conn):
    _, is_admin = teacher_report_scope(conn)
    teacher_clause = "" if is_admin else "AND e.teacher_id=%s"
    params = () if is_admin else (current_teacher_id(),)
    sections = conn.execute(
        f"""SELECT sec.id, sec.name, sec.description,
                  COUNT(DISTINCT st.id) AS student_count,
                  COUNT(DISTINCT s.id) AS subject_count,
                  COUNT(DISTINCT ea.id) AS exam_count
             FROM sections sec
             LEFT JOIN students st ON st.section_id=sec.id AND COALESCE(st.is_archived,0)=0
             LEFT JOIN exam_assignments ea ON ea.section_id=sec.id AND ea.is_active=1
             LEFT JOIN exams e ON e.id=ea.exam_id AND e.is_archived=0 {teacher_clause}
             LEFT JOIN subjects s ON s.id=e.subject_id AND s.is_archived=0
            WHERE sec.is_archived=0
         GROUP BY sec.id
         ORDER BY sec.name""",
        params,
    ).fetchall()
    if is_admin:
        return sections, is_admin
    return [row for row in sections if int(row["subject_count"] or 0) > 0], is_admin


def teacher_section_subject_rows(conn, section_id):
    teacher, is_admin = teacher_report_scope(conn)
    section = conn.execute(
        """SELECT sec.*, COUNT(st.id) AS student_count,
                  SUM(CASE WHEN st.is_active=1 THEN 1 ELSE 0 END) AS active_count
           FROM sections sec LEFT JOIN students st ON st.section_id=sec.id AND COALESCE(st.is_archived,0)=0
           WHERE sec.id=%s AND sec.is_archived=0 GROUP BY sec.id""",
        (section_id,),
    ).fetchone()
    if not section:
        return None, [], False
    teacher_clause = "" if is_admin else "AND e.teacher_id=%s"
    params = [section_id]
    if not is_admin:
        params.append(current_teacher_id())
    subjects = conn.execute(
        f"""SELECT s.id, s.name, s.description, COUNT(DISTINCT ea.id) AS exam_count
             FROM exam_assignments ea
             JOIN exams e ON e.id=ea.exam_id
             JOIN subjects s ON s.id=e.subject_id
            WHERE ea.section_id=%s AND ea.is_active=1 AND e.is_archived=0 {teacher_clause}
         GROUP BY s.id, s.name, s.description
         ORDER BY s.name""",
        tuple(params),
    ).fetchall()
    return section, subjects, is_admin


def teacher_section_subject_report_rows(conn, section_id, subject_id):
    teacher, is_admin = teacher_report_scope(conn)
    section = conn.execute(
        """SELECT sec.*, COUNT(st.id) AS student_count,
                  SUM(CASE WHEN st.is_active=1 THEN 1 ELSE 0 END) AS active_count
           FROM sections sec LEFT JOIN students st ON st.section_id=sec.id AND COALESCE(st.is_archived,0)=0
           WHERE sec.id=%s AND sec.is_archived=0 GROUP BY sec.id""",
        (section_id,),
    ).fetchone()
    if not section:
        return None, None, [], [], False
    subject = conn.execute(
        "SELECT * FROM subjects WHERE id=%s AND is_archived=0 AND teacher_id=%s" if not is_admin else "SELECT * FROM subjects WHERE id=%s AND is_archived=0",
        (subject_id, current_teacher_id()) if not is_admin else (subject_id,),
    ).fetchone()
    if not subject:
        return section, None, [], [], False
    teacher_clause = "" if is_admin else "AND e.teacher_id=%s"
    params = [section_id, subject_id]
    if not is_admin:
        params.append(current_teacher_id())
    exams = conn.execute(
        f"""SELECT ea.id AS assignment_id, ea.exam_id, e.title, ea.fixed_version_id
             FROM exam_assignments ea
             JOIN exams e ON e.id=ea.exam_id
            WHERE ea.section_id=%s AND e.subject_id=%s AND ea.is_active=1 AND e.is_archived=0 {teacher_clause}
         ORDER BY e.title, ea.id""",
        tuple(params),
    ).fetchall()
    students = conn.execute(
        """SELECT st.id, st.full_name, st.student_code, st.email
             FROM students st
            WHERE st.section_id=%s AND st.is_archived=0
          ORDER BY st.full_name""",
        (section_id,),
    ).fetchall()
    evaluation = exams[0] if exams else None
    if not evaluation:
        return section, subject, None, [], is_admin
    attempts = conn.execute(
        f"""SELECT a.id AS attempt_id, a.student_id, a.assignment_id, a.status AS attempt_status,
                  a.grade10, a.percentage, a.submitted_at,
                  GREATEST(0, COALESCE(a.grade10,0) - COALESCE((SELECT SUM(p.points) FROM attempt_penalties p WHERE p.attempt_id=a.id AND p.is_active=1),0)) AS adjusted_grade10
             FROM attempts a
             JOIN exam_assignments ea ON ea.id=a.assignment_id
             JOIN exams e ON e.id=ea.exam_id
             WHERE ea.section_id=%s AND e.subject_id=%s AND ea.id=%s AND a.status='submitted' AND ea.is_active=1 AND e.is_archived=0 {teacher_clause}""",
        tuple([section_id, subject_id, evaluation["assignment_id"]] + ([] if is_admin else [current_teacher_id()])),
    ).fetchall()
    attempts_by_key = {(int(row["student_id"]), int(row["assignment_id"])): row for row in attempts}
    rows = []
    for student in students:
        attempt = attempts_by_key.get((int(student["id"]), int(evaluation["assignment_id"])))
        rows.append({
            "id": int(student["id"]),
            "student_code": student["student_code"],
            "full_name": student["full_name"],
            "grade10": attempt["grade10"] if attempt else None,
            "adjusted_grade10": attempt["adjusted_grade10"] if attempt else None,
            "attempt_id": attempt["attempt_id"] if attempt else None,
            "attempt_status": attempt["attempt_status"] if attempt else None,
        })
    return section, subject, evaluation, rows, is_admin


def teacher_student_report_rows(conn, student_id):
    student, rows = student_exam_rows(conn, student_id)
    if not student:
        return None, [], False
    teacher, is_admin = teacher_report_scope(conn)
    if not is_admin:
        rows = [row for row in rows if int(row["teacher_id"]) == current_teacher_id()]
    return student, rows, is_admin


def teacher_section_report_rows(conn, section_id):
    teacher, is_admin = teacher_report_scope(conn)
    section = conn.execute(
        """SELECT sec.*, COUNT(st.id) AS student_count,
                  SUM(CASE WHEN st.is_active=1 THEN 1 ELSE 0 END) AS active_count
           FROM sections sec LEFT JOIN students st ON st.section_id=sec.id AND COALESCE(st.is_archived,0)=0
           WHERE sec.id=%s AND sec.is_archived=0 GROUP BY sec.id""",
        (section_id,),
    ).fetchone()
    if not section:
        return None, 0, [], False
    teacher_clause = "" if is_admin else "AND e.teacher_id=%s"
    visible_params = [section_id]
    if not is_admin:
        visible_params.append(current_teacher_id())
    visible_exam_count = conn.execute(
        f"""SELECT COUNT(*) AS n
            FROM exam_assignments ea JOIN exams e ON e.id=ea.exam_id
            WHERE ea.section_id=%s AND ea.is_active=1 AND e.is_archived=0 {teacher_clause}""",
        tuple(visible_params),
    ).fetchone()["n"]
    student_rows = conn.execute(
        f"""WITH visible_assignments AS (
                SELECT ea.id, ea.section_id
                FROM exam_assignments ea JOIN exams e ON e.id=ea.exam_id
                WHERE ea.section_id=%s AND ea.is_active=1 AND e.is_archived=0 {teacher_clause}
           )
           SELECT st.id, st.full_name, st.student_code, st.email, st.is_active,
                  COUNT(DISTINCT va.id) AS assigned_exams,
                  COUNT(a.id) AS submitted_exams,
                  AVG(a.percentage) AS avg_percentage,
                  MAX(a.submitted_at) AS last_submitted_at
           FROM students st
           LEFT JOIN visible_assignments va ON va.section_id=st.section_id
           LEFT JOIN attempts a ON a.assignment_id=va.id AND a.student_id=st.id AND a.status='submitted'
           WHERE st.section_id=%s AND st.is_archived=0
           GROUP BY st.id, st.full_name, st.student_code, st.email, st.is_active
           ORDER BY st.full_name""",
        tuple(visible_params + [section_id]),
    ).fetchall()
    return section, visible_exam_count, student_rows, is_admin


def student_dashboard_stats(exams):
    stats = {
        "total_assigned": len(exams), "available": 0, "in_progress": 0, "completed": 0,
        "completion_rate": 0.0, "average_grade": None, "pending": 0, "next_exam": None,
    }
    submitted_grades = []
    first_available = None
    for item in exams:
        status = item["attempt_status"]
        if status == "submitted":
            stats["completed"] += 1
            grade = item["adjusted_grade10"] if "adjusted_grade10" in item.keys() else item["grade10"]
            if grade is not None:
                submitted_grades.append(float(grade))
        elif status == "in_progress":
            stats["in_progress"] += 1
            if stats["next_exam"] is None:
                stats["next_exam"] = item
        else:
            stats["available"] += 1
            if first_available is None:
                first_available = item
    if stats["next_exam"] is None:
        stats["next_exam"] = first_available
    stats["pending"] = stats["available"] + stats["in_progress"]
    if stats["total_assigned"]:
        stats["completion_rate"] = stats["completed"] / stats["total_assigned"] * 100
    if submitted_grades:
        stats["average_grade"] = sum(submitted_grades) / len(submitted_grades)
    return stats


@app.get("/")
def index():
    if session.get("student_authenticated") and session.get("student_id"):
        return redirect(url_for("student_dashboard"))
    settings = current_settings()
    return render_template(
        "index.html", title="Assessment Studio", institution=settings["institution_name"],
        subtitle=settings["assessment_subtitle"].strip(),
    )


@app.get("/health/live")
def health_live():
    return jsonify({"status": "ok"})


@app.get("/health/ready")
def health_ready():
    try:
        with get_db() as conn:
            conn.execute("SELECT 1").fetchone()
    except (DatabaseConnectionError, PoolTimeout):
        return jsonify({"status": "unavailable"}), 503
    return jsonify({"status": "ok"})


@app.post("/start")
def start():
    verify_csrf()
    email = normalize_email(request.form.get("email"))
    password = request.form.get("password", "")
    if not valid_email(email) or not password:
        flash_ui("Enter your email and password to begin.", "error")
        return redirect(url_for("index"))
    with get_db() as conn:
        student = conn.execute(
            """SELECT st.*, sec.name AS section_name
               FROM students st JOIN sections sec ON sec.id=st.section_id
               WHERE st.email=%s  AND st.is_active=1 AND st.is_archived=0 AND sec.is_archived=0""",
            (email,),
        ).fetchone()
        if not student or not student["password_hash"] or not check_password_hash(student["password_hash"], password):
            flash_ui("Invalid email or password, or the account is inactive.", "error")
            return redirect(url_for("index"))
        conn.execute("UPDATE students SET last_login_at=%s,updated_at=%s WHERE id=%s", (now_iso(), now_iso(), student["id"]))
        conn.commit()
        student_id = int(student["id"])
        must_change = bool(student["must_change_password"])
    session.clear()
    session["student_id"] = student_id
    session["student_authenticated"] = True
    csrf_token()
    if must_change:
        flash_ui("Your password must be changed before you start the assessment.", "message")
        return redirect(url_for("student_change_password"))
    return redirect(url_for("student_dashboard"))


@app.get("/student")
def student_dashboard():
    student_id = session.get("student_id")
    if not session.get("student_authenticated") or not student_id:
        return redirect(url_for("index"))
    with get_db() as conn:
        student, exams = student_exam_rows(conn, int(student_id))
        institution = setting(conn, "institution_name")
        rules = current_rules(conn, get_ui_language())
        require_rules = rules_required(conn) and not rules_acknowledged(conn)
    if not student:
        session.clear()
        return redirect(url_for("index"))
    return render_template(
        "student/index.html", student=student, exams=exams, institution=institution,
        dashboard_stats=student_dashboard_stats(exams),
        rules=rules, require_rules=require_rules,
    )


@app.post("/student/rules/acknowledge")
def acknowledge_student_rules():
    verify_csrf()
    if not session.get("student_authenticated") or not session.get("student_id"):
        abort(403)
    with get_db() as conn:
        rules = current_rules(conn)
        if rules_required(conn):
            if request.form.get("rules_version") != rules["version"] or request.form.get("accept_rules") != "1":
                flash_ui("You must accept the current usage rules before continuing.", "error")
                return redirect(url_for("student_dashboard"))
    session["rules_acknowledged_version"] = rules["version"]
    session["rules_acknowledged_at"] = now_iso()
    return redirect(url_for("student_dashboard"))


def choose_assignment_version(conn, assignment, student_id):
    existing = conn.execute(
        "SELECT version_id FROM student_exam_allocations WHERE assignment_id=%s AND student_id=%s",
        (assignment["id"], student_id),
    ).fetchone()
    if existing:
        return int(existing["version_id"])
    if assignment["version_mode"] == "fixed":
        version_id = assignment["fixed_version_id"]
        valid = conn.execute(
            """SELECT 1 FROM exam_versions v WHERE v.id=%s AND v.exam_id=%s AND v.is_active=1
               AND EXISTS(SELECT 1 FROM exam_version_questions q WHERE q.version_id=v.id)""",
            (version_id, assignment["exam_id"]),
        ).fetchone()
        if not valid:
            return None
    else:
        candidate_rows = conn.execute(
            """SELECT v.id,
                      (SELECT COUNT(*) FROM student_exam_allocations sea
                       WHERE sea.assignment_id=%s AND sea.version_id=v.id) AS allocated_count
               FROM exam_versions v WHERE v.exam_id=%s AND v.is_active=1
                 AND EXISTS(SELECT 1 FROM exam_version_questions q WHERE q.version_id=v.id)
               ORDER BY v.id""",
            (assignment["id"], assignment["exam_id"]),
        ).fetchall()
        if not candidate_rows:
            return None
        minimum = min(int(r["allocated_count"] or 0) for r in candidate_rows)
        candidates = [int(r["id"]) for r in candidate_rows if int(r["allocated_count"] or 0) == minimum]
        version_id = secrets.choice(candidates)
    conn.execute(
        "INSERT INTO student_exam_allocations(assignment_id,student_id,version_id,allocated_at) VALUES(%s,%s,%s,%s)",
        (assignment["id"], student_id, version_id, now_iso()),
    )
    return int(version_id)


@app.post("/student/exams/<int:assignment_id>/start")
def start_assigned_exam(assignment_id):
    verify_csrf()
    student_id = session.get("student_id")
    if not session.get("student_authenticated") or not student_id:
        return redirect(url_for("index"))
    with get_db() as conn:
        if not rules_acknowledged(conn):
            flash_ui("You must accept the current usage rules before starting or resuming an exam.", "error")
            return redirect(url_for("student_dashboard"))
        student = conn.execute(
            """SELECT st.*,sec.name AS section_name FROM students st JOIN sections sec ON sec.id=st.section_id
               WHERE st.id=%s AND st.is_active=1 AND st.is_archived=0 AND sec.is_archived=0""",
            (student_id,),
        ).fetchone()
        assignment = conn.execute(
            """SELECT ea.*,e.title,e.description,e.subject_id,e.teacher_id,e.is_published,e.is_archived,s.name AS subject_name
               FROM exam_assignments ea JOIN exams e ON e.id=ea.exam_id JOIN subjects s ON s.id=e.subject_id
               WHERE ea.id=%s AND ea.section_id=%s AND ea.is_active=1 FOR UPDATE OF ea""",
            (assignment_id, student["section_id"] if student else -1),
        ).fetchone()
        if not student or not assignment or assignment["is_archived"] or not assignment["is_published"]:
            flash_ui("This exam is not available for your section.", "error")
            return redirect(url_for("student_dashboard"))
        previous = conn.execute(
            "SELECT * FROM attempts WHERE student_id=%s AND assignment_id=%s ORDER BY started_at DESC LIMIT 1",
            (student_id, assignment_id),
        ).fetchone()
        if previous:
            session["attempt_id"] = previous["id"]
            if previous["status"] == "submitted":
                flash_ui("You already completed this exam.", "message")
                return redirect(url_for("result", attempt_id=previous["id"]))
            return redirect(url_for("exam"))
        version_id = choose_assignment_version(conn, assignment, int(student_id))
        if not version_id:
            conn.rollback()
            flash_ui("No valid exam version is available.", "error")
            return redirect(url_for("student_dashboard"))
        version = conn.execute("SELECT * FROM exam_versions WHERE id=%s", (version_id,)).fetchone()
        questions = version_questions(conn, version_id)
        if not questions:
            conn.rollback()
            flash_ui("This version has no questions yet.", "error")
            return redirect(url_for("student_dashboard"))
        attempt_id = secrets.token_urlsafe(10)
        seed = secrets.randbelow(2_000_000_000)
        labels = category_labels_for(questions)
        rules = current_rules(conn, setting(conn, "ui_language"))
        accepted_at = session.get("rules_acknowledged_at")
        try:
            conn.execute(
                """INSERT INTO attempts
                   (id,student_id,student_email,student_name,section,started_at,seed,status,questions_json,
                     assessment_title,assessment_subject,category_labels_json,ui_language,exam_id,exam_version_id,assignment_id,exam_version_name,teacher_id,
                     policy_version,policy_accepted_at,policy_text)
                   VALUES(%s,%s,%s,%s,%s,%s,%s,'in_progress',%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                (attempt_id, student["id"], student["email"], student["full_name"], student["section_name"], now_iso(), seed,
                 json.dumps(questions), assignment["title"], assignment["subject_name"], json.dumps(labels),
                  setting(conn, "ui_language") or "es", assignment["exam_id"], version_id, assignment_id, version["name"], assignment["teacher_id"],
                  rules["version"], accepted_at, rules["text"]),
            )
            conn.commit()
        except IntegrityError:
            conn.rollback()
            previous = conn.execute(
                "SELECT * FROM attempts WHERE student_id=%s AND assignment_id=%s ORDER BY started_at DESC LIMIT 1",
                (student_id, assignment_id),
            ).fetchone()
            if not previous:
                raise
            attempt_id = previous["id"]
    session["attempt_id"] = attempt_id
    csrf_token()
    return redirect(url_for("exam"))


@app.route("/student/password", methods=["GET", "POST"])
def student_change_password():
    student_id = session.get("student_id")
    if not session.get("student_authenticated") or not student_id:
        return redirect(url_for("index"))
    with get_db() as conn:
        student = conn.execute(
            """SELECT st.*,sec.name AS section_name FROM students st JOIN sections sec ON sec.id=st.section_id WHERE st.id=%s""",
            (student_id,),
        ).fetchone()
    if not student:
        session.clear()
        return redirect(url_for("index"))

    forced = bool(student["must_change_password"])
    if request.method == "POST":
        verify_csrf()
        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")
        if not forced and (not student["password_hash"] or not check_password_hash(student["password_hash"], current_password)):
            flash_ui("Your current password is incorrect.", "error")
        elif new_password != confirm_password:
            flash_ui("The new passwords do not match.", "error")
        elif not valid_student_password(new_password):
            flash_ui("Password must be at least 8 characters and include at least one letter and one number.", "error")
        else:
            with get_db() as conn:
                conn.execute(
                    "UPDATE students SET password_hash=%s,must_change_password=0,password_updated_at=%s,updated_at=%s WHERE id=%s",
                    (generate_password_hash(new_password), now_iso(), now_iso(), student_id),
                )
                conn.commit()
            flash_ui("Password changed. You can continue.", "success")
            return redirect(url_for("student_dashboard"))
    return render_template("student/password.html", student=student, forced=forced)


@app.post("/student/logout")
def student_logout():
    verify_csrf()
    session.clear()
    return redirect(url_for("index"))


@app.get("/exam")
def exam():
    attempt_id = session.get("attempt_id")
    student_id = session.get("student_id")
    if not attempt_id or not session.get("student_authenticated") or not student_id:
        return redirect(url_for("index"))
    with get_db() as conn:
        if not rules_acknowledged(conn):
            flash_ui("You must accept the current usage rules before starting or resuming an exam.", "error")
            return redirect(url_for("student_dashboard"))
        attempt = conn.execute("SELECT * FROM attempts WHERE id=%s AND student_id=%s", (attempt_id, student_id)).fetchone()
        listening_max_plays = int(setting(conn, "listening_max_plays") or 2)
    if not attempt:
        session.clear()
        return redirect(url_for("index"))
    if attempt["status"] == "submitted":
        return redirect(url_for("result", attempt_id=attempt_id))
    questions = attempt_questions(attempt)
    groups = build_exam(attempt["seed"], questions)
    return render_template(
        "exam/index.html", attempt=attempt, groups=groups, type_labels=localized_type_labels(),
        total=len(questions), title=attempt["assessment_title"] or DEFAULT_SETTINGS["assessment_title"],
        listening_max_plays=max(1, min(listening_max_plays, 5)),
        exam_js_i18n=exam_js_strings(),
        server_draft=json.loads(attempt["answers_json"] or "{}"),
        invalid_question=session.pop("invalid_question", None),
    )


@app.get("/api/listening-script/<question_id>")
def listening_script(question_id):
    attempt_id = session.get("attempt_id")
    student_id = session.get("student_id")
    if not attempt_id or not session.get("student_authenticated") or not student_id:
        return jsonify({"ok": False}), 403
    with get_db() as conn:
        attempt = conn.execute("SELECT * FROM attempts WHERE id=%s AND student_id=%s", (attempt_id, student_id)).fetchone()
    if not attempt or attempt["status"] != "in_progress":
        return jsonify({"ok": False}), 403
    for q in attempt_questions(attempt):
        if str(q.get("id")) != str(question_id) or q.get("type") != "listening":
            continue
        source = q.get("listening_source") or ("audio" if q.get("audio") else "tts")
        script = (q.get("script") or "").strip()
        if source != "tts" or not script:
            return jsonify({"ok": False}), 404
        return jsonify({"ok": True, "text": script, "lang": q.get("tts_lang") or "auto"})
    return jsonify({"ok": False}), 404


@app.post("/api/integrity-event")
def integrity_event():
    attempt_id = session.get("attempt_id")
    student_id = session.get("student_id")
    if not attempt_id or not session.get("student_authenticated") or not student_id:
        return jsonify({"ok": False}), 403
    payload = request.get_json(silent=True) or {}
    event_type = str(payload.get("type", ""))[:40]
    if event_type not in SECURITY_EVENT_COLUMNS:
        return jsonify({"ok": False, "error": "unsupported_event"}), 400
    timestamp = now_iso()
    detail = {k: payload.get(k) for k in ("question", "client_time", "visibility", "duration_ms", "shortcut") if payload.get(k) is not None}
    with get_db() as conn:
        attempt = conn.execute(
            "SELECT status FROM attempts WHERE id=%s AND student_id=%s FOR UPDATE",
            (attempt_id, student_id),
        ).fetchone()
        if not attempt or attempt["status"] != "in_progress":
            return jsonify({"ok": False}), 409
        column = SECURITY_EVENT_COLUMNS[event_type]
        conn.execute(f"UPDATE attempts SET {column}={column}+1, last_security_event_at=%s WHERE id=%s", (timestamp, attempt_id))
        if event_type == "focus_return":
            try:
                duration = max(0.0, min(float(payload.get("duration_ms", 0)) / 1000.0, 86400.0))
            except (TypeError, ValueError):
                duration = 0.0
            conn.execute(
                "UPDATE attempts SET away_seconds=away_seconds+%s, longest_away_seconds=GREATEST(longest_away_seconds, %s) WHERE id=%s",
                (duration, duration, attempt_id),
            )
        conn.execute(
            "INSERT INTO integrity_events(attempt_id,event_type,occurred_at,detail_json) VALUES(%s,%s,%s,%s)",
            (attempt_id, event_type, timestamp, json.dumps(detail)),
        )
        conn.commit()
    return jsonify({"ok": True})


@app.post("/submit")
def submit_exam():
    verify_csrf()
    attempt_id = session.get("attempt_id")
    student_id = session.get("student_id")
    if not attempt_id or not session.get("student_authenticated") or not student_id:
        abort(403)
    with get_db() as conn:
        attempt = conn.execute(
            "SELECT * FROM attempts WHERE id=%s AND student_id=%s FOR UPDATE",
            (attempt_id, student_id),
        ).fetchone()
        if not attempt:
            abort(403)
        if attempt["status"] != "in_progress":
            return redirect(url_for("result", attempt_id=attempt_id))
        questions = attempt_questions(attempt)
        merge_security_snapshot(conn, attempt_id, request.form)
        invalid, draft = validate_required_answers(request.form, questions)
        if invalid:
            conn.execute("UPDATE attempts SET answers_json=%s WHERE id=%s", (json.dumps(draft), attempt_id))
            conn.commit()
            session["invalid_question"] = invalid[0]
            flash_ui("Answer every question with a valid response before submitting.", "error")
            return redirect(url_for("exam"))
        score, total, percentage, grade10, category_scores, type_scores, captured = score_attempt(request.form, questions)
        conn.execute(
            """UPDATE attempts SET submitted_at=%s, score=%s, total=%s, percentage=%s, grade10=%s,
                category_scores=%s, type_scores=%s, answers_json=%s, status='submitted'
                WHERE id=%s AND status='in_progress'""",
            (now_iso(), score, total, percentage, grade10, json.dumps(category_scores), json.dumps(type_scores), json.dumps(captured), attempt_id),
        )
        conn.commit()
    return redirect(url_for("result", attempt_id=attempt_id))


def teacher_can_view_attempt(conn, attempt_id):
    teacher_id = current_teacher_id()
    if not teacher_id:
        return False
    active_teacher = conn.execute(
        "SELECT 1 FROM teachers WHERE id=%s AND is_active=1",
        (teacher_id,),
    ).fetchone()
    if not active_teacher:
        session.clear()
        return False
    return bool(conn.execute("SELECT 1 FROM attempts WHERE id=%s AND teacher_id=%s", (attempt_id, teacher_id)).fetchone())


@app.get("/result/<attempt_id>")
def result(attempt_id):
    teacher_view = bool(session.get("teacher_authenticated"))
    if teacher_view:
        with get_db() as _conn:
            if not teacher_can_view_attempt(_conn, attempt_id):
                abort(403)
    else:
        if not session.get("student_authenticated") or not session.get("student_id"):
            abort(403)
        with get_db() as _conn:
            _owned = _conn.execute("SELECT 1 FROM attempts WHERE id=%s AND student_id=%s", (attempt_id, session.get("student_id"))).fetchone()
        if not _owned:
            abort(403)
    with get_db() as conn:
        attempt = conn.execute("SELECT * FROM attempts WHERE id = %s", (attempt_id,)).fetchone()
        institution = setting(conn, "institution_name")
        penalty_total = active_penalty_total(conn, attempt_id) if attempt else 0
        penalties = penalties_for_attempt(conn, attempt_id) if attempt else []
        integrity_events = conn.execute(
            "SELECT * FROM integrity_events WHERE attempt_id=%s ORDER BY occurred_at,id", (attempt_id,)
        ).fetchall() if teacher_view else []
    if not attempt or attempt["status"] != "submitted":
        abort(404)
    labels = json.loads(attempt["category_labels_json"] or "{}")
    return render_template(
        "exam/result.html", attempt=attempt, category_scores=json.loads(attempt["category_scores"] or "{}"),
        type_scores=json.loads(attempt["type_scores"] or "{}"), category_labels=labels,
        type_labels=localized_type_labels(), institution=institution, blocked_actions=blocked_action_total(attempt),
        integrity_status=integrity_label(attempt), focus_changes=focus_change_total(attempt),
        penalties=penalties if teacher_view else [p for p in penalties if p["is_active"]], penalty_total=penalty_total,
        adjusted_grade=adjusted_grade10(attempt["grade10"], penalty_total),
        integrity_review=integrity_event_presentation(integrity_events, attempt, get_ui_language()) if teacher_view else None,
    )


@app.post("/teacher/results/<attempt_id>/penalties")
@teacher_required
def add_attempt_penalty(attempt_id):
    verify_csrf()
    reason = clean_name(request.form.get("reason"), 500)
    try:
        points = round(float(request.form.get("points", "")), 2)
    except (TypeError, ValueError):
        points = 0
    with get_db() as conn:
        attempt = conn.execute(
            "SELECT * FROM attempts WHERE id=%s AND teacher_id=%s AND status='submitted' FOR UPDATE",
            (attempt_id, current_teacher_id()),
        ).fetchone()
        if not attempt:
            abort(404)
        available = adjusted_grade10(attempt["grade10"], active_penalty_total(conn, attempt_id))
        if not reason:
            flash_ui("A reason is required for every penalty.", "error")
        elif not math.isfinite(points) or points <= 0 or points > available:
            flash_ui("Penalty points must be greater than zero and may not exceed the current adjusted grade.", "error")
        else:
            conn.execute(
                """INSERT INTO attempt_penalties(attempt_id,teacher_id,points,reason,created_at,is_active)
                   VALUES(%s,%s,%s,%s,%s,1)""",
                (attempt_id, current_teacher_id(), points, reason, now_iso()),
            )
            conn.commit()
            flash_ui("Penalty applied.", "success")
    return redirect(url_for("result", attempt_id=attempt_id))


@app.post("/teacher/results/<attempt_id>/penalties/<int:penalty_id>/revoke")
@teacher_required
def revoke_attempt_penalty(attempt_id, penalty_id):
    verify_csrf()
    with get_db() as conn:
        penalty = conn.execute(
            """SELECT p.id FROM attempt_penalties p JOIN attempts a ON a.id=p.attempt_id
               WHERE p.id=%s AND p.attempt_id=%s AND p.teacher_id=%s AND a.teacher_id=%s AND p.is_active=1""",
            (penalty_id, attempt_id, current_teacher_id(), current_teacher_id()),
        ).fetchone()
        if not penalty:
            abort(404)
        conn.execute(
            "UPDATE attempt_penalties SET is_active=0,revoked_at=%s,revoked_by=%s WHERE id=%s",
            (now_iso(), current_teacher_id(), penalty_id),
        )
        conn.commit()
    flash_ui("Penalty revoked.", "success")
    return redirect(url_for("result", attempt_id=attempt_id))


@app.get("/report/<attempt_id>.pdf")
def report_pdf(attempt_id):
    if session.get("teacher_authenticated"):
        with get_db() as _conn:
            if not teacher_can_view_attempt(_conn, attempt_id):
                abort(403)
    else:
        if not session.get("student_authenticated") or not session.get("student_id"):
            abort(403)
        with get_db() as _conn:
            _owned = _conn.execute("SELECT 1 FROM attempts WHERE id=%s AND student_id=%s", (attempt_id, session.get("student_id"))).fetchone()
        if not _owned:
            abort(403)
    with get_db() as conn:
        attempt = conn.execute("SELECT * FROM attempts WHERE id = %s", (attempt_id,)).fetchone()
        institution = setting(conn, "institution_name")
        penalty_total = active_penalty_total(conn, attempt_id) if attempt else 0
        penalties = [penalty for penalty in penalties_for_attempt(conn, attempt_id) if penalty["is_active"]] if attempt else []
    if not attempt or attempt["status"] != "submitted":
        abort(404)

    language = attempt["ui_language"] if "ui_language" in attempt.keys() and attempt["ui_language"] in SUPPORTED_UI_LANGUAGES else get_ui_language()
    rt = lambda text: tr(text, language=language)
    pdf_type_labels = localized_type_labels(language)
    labels = json.loads(attempt["category_labels_json"] or "{}")
    category_scores = json.loads(attempt["category_scores"] or "{}")
    type_scores = json.loads(attempt["type_scores"] or "{}")
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=0.65*inch, leftMargin=0.65*inch, topMargin=0.55*inch, bottomMargin=0.55*inch)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ExamInstitution", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=10.5, leading=12, textColor=colors.HexColor("#334155"), spaceAfter=2, keepWithNext=1))
    styles.add(ParagraphStyle("ExamTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=20, spaceAfter=4, keepWithNext=1))
    styles.add(ParagraphStyle("ExamMeta", parent=styles["BodyText"], fontSize=10, leading=12, spaceAfter=1, keepWithNext=1))
    styles.add(ParagraphStyle("ExamVersion", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=14, spaceBefore=6, spaceAfter=4, keepWithNext=1))
    styles.add(ParagraphStyle("ExamSection", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=11, leading=13, spaceBefore=6, spaceAfter=4, keepWithNext=1))
    styles.add(ParagraphStyle("ExamQuestion", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=6, spaceAfter=3, keepWithNext=1))
    styles.add(ParagraphStyle("ExamChoice", parent=styles["BodyText"], fontSize=9.5, leading=11, leftIndent=12, firstLineIndent=0, spaceAfter=0, spaceBefore=0, keepWithNext=1))
    story = [
        Paragraph(institution, styles["Heading2"]),
        Paragraph(attempt["assessment_title"] or DEFAULT_SETTINGS["assessment_title"], styles["Heading1"]),
        Paragraph(attempt["assessment_subject"] or "", styles["Heading3"]),
        Spacer(1, 8),
        Paragraph(f"<b>{rt('Student')}:</b> {attempt['student_name']}", styles["BodyText"]),
        Paragraph(f"<b>{rt('Email')}:</b> {attempt['student_email'] or '—'}", styles["BodyText"]),
        Paragraph(f"<b>{rt('Section')}:</b> {attempt['section']}", styles["BodyText"]),
        Paragraph(f"<b>{rt('Submitted')}:</b> {(attempt['submitted_at'] or '').replace('T', ' ')}", styles["BodyText"]),
        Spacer(1, 12),
        Paragraph(
            f"<b>{rt('Score')}:</b> {attempt['score']:.0f} / {attempt['total']:.0f} &nbsp;&nbsp; "
            f"<b>{rt('Percentage')}:</b> {attempt['percentage']:.1f}% &nbsp;&nbsp; "
            f"<b>{rt('Raw grade')}:</b> {attempt['grade10']:.1f} / 10 &nbsp;&nbsp; "
            f"<b>{rt('Penalty points')}:</b> {penalty_total:.2f} &nbsp;&nbsp; "
            f"<b>{rt('Adjusted final grade')}:</b> {adjusted_grade10(attempt['grade10'], penalty_total):.2f} / 10", styles["Heading3"]),
        Spacer(1, 18), Paragraph(rt("Performance by category"), styles["Heading3"]),
    ]
    category_data = [[rt("Category"), rt("Score")]]
    for key, values in category_scores.items():
        category_data.append([labels.get(str(key), f"{rt('Category')} {key}"), f"{values['earned']:.0f} / {values['total']:.0f}"])
    category_table = Table(category_data, colWidths=[4.8*inch, 1.2*inch])
    category_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#F1F5F9")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#D7DEE8")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("BOTTOMPADDING", (0,0), (-1,-1), 7), ("TOPPADDING", (0,0), (-1,-1), 7),
    ]))
    story += [category_table, Spacer(1, 18), Paragraph(rt("Performance by question type"), styles["Heading3"])]
    type_data = [[rt("Question type"), rt("Score")]]
    for t, values in type_scores.items():
        if values.get("total", 0):
            type_data.append([pdf_type_labels.get(t, t), f"{values['earned']:.0f} / {values['total']:.0f}"])
    type_table = Table(type_data, colWidths=[4.8*inch, 1.2*inch])
    type_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#F1F5F9")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#D7DEE8")),
        ("BOTTOMPADDING", (0,0), (-1,-1), 7), ("TOPPADDING", (0,0), (-1,-1), 7),
    ]))
    story += [
        type_table, Spacer(1, 18), Paragraph(rt("Exam integrity summary"), styles["Heading3"]),
        Paragraph(
            f"<b>{rt('Estimated window/app changes')}:</b> {focus_change_total(attempt)} &nbsp;&nbsp; "
            f"<b>{rt('Total detected time away')}:</b> {float(attempt['away_seconds'] or 0):.1f} s &nbsp;&nbsp; "
            f"<b>{rt('Blocked actions')}:</b> {blocked_action_total(attempt)}", styles["BodyText"]),
        Spacer(1, 10), Paragraph(rt("Integrity events are technical indicators for teacher review and are not, by themselves, proof of misconduct."), styles["BodyText"]),
        Spacer(1, 18), Paragraph(rt("This report records the student's submitted result. It does not display the answer key."), styles["BodyText"]),
    ]
    active_reasons = [penalty["reason"] for penalty in penalties]
    if active_reasons:
        story += [Spacer(1, 14), Paragraph(rt("Active penalty reasons"), styles["Heading3"])]
        story.extend(Paragraph(f"- {reason}", styles["BodyText"]) for reason in active_reasons)
    doc.build(story)
    buffer.seek(0)
    prefix = "Informe_Evaluacion" if language == "es" else "Assessment_Report"
    filename = f"{prefix}_{attempt['student_name'].replace(' ', '_')}.pdf"
    return send_file(buffer, as_attachment=True, download_name=filename, mimetype="application/pdf")


def result_filter_values(args=None):
    args = args or request.args
    return {
        "q": clean_name(args.get("q", ""), 120),
        "section": clean_name(args.get("section", ""), 80),
        "subject": clean_name(args.get("subject", ""), 120),
        "assessment": clean_name(args.get("assessment", ""), 180),
        "version": clean_name(args.get("version", ""), 80),
        "integrity": args.get("integrity", "").strip(),
        "date_from": args.get("date_from", "").strip()[:10],
        "date_to": args.get("date_to", "").strip()[:10],
    }


def submitted_where(filters):
    where = ["status='submitted'"]
    params = []
    teacher_id = current_teacher_id()
    if teacher_id:
        where.append("teacher_id=%s")
        params.append(teacher_id)
    if filters.get("q"):
        like = f"%{filters['q']}%"
        where.append("(student_name ILIKE %s OR COALESCE(student_email,'') ILIKE %s)")
        params.extend([like, like])
    if filters.get("section"):
        where.append("section=%s")
        params.append(filters["section"])
    if filters.get("subject"):
        where.append("COALESCE(assessment_subject,'')=%s")
        params.append(filters["subject"])
    if filters.get("assessment"):
        where.append("COALESCE(assessment_title,'')=%s")
        params.append(filters["assessment"])
    if filters.get("version"):
        where.append("COALESCE(exam_version_name,'')=%s")
        params.append(filters["version"])
    if filters.get("date_from"):
        where.append("substr(submitted_at,1,10)>=%s")
        params.append(filters["date_from"])
    if filters.get("date_to"):
        where.append("substr(submitted_at,1,10)<=%s")
        params.append(filters["date_to"])
    integrity = filters.get("integrity")
    incident_expr = "(COALESCE(focus_departures,0)+COALESCE(blur_events,0)+COALESCE(context_menu_attempts,0)+COALESCE(copy_attempts,0)+COALESCE(cut_attempts,0)+COALESCE(paste_attempts,0)+COALESCE(shortcut_attempts,0)+COALESCE(fullscreen_exits,0))"
    if integrity == "clean":
        where.append(f"{incident_expr}=0")
    elif integrity == "review":
        where.append(f"{incident_expr}>0")
    return where, params


def fetch_submitted_rows(filters=None, order="ASC"):
    filters = filters if filters is not None else result_filter_values()
    where, params = submitted_where(filters)
    direction = "DESC" if str(order).upper() == "DESC" else "ASC"
    with get_db() as conn:
        return conn.execute(
            f"""SELECT attempts.*,
                COALESCE((SELECT SUM(p.points) FROM attempt_penalties p WHERE p.attempt_id=attempts.id AND p.is_active=1),0) AS penalty_points,
                GREATEST(0,COALESCE(grade10,0)-COALESCE((SELECT SUM(p.points) FROM attempt_penalties p WHERE p.attempt_id=attempts.id AND p.is_active=1),0)) AS adjusted_grade10
                FROM attempts WHERE {' AND '.join(where)} ORDER BY submitted_at {direction}""", params
        ).fetchall()


def result_filter_options(conn):
    teacher_id = current_teacher_id()
    where = "status='submitted'"
    params = []
    if teacher_id:
        where += " AND teacher_id=%s"
        params.append(teacher_id)
    return {
        "sections": [r["section"] for r in conn.execute(f"SELECT DISTINCT section FROM attempts WHERE {where} AND TRIM(section)<>'' ORDER BY section ", params).fetchall()],
        "subjects": [r["assessment_subject"] for r in conn.execute(f"SELECT DISTINCT assessment_subject FROM attempts WHERE {where} AND COALESCE(TRIM(assessment_subject),'')<>'' ORDER BY assessment_subject ", params).fetchall()],
        "assessments": [r["assessment_title"] for r in conn.execute(f"SELECT DISTINCT assessment_title FROM attempts WHERE {where} AND COALESCE(TRIM(assessment_title),'')<>'' ORDER BY assessment_title ", params).fetchall()],
        "versions": [r["exam_version_name"] for r in conn.execute(f"SELECT DISTINCT exam_version_name FROM attempts WHERE {where} AND COALESCE(TRIM(exam_version_name),'')<>'' ORDER BY exam_version_name ", params).fetchall()],
    }


def teacher_dashboard_stats(conn, teacher_id, visible_rows):
    incident_expr = "(COALESCE(focus_departures,0)+COALESCE(blur_events,0)+COALESCE(context_menu_attempts,0)+COALESCE(copy_attempts,0)+COALESCE(cut_attempts,0)+COALESCE(paste_attempts,0)+COALESCE(shortcut_attempts,0)+COALESCE(fullscreen_exits,0))"
    summary = conn.execute(
        f"""SELECT COUNT(*) AS submitted,
                   COALESCE(AVG(percentage),0) AS average_percentage,
                   COALESCE(AVG(GREATEST(0,COALESCE(grade10,0)-COALESCE((SELECT SUM(p.points) FROM attempt_penalties p WHERE p.attempt_id=attempts.id AND p.is_active=1),0))),0) AS average_grade,
                   COALESCE(SUM(CASE WHEN {incident_expr}=0 THEN 1 ELSE 0 END),0) AS clean_count,
                   COALESCE(SUM(CASE WHEN {incident_expr}>0 THEN 1 ELSE 0 END),0) AS review_count,
                   COALESCE(SUM(COALESCE(focus_departures,0)+COALESCE(blur_events,0)),0) AS focus_total,
                   COALESCE(SUM(COALESCE(context_menu_attempts,0)+COALESCE(copy_attempts,0)+COALESCE(cut_attempts,0)+COALESCE(paste_attempts,0)+COALESCE(shortcut_attempts,0)),0) AS blocked_total,
                   COALESCE(SUM(CASE WHEN percentage < 60 THEN 1 ELSE 0 END),0) AS band_needs_support,
                   COALESCE(SUM(CASE WHEN percentage >= 60 AND percentage < 80 THEN 1 ELSE 0 END),0) AS band_developing,
                   COALESCE(SUM(CASE WHEN percentage >= 80 THEN 1 ELSE 0 END),0) AS band_strong
            FROM attempts WHERE teacher_id=%s AND status='submitted'""",
        (teacher_id,),
    ).fetchone()
    stats = dict(summary)
    stats["clean_rate"] = stats["clean_count"] / stats["submitted"] * 100 if stats["submitted"] else 0.0
    stats["active_exams"] = conn.execute(
        "SELECT COUNT(*) AS n FROM exams WHERE teacher_id=%s AND is_published=1 AND is_archived=0", (teacher_id,)
    ).fetchone()["n"]
    stats["question_bank"] = conn.execute(
        "SELECT COUNT(*) AS n FROM question_bank WHERE teacher_id=%s AND is_archived=0", (teacher_id,)
    ).fetchone()["n"]
    stats["roster"] = conn.execute(
        "SELECT COUNT(*) AS n FROM students WHERE is_active=1 AND COALESCE(is_archived,0)=0"
    ).fetchone()["n"]
    stats["sections"] = conn.execute(
        "SELECT COUNT(*) AS n FROM sections WHERE is_archived=0"
    ).fetchone()["n"]

    visible = {"submitted": len(visible_rows), "clean_count": 0, "review_count": 0, "focus_total": 0, "blocked_total": 0}
    percentages = []
    for row in visible_rows:
        focus = focus_change_total(row)
        blocked = blocked_action_total(row)
        visible["focus_total"] += focus
        visible["blocked_total"] += blocked
        visible["clean_count" if integrity_label(row) == "No incidents" else "review_count"] += 1
        if row["percentage"] is not None:
            percentages.append(float(row["percentage"]))
    visible["average_percentage"] = sum(percentages) / len(percentages) if percentages else 0.0
    visible["clean_rate"] = visible["clean_count"] / visible["submitted"] * 100 if visible["submitted"] else 0.0
    stats["visible"] = visible
    return stats


@app.route("/teacher", methods=["GET", "POST"])
def teacher():
    if request.method == "POST" and not session.get("teacher_authenticated"):
        email = normalize_email(request.form.get("email"))
        password = request.form.get("password", "")
        with get_db() as conn:
            account = conn.execute("SELECT * FROM teachers WHERE email=%s ", (email,)).fetchone()
            valid = bool(account and account["is_active"] and check_password_hash(account["password_hash"], password))
            if valid:
                conn.execute("UPDATE teachers SET last_login_at=%s,updated_at=%s WHERE id=%s", (now_iso(), now_iso(), account["id"]))
                conn.commit()
        if valid:
            session.clear()
            session["teacher_authenticated"] = True
            session["teacher_id"] = int(account["id"])
            session["teacher_role"] = account["role"]
            session["teacher_name"] = account["full_name"]
            csrf_token()
            return redirect(url_for("teacher"))
        flash_ui("Invalid teacher email or password.", "error")
    if not session.get("teacher_authenticated") or not current_teacher_id():
        return render_template("teacher/login.html", institution=current_settings()["institution_name"])
    filters = result_filter_values()
    rows = fetch_submitted_rows(filters, order="DESC")
    teacher_id = current_teacher_id()
    with get_db() as conn:
        teacher_user = conn.execute("SELECT * FROM teachers WHERE id=%s AND is_active=1", (teacher_id,)).fetchone()
        if not teacher_user:
            session.clear()
            return redirect(url_for("teacher"))
        dashboard_stats = teacher_dashboard_stats(conn, teacher_id, rows)
        options = result_filter_options(conn)
        title = "Assessment Studio"
    return render_template(
        "teacher/index.html", rows=rows, title=title, dashboard_stats=dashboard_stats,
        filters=filters, filter_options=options, teacher_user=teacher_user,
        subject_name=tr("Published exams"),
    )


@app.post("/teacher/logout")
def teacher_logout():
    if session.get("teacher_authenticated"):
        verify_csrf()
    session.clear()
    return redirect(url_for("teacher"))



def clean_name(value, max_len=120):
    return " ".join(str(value or "").strip().split())[:max_len]



def normalize_csv_key(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")


def read_csv_upload(upload):
    if not upload or not upload.filename:
        raise ValueError(tr("Choose CSV file"))
    if not upload.filename.lower().endswith(".csv"):
        raise ValueError("Only .csv files are accepted.")
    raw = upload.read()
    if not raw:
        raise ValueError("The CSV file is empty.")
    text = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("The CSV encoding could not be read.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        raise ValueError("The CSV does not contain headers.")
    fields = [normalize_csv_key(name) for name in reader.fieldnames]
    rows = []
    for raw_row in reader:
        row = {}
        for original, normalized in zip(reader.fieldnames, fields):
            row[normalized] = str(raw_row.get(original, "") or "").strip()
        if any(row.values()):
            rows.append(row)
    return fields, rows


def csv_cell(row, *names):
    for name in names:
        key = normalize_csv_key(name)
        if row.get(key, "").strip():
            return row[key].strip()
    return ""


def has_csv_column(fields, *names):
    return any(normalize_csv_key(name) in fields for name in names)


def csv_download(filename, headers, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    payload = io.BytesIO(("\ufeff" + stream.getvalue()).encode("utf-8"))
    payload.seek(0)
    return send_file(payload, as_attachment=True, download_name=filename, mimetype="text/csv; charset=utf-8")


def truthy_csv(value):
    return normalize_csv_key(value) in {"1", "true", "yes", "si", "s", "x", "verdadero"}


def split_pipe(value):
    return [part.strip() for part in str(value or "").split("|") if part.strip()]


def csv_teacher_scope(conn, teacher_id=None):
    teacher_id = teacher_id or (current_teacher_id() if has_request_context() else 0)
    if teacher_id:
        return teacher_id
    row = conn.execute("SELECT id FROM teachers WHERE is_active=1 ORDER BY id LIMIT 1").fetchone()
    return int(row["id"]) if row else 0


def table_has_column(conn, table, column):
    return column in table_columns(conn, table)


def subject_teacher_clause(conn, alias, teacher_id):
    if table_has_column(conn, "subjects", "teacher_id"):
        return f" AND {alias}.teacher_id=%s", [teacher_id]
    return "", []


def category_teacher_clause(conn, alias, teacher_id):
    if table_has_column(conn, "categories", "teacher_id"):
        return f" AND {alias}.teacher_id=%s", [teacher_id]
    return "", []




from routes import teacher_admin, teacher_catalog, teacher_reports, teacher_students  # noqa: E402,F401


def insert_question_row(conn, qid, values, is_active, prompt_override=None):
    timestamp = now_iso()
    columns = table_columns(conn, "question_bank")
    prompt = prompt_override if prompt_override is not None else values["prompt"]
    if "unit" in columns and "unit_label" in columns:
        category = conn.execute(
            "SELECT name FROM categories WHERE id=%s AND teacher_id=%s",
            (values["category_id"], current_teacher_id()),
        ).fetchone()
        unit_label = category["name"] if category else "Category"
        conn.execute(
            """INSERT INTO question_bank
            (id,unit,unit_label,teacher_id,subject_id,category_id,type,prompt,data_json,answer_json,audio,script,is_active,is_archived,created_at,updated_at)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,0,%s,%s)""",
            (qid, values["category_id"], unit_label, current_teacher_id(), values["subject_id"], values["category_id"], values["type"], prompt,
             values["data_json"], values["answer_json"], values["audio"], values["script"], is_active, timestamp, timestamp),
        )
    else:
        conn.execute(
            """INSERT INTO question_bank
            (id,teacher_id,subject_id,category_id,type,prompt,data_json,answer_json,audio,script,is_active,is_archived,created_at,updated_at)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,0,%s,%s)""",
            (qid, current_teacher_id(), values["subject_id"], values["category_id"], values["type"], prompt, values["data_json"],
             values["answer_json"], values["audio"], values["script"], is_active, timestamp, timestamp),
        )


def normalize_question_form(existing=None):
    qtype = request.form.get("type", "").strip()
    if qtype not in TYPE_LABELS:
        raise ValueError(tr("Select a valid question type."))
    try:
        subject_id = int(request.form.get("subject_id", "0"))
        category_id = int(request.form.get("category_id", "0"))
    except ValueError:
        raise ValueError(tr("Select a valid subject and category."))
    with get_db() as conn:
        category = conn.execute(
            """SELECT c.id, c.subject_id FROM categories c JOIN subjects s ON s.id=c.subject_id
               WHERE c.id=%s AND c.subject_id=%s AND c.teacher_id=%s AND c.is_archived=0 AND s.is_archived=0 AND s.teacher_id=%s""",
            (category_id, subject_id, current_teacher_id(), current_teacher_id()),
        ).fetchone()
    if not category:
        raise ValueError(tr("The selected category does not belong to that subject."))

    prompt = " ".join(request.form.get("prompt", "").strip().split())
    if len(prompt) < 3:
        raise ValueError(tr("Write the question prompt."))
    data = {}
    answer = None
    script = request.form.get("script", "").strip() or None
    audio = existing["audio"] if existing is not None and existing["audio"] else None

    if qtype in {"multiple_choice", "listening"}:
        raw_texts = [value.strip() for value in request.form.getlist("choice_text")]
        try:
            selected_raw_index = int(request.form.get("correct_choice", "-1"))
        except ValueError:
            selected_raw_index = -1
        kept = [(idx, text) for idx, text in enumerate(raw_texts) if text]
        if len(kept) < 2:
            raise ValueError(tr("This question type needs at least two answer options."))
        if selected_raw_index < 0 or selected_raw_index >= len(raw_texts) or not raw_texts[selected_raw_index]:
            raise ValueError(tr("Select a non-empty correct answer."))
        choices = [[f"c{idx+1}", text] for idx, (_, text) in enumerate(kept)]
        correct_position = next(i for i, (raw_idx, _) in enumerate(kept) if raw_idx == selected_raw_index)
        data["choices"] = choices
        answer = choices[correct_position][0]
        if qtype == "listening":
            listening_source = request.form.get("listening_source", "").strip().lower()
            if listening_source not in {"audio", "tts"}:
                listening_source = "audio" if audio else ("tts" if script else "audio")
            tts_lang = request.form.get("tts_lang", "auto").strip() or "auto"
            if tts_lang not in {"auto", "en-US", "es-MX", "es-ES"}:
                tts_lang = "auto"
            data["listening_source"] = listening_source
            data["tts_lang"] = tts_lang
            if listening_source == "audio":
                upload = request.files.get("audio_file")
                if upload and upload.filename:
                    filename = secure_filename(upload.filename)
                    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
                    if ext not in ALLOWED_AUDIO_EXTENSIONS:
                        raise ValueError(tr("Audio must be WAV, MP3, M4A, OGG, or AAC."))
                    unique_name = f"audio_{secrets.token_hex(6)}.{ext}"
                    upload.save(os.path.join(AUDIO_DIR, unique_name))
                    audio = unique_name
                if not audio:
                    raise ValueError(tr("Upload an audio file for the audio/listening question."))
            else:
                audio = None
                if not script or len(script.strip()) < 2:
                    raise ValueError(tr("Required when browser voice is selected."))
        else:
            audio = None
            script = None
    elif qtype == "true_false":
        answer = request.form.get("true_false_answer", "true")
        if answer not in {"true", "false"}:
            raise ValueError(tr("Choose True or False as the correct answer."))
        data["choices"] = [["true", "True"], ["false", "False"]]
        audio = None
        script = None
    elif qtype == "short_answer":
        answers = [line.strip() for line in request.form.get("accepted_answers", "").splitlines() if line.strip()]
        if not answers:
            raise ValueError(tr("Add at least one accepted answer, one per line."))
        answer = answers
        data["case_sensitive"] = bool(request.form.get("case_sensitive"))
        audio = None
        script = None
    elif qtype == "numeric":
        try:
            value = float(request.form.get("numeric_answer", "").replace(",", "."))
            tolerance = max(0.0, float((request.form.get("numeric_tolerance", "0") or "0").replace(",", ".")))
        except ValueError:
            raise ValueError(tr("Numeric answer and tolerance must be valid numbers."))
        answer = {"value": value, "tolerance": tolerance}
        audio = None
        script = None
    elif qtype == "order":
        items = [line.strip() for line in request.form.get("order_items", "").splitlines() if line.strip()]
        if len(items) < 2:
            raise ValueError(tr("Ordering questions need at least two items, one per line."))
        data["items"] = [[str(i + 1), text] for i, text in enumerate(items)]
        answer = [str(i + 1) for i in range(len(items))]
        audio = None
        script = None
    elif qtype == "matching":
        lefts = request.form.getlist("match_left")
        rights = request.form.getlist("match_right")
        pairs = [(l.strip(), r.strip()) for l, r in zip(lefts, rights) if l.strip() and r.strip()]
        if len(pairs) < 2:
            raise ValueError(tr("Matching questions need at least two complete pairs."))
        left = [[f"l{i+1}", pair[0]] for i, pair in enumerate(pairs)]
        right = [[f"r{i+1}", pair[1]] for i, pair in enumerate(pairs)]
        data["left"] = left
        data["right"] = right
        answer = {f"l{i+1}": f"r{i+1}" for i in range(len(pairs))}
        audio = None
        script = None

    return {
        "subject_id": subject_id, "category_id": category_id, "type": qtype, "prompt": prompt,
        "data_json": json.dumps(data), "answer_json": json.dumps(answer), "audio": audio, "script": script,
    }


def question_form_view(question=None):
    with get_db() as conn:
        subject_filter, subject_params = subject_teacher_clause(conn, "s", current_teacher_id())
        category_filter, category_params = category_teacher_clause(conn, "c", current_teacher_id())
        subjects = conn.execute(
            f"SELECT * FROM subjects s WHERE s.is_archived=0{subject_filter} ORDER BY s.name",
            tuple(subject_params),
        ).fetchall()
        categories = conn.execute(
            f"""SELECT c.*, s.name AS subject_name FROM categories c JOIN subjects s ON s.id=c.subject_id
               WHERE c.is_archived=0 AND s.is_archived=0{category_filter}{subject_filter} ORDER BY s.name, c.sort_order, c.name""",
            tuple(category_params + subject_params),
        ).fetchall()
    form_question = None
    if question is not None:
        data = json.loads(question["data_json"] or "{}")
        answer = json.loads(question["answer_json"])
        right_lookup = dict(data.get("right", []))
        pairs = []
        if question["type"] == "matching" and isinstance(answer, dict):
            for left_key, left_text in data.get("left", []):
                pairs.append([left_text, right_lookup.get(answer.get(left_key), "")])
        form_question = {
            "id": question["id"], "subject_id": question["subject_id"], "category_id": question["category_id"],
            "type": question["type"], "prompt": question["prompt"], "audio": question["audio"],
            "script": question["script"] or "", "is_active": bool(question["is_active"]),
            "listening_source": data.get("listening_source") or ("audio" if question["audio"] else ("tts" if question["script"] else "audio")),
            "tts_lang": data.get("tts_lang") or "auto",
            "choices": data.get("choices", []), "items": data.get("items", []), "pairs": pairs,
            "answer": answer, "case_sensitive": bool(data.get("case_sensitive", False)),
            "order_items_text": "\n".join(item[1] for item in data.get("items", [])) if question["type"] == "order" else "",
        }
    return render_template(
        "teacher/catalog/question_form.html", question=form_question, type_labels=localized_type_labels(),
        subjects=subjects, categories=categories,
    )


@app.get("/teacher/questions")
@teacher_required
def teacher_questions():
    search = request.args.get("q", "").strip()
    type_filter = request.args.get("type", "").strip()
    subject_filter = request.args.get("subject", "").strip()
    category_filter = request.args.get("category", "").strip()
    status_filter = "all"
    clauses = ["q.is_archived=0", "s.is_archived=0", "c.is_archived=0", "q.teacher_id=%s"]
    params = [current_teacher_id()]
    if search:
        clauses.append("(q.prompt ILIKE %s OR s.name ILIKE %s OR c.name ILIKE %s)")
        params += [f"%{search}%", f"%{search}%", f"%{search}%"]
    if type_filter in TYPE_LABELS:
        clauses.append("q.type=%s")
        params.append(type_filter)
    if subject_filter.isdigit():
        clauses.append("q.subject_id=%s")
        params.append(int(subject_filter))
    if category_filter.isdigit():
        clauses.append("q.category_id=%s")
        params.append(int(category_filter))

    sql = f"""
        SELECT q.*, s.name AS subject_name, c.name AS category_name
        FROM question_bank q JOIN subjects s ON s.id=q.subject_id JOIN categories c ON c.id=q.category_id
        WHERE {' AND '.join(clauses)} ORDER BY s.name, c.sort_order, c.name, q.type, q.updated_at DESC
    """
    with get_db() as conn:
        subject_filter, subject_params = subject_teacher_clause(conn, "s", current_teacher_id())
        category_filter, category_params = category_teacher_clause(conn, "c", current_teacher_id())
        rows = conn.execute(sql, params).fetchall()
        subjects = conn.execute(
            f"SELECT * FROM subjects s WHERE s.is_archived=0{subject_filter} ORDER BY s.name",
            tuple(subject_params),
        ).fetchall()
        categories = conn.execute(
            f"""SELECT c.*, s.name AS subject_name FROM categories c JOIN subjects s ON s.id=c.subject_id
               WHERE c.is_archived=0 AND s.is_archived=0{category_filter}{subject_filter} ORDER BY s.name, c.sort_order, c.name""",
            tuple(category_params + subject_params),
        ).fetchall()
        active_counts = {row["type"]: row["n"] for row in conn.execute(
            "SELECT type, COUNT(*) AS n FROM question_bank WHERE teacher_id=%s AND is_archived=0 GROUP BY type", (current_teacher_id(),)
        ).fetchall()}
    return render_template(
        "teacher/catalog/question_bank.html", rows=rows, subjects=subjects, categories=categories, type_labels=localized_type_labels(),
        filters={"q": search, "type": type_filter, "subject": subject_filter, "category": category_filter, "status": status_filter},
        active_counts=active_counts,
    )



QUESTION_TYPE_ALIASES = {
    "multiple_choice": "multiple_choice", "opcion_multiple": "multiple_choice", "opcion": "multiple_choice", "mcq": "multiple_choice",
    "true_false": "true_false", "verdadero_falso": "true_false", "vf": "true_false",
    "short_answer": "short_answer", "respuesta_corta": "short_answer",
    "numeric": "numeric", "numeric_answer": "numeric", "numerica": "numeric", "respuesta_numerica": "numeric",
    "order": "order", "ordering": "order", "ordenar": "order", "put_in_order": "order",
    "matching": "matching", "relacionar": "matching", "match": "matching",
    "listening": "listening", "audio": "listening", "audio_listening": "listening",
}


def resolve_csv_subject_category(conn, row, default_subject_id=None, default_category_id=None, teacher_id=None):
    teacher_id = csv_teacher_scope(conn, teacher_id)
    subject_filter, subject_params = subject_teacher_clause(conn, "s", teacher_id)
    category_filter, category_params = category_teacher_clause(conn, "c", teacher_id)
    subject_value = csv_cell(row, "Asignatura", "Subject")
    category_value = csv_cell(row, "Categoria", "Category")
    subject = None
    if subject_value:
        if subject_value.isdigit():
            subject = conn.execute(
                f"SELECT * FROM subjects s WHERE s.id=%s AND s.is_archived=0{subject_filter}",
                (int(subject_value), *subject_params),
            ).fetchone()
        if not subject:
            subject = conn.execute(
                f"SELECT * FROM subjects s WHERE s.name=%s AND s.is_archived=0{subject_filter}",
                (subject_value, *subject_params),
            ).fetchone()
    elif default_subject_id:
        subject = conn.execute(
            f"SELECT * FROM subjects s WHERE s.id=%s AND s.is_archived=0{subject_filter}",
            (default_subject_id, *subject_params),
        ).fetchone()
    elif default_category_id:
        subject = conn.execute(
            f"""SELECT s.* FROM subjects s JOIN categories c ON c.subject_id=s.id
                                  WHERE c.id=%s AND c.is_archived=0 AND s.is_archived=0{category_filter}{subject_filter}""",
            (default_category_id, *category_params, *subject_params),
        ).fetchone()
    if not subject:
        raise ValueError("Asignatura no encontrada" if get_ui_language() == "es" else "Subject not found")
    category = None
    if category_value:
        if category_value.isdigit():
            category = conn.execute(
                f"SELECT * FROM categories c WHERE c.id=%s AND c.subject_id=%s AND c.is_archived=0{category_filter}",
                (int(category_value), subject["id"], *category_params),
            ).fetchone()
        if not category:
            category = conn.execute(
                f"SELECT * FROM categories c WHERE c.subject_id=%s AND c.name=%s AND c.is_archived=0{category_filter}",
                (subject["id"], category_value, *category_params),
            ).fetchone()
    elif default_category_id:
        category = conn.execute(
            f"SELECT * FROM categories c WHERE c.id=%s AND c.subject_id=%s AND c.is_archived=0{category_filter}",
            (default_category_id, subject["id"], *category_params),
        ).fetchone()
    if not category:
        raise ValueError("Categoría no encontrada" if get_ui_language() == "es" else "Category not found")
    return subject, category


def question_values_from_csv(conn, row, default_subject_id=None, default_category_id=None, teacher_id=None):
    subject, category = resolve_csv_subject_category(conn, row, default_subject_id, default_category_id, teacher_id)
    raw_type = normalize_csv_key(csv_cell(row, "Tipo", "Type"))
    qtype = QUESTION_TYPE_ALIASES.get(raw_type)
    if not qtype:
        raise ValueError((f"Tipo de pregunta no válido: {csv_cell(row, 'Tipo', 'Type')}" if get_ui_language() == "es" else f"Invalid question type: {csv_cell(row, 'Tipo', 'Type')}"))
    prompt = clean_name(csv_cell(row, "Pregunta", "Question", "Prompt"), 1200)
    if len(prompt) < 3:
        raise ValueError("Pregunta vacía o demasiado corta" if get_ui_language() == "es" else "Question prompt is empty or too short")
    options = split_pipe(csv_cell(row, "Opciones", "Options", "Choices"))
    answer_raw = csv_cell(row, "Respuesta", "Answer", "Correcta", "Correct")
    pairs_raw = csv_cell(row, "Pares", "Pairs", "Matching")
    order_raw = csv_cell(row, "Orden", "Order", "Sequence")
    tolerance_raw = csv_cell(row, "Tolerancia", "Tolerance") or "0"
    script = csv_cell(row, "Script", "Guion") or None
    audio_name = secure_filename(csv_cell(row, "Audio", "Audio_file", "Archivo_audio")) or None
    audio = audio_name if audio_name and os.path.isfile(os.path.join(AUDIO_DIR, audio_name)) else None
    raw_listening_source = normalize_csv_key(csv_cell(row, "FuenteAudio", "AudioSource", "ListeningSource", "Listening_source"))
    raw_tts_lang = csv_cell(row, "IdiomaVoz", "VoiceLanguage", "SpeechLanguage", "TTSLang") or "auto"
    data = {}
    answer = None
    if qtype in {"multiple_choice", "listening"}:
        if len(options) < 2:
            raise ValueError("Se requieren al menos dos opciones" if get_ui_language() == "es" else "At least two options are required")
        choices = [[f"c{i+1}", text] for i, text in enumerate(options)]
        correct_idx = None
        for i, text in enumerate(options):
            if text.casefold() == answer_raw.casefold():
                correct_idx = i
                break
        if correct_idx is None and answer_raw.isdigit() and 1 <= int(answer_raw) <= len(options):
            correct_idx = int(answer_raw) - 1
        if correct_idx is None:
            raise ValueError("La respuesta correcta no coincide con las opciones" if get_ui_language() == "es" else "The correct answer does not match any option")
        data["choices"] = choices
        answer = choices[correct_idx][0]
        if qtype == "multiple_choice":
            script = None
            audio = None
        else:
            if raw_listening_source in {"tts", "voz", "script", "browser_voice"}:
                source = "tts"
            elif raw_listening_source in {"audio", "archivo", "file"}:
                source = "audio"
            else:
                source = "audio" if audio else ("tts" if script else "audio")
            if source == "tts" and script:
                audio = None
            data["listening_source"] = source
            data["tts_lang"] = raw_tts_lang if raw_tts_lang in {"auto", "en-US", "es-MX", "es-ES"} else "auto"
    elif qtype == "true_false":
        normalized = normalize_csv_key(answer_raw)
        if normalized in {"true", "verdadero", "1", "yes", "si"}:
            answer = "true"
        elif normalized in {"false", "falso", "0", "no"}:
            answer = "false"
        else:
            raise ValueError("Respuesta verdadero/falso no válida" if get_ui_language() == "es" else "Invalid true/false answer")
        data["choices"] = [["true", "True"], ["false", "False"]]
        audio = None; script = None
    elif qtype == "short_answer":
        answers = split_pipe(answer_raw)
        if not answers:
            raise ValueError("Agrega al menos una respuesta aceptada" if get_ui_language() == "es" else "Add at least one accepted answer")
        answer = answers
        data["case_sensitive"] = truthy_csv(csv_cell(row, "SensibleMayusculas", "CaseSensitive", "Case_sensitive"))
        audio = None; script = None
    elif qtype == "numeric":
        try:
            value = float(answer_raw.replace(",", "."))
            tolerance = max(0.0, float(tolerance_raw.replace(",", ".")))
        except ValueError:
            raise ValueError("Respuesta numérica o tolerancia no válida" if get_ui_language() == "es" else "Invalid numeric answer or tolerance")
        answer = {"value": value, "tolerance": tolerance}
        audio = None; script = None
    elif qtype == "order":
        items = split_pipe(order_raw or csv_cell(row, "Opciones", "Options"))
        if len(items) < 2:
            raise ValueError("Se requieren al menos dos elementos para ordenar" if get_ui_language() == "es" else "At least two ordering items are required")
        data["items"] = [[str(i+1), text] for i, text in enumerate(items)]
        answer = [str(i+1) for i in range(len(items))]
        audio = None; script = None
    elif qtype == "matching":
        pair_parts = split_pipe(pairs_raw)
        pairs = []
        for part in pair_parts:
            if "=>" in part:
                left, right = part.split("=>", 1)
            elif "=" in part:
                left, right = part.split("=", 1)
            else:
                continue
            if left.strip() and right.strip():
                pairs.append((left.strip(), right.strip()))
        if len(pairs) < 2:
            raise ValueError("Se requieren al menos dos pares usando izquierda=>derecha" if get_ui_language() == "es" else "At least two pairs are required using left=>right")
        data["left"] = [[f"l{i+1}", pair[0]] for i, pair in enumerate(pairs)]
        data["right"] = [[f"r{i+1}", pair[1]] for i, pair in enumerate(pairs)]
        answer = {f"l{i+1}": f"r{i+1}" for i in range(len(pairs))}
        audio = None; script = None
    missing_listening_source = False
    if qtype == "listening":
        source = data.get("listening_source") or ("audio" if audio else ("tts" if script else "audio"))
        missing_listening_source = (source == "audio" and not audio) or (source == "tts" and not script)
    return {
        "subject_id": subject["id"], "category_id": category["id"], "type": qtype, "prompt": prompt,
        "data_json": json.dumps(data, ensure_ascii=False), "answer_json": json.dumps(answer, ensure_ascii=False),
        "audio": audio, "script": script,
    }, missing_listening_source


@app.get("/teacher/questions/import/template")
@teacher_required
def teacher_questions_import_template():
    english = get_ui_language() == "en"
    headers = (["Subject", "Category", "Type", "Question", "Options", "Answer", "Pairs", "Order", "Tolerance", "CaseSensitive", "Script", "Audio", "AudioSource", "VoiceLanguage"]
               if english else ["Asignatura", "Categoria", "Tipo", "Pregunta", "Opciones", "Respuesta", "Pares", "Orden", "Tolerancia", "SensibleMayusculas", "Script", "Audio", "FuenteAudio", "IdiomaVoz"])
    rows = ([
        ["", "", "multiple_choice", "Which option is correct?", "A|B|C|D", "B", "", "", "", "", "", "", "", ""],
        ["", "", "true_false", "This statement is true.", "", "true", "", "", "", "", "", "", "", ""],
        ["", "", "short_answer", "Enter an accepted answer.", "", "answer|another answer", "", "", "", "no", "", "", "", ""],
        ["", "", "numeric", "What is the result?", "", "10", "", "", "0.5", "", "", "", "", ""],
        ["", "", "order", "Arrange the steps.", "", "", "", "Step 1|Step 2|Step 3", "", "", "", "", "", ""],
        ["", "", "matching", "Match the items.", "", "", "A=>1|B=>2|C=>3", "", "", "", "", "", "", ""],
        ["", "", "listening", "Listen and select.", "Yes|No", "Yes", "", "", "", "", "This sentence is read by the browser.", "", "tts", "en-US"],
    ] if english else [
        ["", "", "multiple_choice", "¿Cuál opción es correcta?", "A|B|C|D", "B", "", "", "", "", "", "", "", ""],
        ["", "", "true_false", "La afirmación es verdadera.", "", "verdadero", "", "", "", "", "", "", "", ""],
        ["", "", "short_answer", "Escribe una respuesta válida.", "", "respuesta|otra respuesta", "", "", "", "no", "", "", "", ""],
        ["", "", "numeric", "¿Cuál es el resultado?", "", "10", "", "", "0.5", "", "", "", "", ""],
        ["", "", "order", "Ordena los pasos.", "", "", "", "Paso 1|Paso 2|Paso 3", "", "", "", "", "", ""],
        ["", "", "matching", "Relaciona correctamente.", "", "", "A=>1|B=>2|C=>3", "", "", "", "", "", "", ""],
        ["", "", "listening", "Escucha y selecciona.", "Yes|No", "Yes", "", "", "", "", "This sentence is read by the browser.", "", "tts", "en-US"],
    ])
    return csv_download("question_bank_import_template.csv" if english else "plantilla_banco_preguntas.csv", headers, rows)


@app.post("/teacher/questions/import")
@teacher_required
def teacher_questions_import():
    verify_csrf()
    try:
        default_subject_id = int(request.form.get("default_subject_id", "0") or 0) or None
    except ValueError:
        default_subject_id = None
    try:
        default_category_id = int(request.form.get("default_category_id", "0") or 0) or None
    except ValueError:
        default_category_id = None
    try:
        fields, rows = read_csv_upload(request.files.get("csv_file"))
    except ValueError as exc:
        flash_ui(str(exc), "error")
        return redirect(url_for("teacher_questions"))
    if not (has_csv_column(fields, "Tipo", "Type") and has_csv_column(fields, "Pregunta", "Question", "Prompt")):
        flash_ui("The CSV must include type and question columns.", "error")
        return redirect(url_for("teacher_questions"))
    imported = 0
    needs_audio = 0
    errors = []
    with get_db() as conn:
        for idx, row in enumerate(rows, start=2):
            try:
                with conn.transaction():
                    values, missing_audio = question_values_from_csv(conn, row, default_subject_id, default_category_id)
                    qid = f"q_{secrets.token_hex(6)}"
                    insert_question_row(conn, qid, values, 0)
                imported += 1
                if missing_audio:
                    needs_audio += 1
            except (ValueError, IntegrityError) as exc:
                errors.append(f"Fila {idx}: {exc}" if get_ui_language() == "es" else f"Row {idx}: {exc}")
        conn.commit()
    if not imported:
        flash_ui("No valid questions were imported.", "error")
    else:
        flash_ui("Question import complete", "success")
    if errors:
        flash(f"{tr('Rows skipped')}: {len(errors)} · " + " | ".join(errors[:4]), "warning")
    if needs_audio:
        flash(f"{tr('Listening questions needing a source')}: {needs_audio}", "warning")
    return redirect(url_for("teacher_questions"))

@app.route("/teacher/questions/new", methods=["GET", "POST"])
@teacher_required
def teacher_question_new():
    if request.method == "POST":
        verify_csrf()
        try:
            values = normalize_question_form()
        except ValueError as exc:
            flash(str(exc), "error")
            return question_form_view()
        qid = f"q_{secrets.token_hex(6)}"
        with get_db() as conn:
            insert_question_row(conn, qid, values, 0)
            conn.commit()
        flash_ui("Question added to the bank.", "success")
        return redirect(url_for("teacher_questions"))
    return question_form_view()


@app.route("/teacher/questions/<qid>/edit", methods=["GET", "POST"])
@teacher_required
def teacher_question_edit(qid):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM question_bank WHERE id=%s AND teacher_id=%s AND is_archived=0", (qid, current_teacher_id())).fetchone()
    if not row:
        abort(404)
    if request.method == "POST":
        verify_csrf()
        try:
            values = normalize_question_form(row)
        except ValueError as exc:
            flash(str(exc), "error")
            return question_form_view(row)
        with get_db() as conn:
            conn.execute(
                """UPDATE question_bank SET subject_id=%s,category_id=%s,type=%s,prompt=%s,data_json=%s,answer_json=%s,
                   audio=%s,script=%s,updated_at=%s WHERE id=%s AND teacher_id=%s""",
                (values["subject_id"], values["category_id"], values["type"], values["prompt"], values["data_json"],
                 values["answer_json"], values["audio"], values["script"], now_iso(), qid, current_teacher_id()),
            )
            conn.commit()
        flash_ui("Question updated. Existing student attempts keep their original snapshot.", "success")
        return redirect(url_for("teacher_questions"))
    return question_form_view(row)


@app.post("/teacher/questions/<qid>/toggle")
@teacher_required
def teacher_question_toggle(qid):
    verify_csrf()
    with get_db() as conn:
        row = conn.execute("SELECT is_active FROM question_bank WHERE id=%s AND teacher_id=%s AND is_archived=0", (qid, current_teacher_id())).fetchone()
        if not row:
            abort(404)
        conn.execute("UPDATE question_bank SET is_active=%s,updated_at=%s WHERE id=%s AND teacher_id=%s", (0 if row["is_active"] else 1, now_iso(), qid, current_teacher_id()))
        conn.commit()
    flash_ui("Question availability updated.", "success")
    return redirect(request.referrer or url_for("teacher_questions"))


def clone_question_row(conn, source_row, new_id):
    """Clone a question while preserving ownership and server-side answer data."""
    timestamp = now_iso()
    copy_suffix = " (copy)"
    if has_request_context():
        copy_suffix = " (copia)" if get_ui_language() == "es" else " (copy)"
    conn.execute(
        """INSERT INTO question_bank
           (id,teacher_id,subject_id,category_id,type,prompt,data_json,answer_json,audio,script,
            is_active,is_archived,created_at,updated_at)
           VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,0,0,%s,%s)""",
        (new_id, source_row["teacher_id"], source_row["subject_id"], source_row["category_id"],
         source_row["type"], f"{source_row['prompt']}{copy_suffix}", source_row["data_json"],
         source_row["answer_json"], source_row["audio"], source_row["script"], timestamp, timestamp),
    )


@app.post("/teacher/questions/<qid>/duplicate")
@teacher_required
def teacher_question_duplicate(qid):
    verify_csrf()
    new_id = f"q_{secrets.token_hex(6)}"
    try:
        with get_db() as conn:
            row = conn.execute("SELECT * FROM question_bank WHERE id=%s AND teacher_id=%s AND is_archived=0", (qid, current_teacher_id())).fetchone()
            if not row:
                abort(404)
            clone_question_row(conn, row, new_id)
            conn.commit()
    except DatabaseError as exc:
        app.logger.exception("Could not duplicate question %s", qid)
        message = (
            "No se pudo duplicar la pregunta. Revisa la estructura de la base de datos e inténtalo nuevamente."
            if get_ui_language() == "es" else
            "The question could not be duplicated. Check the database structure and try again."
        )
        flash(message, "error")
        return redirect(request.referrer or url_for("teacher_questions"))
    flash_ui("Question duplicated as bank-only so you can edit it safely.", "success")
    return redirect(url_for("teacher_question_edit", qid=new_id))


@app.post("/teacher/questions/<qid>/archive")
@teacher_required
def teacher_question_archive(qid):
    verify_csrf()
    with get_db() as conn:
        row = conn.execute("SELECT id FROM question_bank WHERE id=%s AND teacher_id=%s AND is_archived=0", (qid, current_teacher_id())).fetchone()
        if not row:
            abort(404)
        conn.execute(
            "UPDATE question_bank SET is_active=0,is_archived=1,updated_at=%s WHERE id=%s AND teacher_id=%s",
            (now_iso(), qid, current_teacher_id()),
        )
        conn.commit()
    flash_ui("Question archived.", "success")
    return redirect(request.referrer or url_for("teacher_questions"))


@app.post("/teacher/questions/bulk")
@teacher_required
def teacher_questions_bulk():
    verify_csrf()
    ids = [qid for qid in request.form.getlist("question_ids") if qid]
    action = request.form.get("action", "")
    if not ids:
        flash_ui("Select at least one question.", "error")
        return redirect(url_for("teacher_questions"))
    placeholders = ",".join("%s" for _ in ids)
    with get_db() as conn:
        if action == "activate":
            conn.execute(f"UPDATE question_bank SET is_active=1,updated_at=%s WHERE teacher_id=%s AND id IN ({placeholders})", [now_iso(), current_teacher_id(), *ids])
        elif action == "deactivate":
            conn.execute(f"UPDATE question_bank SET is_active=0,updated_at=%s WHERE teacher_id=%s AND id IN ({placeholders})", [now_iso(), current_teacher_id(), *ids])
        elif action == "archive":
            conn.execute(f"UPDATE question_bank SET is_active=0,is_archived=1,updated_at=%s WHERE teacher_id=%s AND id IN ({placeholders})", [now_iso(), current_teacher_id(), *ids])
        else:
            flash_ui("Choose a valid bulk action.", "error")
            return redirect(url_for("teacher_questions"))
        conn.commit()
    flash((f"Se actualizaron {len(ids)} pregunta(s)." if get_ui_language() == "es" else f"Updated {len(ids)} question(s)."), "success")
    return redirect(url_for("teacher_questions"))



def exam_row(conn, exam_id):
    return conn.execute(
        """SELECT e.*,s.name AS subject_name FROM exams e JOIN subjects s ON s.id=e.subject_id
           WHERE e.id=%s AND e.teacher_id=%s AND e.is_archived=0""", (exam_id, current_teacher_id())
    ).fetchone()


def printable_question_row(row, rng=None):
    data = json.loads(row["data_json"] or "{}")
    item = {"type": row["type"], "prompt": row["prompt"]}
    if row["type"] in {"multiple_choice", "true_false", "listening"}:
        choices = [list(choice) for choice in data.get("choices", [])]
        if rng is not None:
            rng.shuffle(choices)
        item["choices"] = choices
    elif row["type"] == "order":
        items = [list(value) for value in data.get("items", [])]
        if rng is not None:
            rng.shuffle(items)
        item["items"] = items
    elif row["type"] == "matching":
        left = [list(value) for value in data.get("left", [])]
        right = [list(value) for value in data.get("right", [])]
        if rng is not None:
            rng.shuffle(left)
            rng.shuffle(right)
        item["left"] = left
        item["right"] = right
    return item


def exam_version_questions_pdf_rows(conn, version_id):
    return conn.execute(
        """SELECT q.*, s.name AS subject_name, c.name AS category_name
           FROM exam_version_questions evq
           JOIN question_bank q ON q.id=evq.question_id
           JOIN subjects s ON s.id=q.subject_id
           JOIN categories c ON c.id=q.category_id
          WHERE evq.version_id=%s AND q.teacher_id=%s
          ORDER BY evq.position, q.created_at, q.id""",
        (version_id, current_teacher_id()),
    ).fetchall()


def exam_export_payload(conn, exam_id):
    exam = exam_row(conn, exam_id)
    if not exam:
        abort(404)
    versions = conn.execute(
        """SELECT id,name FROM exam_versions
           WHERE exam_id=%s AND is_active=1 ORDER BY id""",
        (exam_id,),
    ).fetchall()
    if not versions:
        abort(404)
    institution = setting(conn, "institution_name")
    teacher = current_teacher(conn)
    teacher_name = teacher["full_name"] if teacher else session.get("teacher_name", "Teacher")
    language = get_ui_language()
    rt = lambda text: tr(text, language=language)
    version_payloads = []
    for version in versions:
        rows = exam_version_questions_pdf_rows(conn, version["id"])
        questions = []
        for row in rows:
            rng = random.Random(f"{exam['id']}:{version['id']}:{row['id']}")
            questions.append(printable_question_row(row, rng))
        version_payloads.append({"version": version, "questions": questions})
    return exam, institution, teacher_name, rt, version_payloads


def exam_export_filename(exam, extension):
    base = secure_filename(exam["title"]) or "exam"
    return f"{base}.{extension}"


def markdown_escape(text):
    return str(text or "").replace("\r", "").strip()


def docx_paragraph(runs, keep_next=False, keep_lines=False, space_before=0, space_after=0):
    props = []
    if keep_next:
        props.append("<w:keepNext/>")
    if keep_lines:
        props.append("<w:keepLines/>")
    if space_before or space_after:
        props.append(f'<w:spacing w:before="{space_before}" w:after="{space_after}"/>')
    ppr = f"<w:pPr>{''.join(props)}</w:pPr>" if props else ""
    return f"<w:p>{ppr}{''.join(runs)}</w:p>"


def docx_table(cells, col_widths):
    grid = ''.join(f'<w:gridCol w:w="{width}"/>' for width in col_widths)
    rows = []
    for row_index, row in enumerate(cells):
        row_props = '<w:trPr><w:tblHeader/><w:cantSplit/></w:trPr>' if row_index == 0 else '<w:trPr><w:cantSplit/></w:trPr>'
        row_cells = []
        for col_index, cell in enumerate(row):
            width = col_widths[col_index]
            row_cells.append(
                f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/></w:tcPr>'
                f'{docx_paragraph([docx_run(cell["text"], bold=cell.get("bold", False))], space_before=0, space_after=0)}'
                '</w:tc>'
            )
        rows.append(f'<w:tr>{row_props}{"".join(row_cells)}</w:tr>')
    return (
        '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/>'
        '<w:tblBorders>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:insideH w:val="single" w:sz="2" w:space="0" w:color="E2E8F0"/>'
        '<w:insideV w:val="single" w:sz="2" w:space="0" w:color="E2E8F0"/>'
        '</w:tblBorders></w:tblPr>'
        f'<w:tblGrid>{grid}</w:tblGrid>'
        f'{"".join(rows)}</w:tbl>'
    )


def exam_question_markdown(question, question_number, rt):
    lines = [f"### {question_number}. {markdown_escape(question['prompt'])}"]
    if question["type"] in {"multiple_choice", "true_false", "listening"}:
        lines.append("")
        lines.append("#### Choices")
        for choice_index, (_, choice_text) in enumerate(question.get("choices", [])):
            lines.append(f"- {pdf_choice_letter(choice_index)}. {markdown_escape(choice_text)}")
    elif question["type"] == "order":
        lines.append("")
        lines.append(f"#### {rt('Order line')}")
        for _, item_text in question.get("items", []):
            lines.append(f"- {markdown_escape(item_text)}")
        lines.append("- ________________________________________________")
    elif question["type"] == "matching":
        lines.append("")
        lines.append(f"#### {rt('Left items')}")
        for _, item_text in question.get("left", []):
            lines.append(f"- {markdown_escape(item_text)}")
        lines.append("")
        lines.append(f"#### {rt('Right items')}")
        for _, item_text in question.get("right", []):
            lines.append(f"- {markdown_escape(item_text)}")
    else:
        lines.append("")
        lines.append(f"#### {rt('Answer line')}")
        lines.append("_______________________________________________")
    return "\n".join(lines)


def exam_markdown_export(exam, institution, teacher_name, rt, version_payloads):
    parts = [
        f"# {markdown_escape(exam['title'])}",
        "",
        f"- {rt('Institution')}: {markdown_escape(institution)}",
        f"- {rt('Subject')}: {markdown_escape(exam['subject_name'] or '')}",
        f"- {rt('Teacher')}: {markdown_escape(teacher_name)}",
    ]
    for index, payload in enumerate(version_payloads):
        version = payload["version"]
        questions = payload["questions"]
        parts.extend([
            "",
            f"## {rt('Version')} {markdown_escape(version['name'])}",
            "",
            f"### {rt('Questions')} ({len(questions)})",
        ])
        for q_index, question in enumerate(questions, start=1):
            parts.extend(["", exam_question_markdown(question, q_index, rt)])
        if index < len(version_payloads) - 1:
            parts.append("")
    parts.append("")
    return "\n".join(parts)


def docx_run(text, bold=False):
    props = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f"<w:r>{props}<w:t xml:space=\"preserve\">{xml_escape(str(text or ''))}</w:t></w:r>"
def docx_page_break():
    return "<w:p><w:r><w:br w:type=\"page\"/></w:r></w:p>"


def exam_docx_export(exam, institution, teacher_name, rt, version_payloads):
    paragraphs = [
        docx_paragraph([docx_run(institution, bold=True)], keep_next=True, space_after=30),
        docx_paragraph([docx_run(exam["title"], bold=True)], keep_next=True, space_after=90),
        docx_paragraph([docx_run(f"{rt('Institution')}: ", bold=True), docx_run(institution)], keep_next=True, space_after=0),
        docx_paragraph([docx_run(f"{rt('Subject')}: ", bold=True), docx_run(exam["subject_name"] or "")], keep_next=True, space_after=0),
        docx_paragraph([docx_run(f"{rt('Teacher')}: ", bold=True), docx_run(teacher_name)], keep_next=True, space_after=0),
    ]
    for index, payload in enumerate(version_payloads):
        version = payload["version"]
        questions = payload["questions"]
        paragraphs.extend([
            docx_paragraph([docx_run(f"{rt('Version')} {version['name']}", bold=True)], keep_next=True, space_before=120, space_after=30),
            docx_paragraph([docx_run(f"{rt('Questions')} ({len(questions)})", bold=True)], keep_next=True, space_after=60),
        ])
        for q_index, question in enumerate(questions, start=1):
            paragraphs.append(docx_paragraph([docx_run(f"{q_index}. {question['prompt']}", bold=True)], keep_next=True, space_before=90, space_after=30))
            if question["type"] in {"multiple_choice", "true_false", "listening"}:
                for choice_index, (_, choice_text) in enumerate(question.get("choices", [])):
                    paragraphs.append(docx_paragraph([docx_run(f"{pdf_choice_letter(choice_index)}. ", bold=True), docx_run(choice_text)], space_before=0, space_after=0))
            elif question["type"] == "order":
                for _, item_text in question.get("items", []):
                    paragraphs.append(docx_paragraph([docx_run(f"• {item_text}")], space_before=0, space_after=0))
                paragraphs.append(docx_paragraph([docx_run(f"{rt('Order line')}: ________________________________________________")], space_before=0, space_after=0))
            elif question["type"] == "matching":
                left_items = [text for _, text in question.get("left", [])]
                right_items = [text for _, text in question.get("right", [])]
                max_rows = max(len(left_items), len(right_items))
                table_rows = [[{"text": rt('Left items'), "bold": True}, {"text": rt('Right items'), "bold": True}]]
                for idx in range(max_rows):
                    table_rows.append([
                        {"text": left_items[idx] if idx < len(left_items) else ""},
                        {"text": right_items[idx] if idx < len(right_items) else ""},
                    ])
                paragraphs.append(docx_table(table_rows, [4320, 4320]))
            else:
                paragraphs.append(docx_paragraph([docx_run(f"{rt('Answer line')}: ________________________________________________")], space_before=0, space_after=0))
        if index < len(version_payloads) - 1:
            paragraphs.append(docx_page_break())
    body = "".join(paragraphs) + (
        '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="864" w:right="936" w:bottom="864" w:left="936" w:header="0" w:footer="0" w:gutter="0"/></w:sectPr>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:body>{body}</w:body></w:document>'
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '</Types>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        '</Relationships>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document)
    buffer.seek(0)
    return buffer


def exam_pdf_page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 9)
    canvas.drawRightString(letter[0] - 0.65 * inch, 0.4 * inch, f"Page {canvas.getPageNumber()}")
    canvas.restoreState()


def pdf_choice_letter(index):
    label = ""
    while True:
        index, remainder = divmod(index, 26)
        label = chr(65 + remainder) + label
        if index == 0:
            return label
        index -= 1


from routes import teacher_exams  # noqa: E402,F401


def export_categories(rows):
    labels = {}
    for row in rows:
        try:
            row_labels = json.loads(row["category_labels_json"] or "{}")
            scores = json.loads(row["category_scores"] or "{}")
        except (json.JSONDecodeError, TypeError):
            continue
        subject = (row["assessment_subject"] or tr("Subject")).strip()
        for key in scores:
            category = row_labels.get(str(key), f"{tr('Category')} {key}")
            labels.setdefault(str(key), f"{subject} · {category}")
    return sorted(labels.items(), key=lambda item: (item[1].casefold(), item[0]))


def score_text(data):
    return f"{data.get('earned', 0)}/{data.get('total', 0)}"


@app.get("/teacher/export.csv")
@teacher_required
def export_csv():
    rows = fetch_submitted_rows()
    categories = export_categories(rows)
    output = io.StringIO()
    writer = csv.writer(output)
    export_type_labels = localized_type_labels()
    headers = [
        tr("Attempt ID"), tr("Assessment"), tr("Version"), tr("Subject"), tr("Student"), tr("Email"), tr("Section"), tr("Started"), tr("Submitted"), tr("Score"), tr("Total"), tr("Percentage"), f"{tr('Raw grade')} / 10", tr("Penalty points"), f"{tr('Adjusted final grade')} / 10",
        *[label for _, label in categories], *export_type_labels.values(),
        tr("Integrity Status"), tr("Estimated window/app changes"), tr("Tab/App Visibility Changes"), tr("Total Time Away (s)"),
        tr("Longest Time Away (s)"), tr("Window Blur Events"), tr("Page Hide Events"), tr("Right-click Attempts"), tr("Copy Attempts"), tr("Cut Attempts"),
        tr("Paste Attempts"), tr("Blocked Shortcuts"), tr("Fullscreen Exits")
    ]
    writer.writerow(headers)
    for r in rows:
        category_scores = json.loads(r["category_scores"] or "{}")
        types = json.loads(r["type_scores"] or "{}")
        writer.writerow([
            r["id"], r["assessment_title"] or tr("Assessment"), r["exam_version_name"] or "", r["assessment_subject"] or "", r["student_name"], r["student_email"] or "", r["section"],
            r["started_at"], r["submitted_at"], r["score"], r["total"], r["percentage"], r["grade10"], r["penalty_points"], r["adjusted_grade10"],
            *[score_text(category_scores.get(key, {})) for key, _ in categories],
            *[score_text(types.get(key, {})) for key in TYPE_LABELS],
            tr(integrity_label(r)), focus_change_total(r), r["focus_departures"], round(float(r["away_seconds"] or 0), 1),
            round(float(r["longest_away_seconds"] or 0), 1), r["blur_events"], r["pagehide_events"], r["context_menu_attempts"],
            r["copy_attempts"], r["cut_attempts"], r["paste_attempts"], r["shortcut_attempts"], r["fullscreen_exits"],
        ])
    response = make_response(output.getvalue())
    filters = result_filter_values()
    section_slug = secure_filename(filters.get("section") or "all") or "all"
    response.headers["Content-Disposition"] = f"attachment; filename=assessment_results_{section_slug}.csv"
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    return response


@app.get("/teacher/export.xlsx")
@teacher_required
def export_xlsx():
    rows = fetch_submitted_rows()
    categories = export_categories(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = tr("Results")[:31]
    export_type_labels = localized_type_labels()
    headers = [
        tr("Attempt ID"), tr("Assessment"), tr("Version"), tr("Subject"), tr("Student"), tr("Email"), tr("Section"), tr("Started"), tr("Submitted"), tr("Score"), tr("Total"), tr("Percentage"), f"{tr('Raw grade')} / 10", tr("Penalty points"), f"{tr('Adjusted final grade')} / 10",
        *[label for _, label in categories], *export_type_labels.values(),
        tr("Integrity Status"), tr("Estimated window/app changes"), tr("Tab/App Visibility Changes"), tr("Total Time Away (s)"), tr("Longest Time Away (s)"),
        tr("Window Blur Events"), tr("Page Hide Events"), tr("Right-click Attempts"), tr("Copy Attempts"), tr("Cut Attempts"), tr("Paste Attempts"), tr("Blocked Shortcuts"), tr("Fullscreen Exits")
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="E8EEF8")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        category_scores = json.loads(r["category_scores"] or "{}")
        types = json.loads(r["type_scores"] or "{}")
        ws.append([
            r["id"], r["assessment_title"] or tr("Assessment"), r["exam_version_name"] or "", r["assessment_subject"] or "", r["student_name"], r["student_email"] or "", r["section"],
            r["started_at"], r["submitted_at"], r["score"], r["total"], r["percentage"], r["grade10"], r["penalty_points"], r["adjusted_grade10"],
            *[score_text(category_scores.get(key, {})) for key, _ in categories],
            *[score_text(types.get(key, {})) for key in TYPE_LABELS],
            tr(integrity_label(r)), focus_change_total(r), r["focus_departures"], round(float(r["away_seconds"] or 0), 1),
            round(float(r["longest_away_seconds"] or 0), 1), r["blur_events"], r["pagehide_events"], r["context_menu_attempts"],
            r["copy_attempts"], r["cut_attempts"], r["paste_attempts"], r["shortcut_attempts"], r["fullscreen_exits"],
        ])
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 18
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    events_ws = wb.create_sheet(tr("Integrity Events")[:31])
    events_ws.append([tr("Attempt ID"), tr("Assessment"), tr("Version"), tr("Subject"), tr("Student"), tr("Email"), tr("Section"), tr("Event Type"), tr("Server Time"), tr("Details")])
    for cell in events_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    attempt_ids = [r["id"] for r in rows]
    with get_db() as conn:
        if attempt_ids:
            placeholders = ",".join("%s" for _ in attempt_ids)
            events = conn.execute(
                f"""SELECT e.attempt_id,a.assessment_title,a.exam_version_name,a.assessment_subject,a.student_name,a.student_email,a.section,e.event_type,e.occurred_at,e.detail_json
                    FROM integrity_events e JOIN attempts a ON a.id=e.attempt_id
                    WHERE a.status='submitted' AND a.id IN ({placeholders})
                    ORDER BY e.occurred_at ASC,e.id ASC""",
                attempt_ids,
            ).fetchall()
        else:
            events = []
    for event in events:
        events_ws.append([
            event["attempt_id"], event["assessment_title"], event["exam_version_name"] or "", event["assessment_subject"], event["student_name"], event["student_email"] or "", event["section"],
            event["event_type"], event["occurred_at"], event["detail_json"] or "{}"
        ])
    for col, width in {"A":18,"B":34,"C":22,"D":30,"E":30,"F":14,"G":22,"H":21,"I":60}.items():
        events_ws.column_dimensions[col].width = width
    events_ws.freeze_panes = "A2"
    events_ws.auto_filter.ref = events_ws.dimensions

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    filters = result_filter_values()
    section_slug = secure_filename(filters.get("section") or "all") or "all"
    return send_file(stream, as_attachment=True, download_name=f"assessment_results_{section_slug}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=os.getenv("FLASK_DEBUG") == "1")
