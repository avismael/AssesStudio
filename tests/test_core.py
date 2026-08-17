import ast
import csv
import io
import json
import os
import re
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

_tmp = tempfile.NamedTemporaryFile(suffix='.db', delete=False)
_tmp.close()
os.environ['DATABASE_PATH'] = _tmp.name

import app  # noqa: E402
import policy_defaults  # noqa: E402
import seed  # noqa: E402


def test_default_catalog_exists():
    with app.get_db() as conn:
        subject = conn.execute("SELECT * FROM subjects WHERE id=1").fetchone()
        category = conn.execute("SELECT * FROM categories WHERE id=1").fetchone()
    assert subject['name'] == 'General'
    assert category['name'] == 'General'


def test_short_answer_scoring_is_case_insensitive():
    q = [{
        'id':'q1','category_id':1,'type':'short_answer','answer':['Photosynthesis'],
        'case_sensitive':False
    }]
    class F:
        def get(self, key, default=''):
            return '  photosynthesis  ' if key == 'q_q1' else default
    score, total, *_ = app.score_attempt(F(), q)
    assert score == 1.0 and total == 1.0


def test_numeric_tolerance():
    q = [{'id':'q2','category_id':1,'type':'numeric','answer':{'value':10.0,'tolerance':0.5}}]
    class F:
        def get(self, key, default=''):
            return '10.4' if key == 'q_q2' else default
    score, total, *_ = app.score_attempt(F(), q)
    assert score == 1.0 and total == 1.0


def test_build_exam_hides_answers():
    q = [{'id':'q3','subject_id':1,'subject_name':'Math','category_id':2,'category_name':'Algebra','type':'multiple_choice','prompt':'2+2?','choices':[['a','3'],['b','4']],'answer':'b'}]
    groups = app.build_exam(123, q)
    rendered = groups['multiple_choice'][0]
    assert 'answer' not in rendered
    assert rendered['category_name'] == 'Algebra'


def test_student_roster_schema_exists():
    with app.get_db() as conn:
        section_columns = {row[1] for row in conn.execute("PRAGMA table_info(sections)").fetchall()}
        student_columns = {row[1] for row in conn.execute("PRAGMA table_info(students)").fetchall()}
    assert {"id", "name", "is_archived"}.issubset(section_columns)
    assert {"id", "full_name", "section_id", "student_code", "email", "password_hash", "must_change_password", "is_active"}.issubset(student_columns)


def test_student_access_defaults_to_accounts():
    with app.get_db() as conn:
        assert app.setting(conn, "student_access_mode") == "accounts"


def test_student_password_policy_and_email_normalization():
    assert app.normalize_email(" Student@Test.COM ") == "student@test.com"
    assert app.valid_email("student@test.com")
    assert not app.valid_student_password("short")
    assert app.valid_student_password("Secure123")
    generated = app.generate_temp_password()
    assert app.valid_student_password(generated)


def test_ui_language_defaults_to_spanish():
    with app.get_db() as conn:
        assert app.setting(conn, "ui_language") == "es"


def test_ui_language_is_safe_without_request_context():
    assert app.get_ui_language() == 'es'


def test_question_clone_keeps_legacy_safe_fields_and_deactivates_copy():
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute(
            "INSERT OR IGNORE INTO subjects(id,name,description,created_at,updated_at) VALUES(22,'Clone Test','',?,?)",
            (stamp, stamp),
        )
        conn.execute(
            "INSERT OR IGNORE INTO categories(id,subject_id,name,description,sort_order,created_at,updated_at) VALUES(22,22,'Clone Category','',0,?,?)",
            (stamp, stamp),
        )
        conn.execute(
            """INSERT OR REPLACE INTO question_bank
               (id,subject_id,category_id,type,prompt,data_json,answer_json,audio,script,is_active,is_archived,created_at,updated_at)
               VALUES('clone_source',22,22,'multiple_choice','Original prompt','{\"choices\":[[\"c1\",\"A\"],[\"c2\",\"B\"]]}','\"c1\"',NULL,NULL,1,0,?,?)""",
            (stamp, stamp),
        )
        source = conn.execute("SELECT * FROM question_bank WHERE id='clone_source'").fetchone()
        app.clone_question_row(conn, source, 'clone_copy')
        conn.commit()
        copy = conn.execute("SELECT * FROM question_bank WHERE id='clone_copy'").fetchone()
    assert copy is not None
    assert copy['subject_id'] == 22
    assert copy['category_id'] == 22
    assert copy['is_active'] == 0
    assert copy['is_archived'] == 0
    assert copy['data_json'] == source['data_json']
    assert copy['answer_json'] == source['answer_json']


def test_multi_exam_schema_exists():
    with app.get_db() as conn:
        for table in ("exams", "exam_versions", "exam_version_questions", "exam_assignments", "student_exam_allocations"):
            assert conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone()
        attempt_columns = {row[1] for row in conn.execute("PRAGMA table_info(attempts)").fetchall()}
    assert {"exam_id", "exam_version_id", "assignment_id", "exam_version_name"}.issubset(attempt_columns)


def test_student_login_forces_password_change_then_shows_exam_dashboard():
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute("INSERT OR IGNORE INTO sections(id,name,description,is_archived,created_at,updated_at) VALUES(91,'Login Test','',0,?,?)", (stamp, stamp))
        conn.execute(
            """INSERT OR REPLACE INTO students
               (id,full_name,section_id,student_code,email,password_hash,must_change_password,password_updated_at,notes,is_active,is_archived,created_at,updated_at)
               VALUES(91,'Login Student',91,'LOGIN-91','login.student@example.com',?,1,?,NULL,1,0,?,?)""",
            (app.generate_password_hash('TempPass123'), stamp, stamp, stamp),
        )
        conn.execute("INSERT OR IGNORE INTO subjects(id,name,description,is_archived,created_at,updated_at) VALUES(92,'Login Subject','',0,?,?)", (stamp, stamp))
        conn.execute("INSERT OR IGNORE INTO categories(id,subject_id,name,description,sort_order,is_archived,created_at,updated_at) VALUES(92,92,'Login Category','',0,0,?,?)", (stamp, stamp))
        conn.execute(
            """INSERT OR REPLACE INTO question_bank
               (id,subject_id,category_id,type,prompt,data_json,answer_json,audio,script,is_active,is_archived,created_at,updated_at)
               VALUES('login_q1',92,92,'multiple_choice','2 + 2?','{"choices":[["a","3"],["b","4"]]}','"b"',NULL,NULL,0,0,?,?)""",
            (stamp, stamp),
        )
        cur = conn.execute("INSERT INTO exams(subject_id,title,description,is_published,is_archived,created_at,updated_at) VALUES(92,'Login Exam','',1,0,?,?)", (stamp, stamp))
        exam_id = cur.lastrowid
        cur = conn.execute("INSERT INTO exam_versions(exam_id,name,is_active,created_at,updated_at) VALUES(?,'A',1,?,?)", (exam_id, stamp, stamp))
        version_id = cur.lastrowid
        conn.execute("INSERT INTO exam_version_questions(version_id,question_id,position) VALUES(?,'login_q1',0)", (version_id,))
        cur = conn.execute("INSERT INTO exam_assignments(exam_id,section_id,version_mode,fixed_version_id,is_active,created_at,updated_at) VALUES(?,91,'fixed',?,1,?,?)", (exam_id, version_id, stamp, stamp))
        assignment_id = cur.lastrowid
        conn.commit()

    client = app.app.test_client()
    with client.session_transaction() as sess:
        sess['csrf_token'] = 'account-test-token'
    response = client.post('/start', data={'csrf_token':'account-test-token','email':'LOGIN.STUDENT@example.com','password':'TempPass123'})
    assert response.status_code == 302 and response.headers['Location'].endswith('/student/password')
    with client.session_transaction() as sess:
        token = sess['csrf_token']
        assert sess['student_id'] == 91
    response = client.post('/student/password', data={'csrf_token':token,'new_password':'MyNewPass456','confirm_password':'MyNewPass456'})
    assert response.status_code == 302 and response.headers['Location'].endswith('/student')
    with client.session_transaction() as sess:
        token = sess['csrf_token']
    with app.get_db() as conn:
        rules = app.current_rules(conn)
    response = client.post('/student/rules/acknowledge', data={'csrf_token':token, 'rules_version':rules['version'], 'accept_rules':'1'})
    assert response.status_code == 302
    response = client.post(f'/student/exams/{assignment_id}/start', data={'csrf_token':token})
    assert response.status_code == 302 and response.headers['Location'].endswith('/exam')
    with app.get_db() as conn:
        student = conn.execute("SELECT * FROM students WHERE id=91").fetchone()
        attempt = conn.execute("SELECT * FROM attempts WHERE student_id=91 AND assignment_id=?", (assignment_id,)).fetchone()
    with client.session_transaction() as sess:
        accepted_at = sess['rules_acknowledged_at']
    assert student['must_change_password'] == 0
    assert app.check_password_hash(student['password_hash'], 'MyNewPass456')
    assert attempt['student_email'] == 'login.student@example.com'
    assert attempt['exam_version_name'] == 'A'
    assert attempt['ui_language'] in app.SUPPORTED_UI_LANGUAGES
    assert attempt['policy_version'] == rules['version']
    assert attempt['policy_accepted_at'] == accepted_at
    assert attempt['policy_text'] == rules['text']


def test_one_attempt_unique_per_student_assignment():
    with app.get_db() as conn:
        indexes = conn.execute("PRAGMA index_list(attempts)").fetchall()
        names = {row[1] for row in indexes}
    assert 'idx_attempt_once_per_assignment' in names



def _teacher_client():
    client = app.app.test_client()
    with app.get_db() as conn:
        teacher = conn.execute("SELECT id FROM teachers ORDER BY id LIMIT 1").fetchone()
    with client.session_transaction() as sess:
        sess['teacher_authenticated'] = True
        sess['teacher_id'] = teacher['id']
        sess['csrf_token'] = 'teacher-test-token'
    return client


def _admin_client(csrf_token='admin-test-token'):
    client = app.app.test_client()
    with app.get_db() as conn:
        admin = conn.execute("SELECT id FROM teachers WHERE role='admin' AND is_active=1 ORDER BY id LIMIT 1").fetchone()
    with client.session_transaction() as sess:
        sess.update(teacher_authenticated=True, teacher_id=admin['id'], csrf_token=csrf_token)
    return client


def test_bulk_student_csv_import_generates_email_from_nie():
    import io
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute("INSERT OR IGNORE INTO sections(id,name,description,is_archived,created_at,updated_at) VALUES(301,'CSV Students','',0,?,?)", (stamp, stamp))
        conn.commit()
    payload = 'NIE,Nombre,Apellido,Correo\n990001,Ana,Prueba,\n990002,Carlos,Prueba,\n'.encode('utf-8')
    client = _teacher_client()
    response = client.post('/teacher/students/import', data={
        'csrf_token':'teacher-test-token', 'section_id':'301', 'generate_email_from_nie':'on',
        'email_domain':'clases.edu', 'batch_password_mode':'manual', 'batch_password':'Temporal123',
        'must_change_password':'on', 'csv_file':(io.BytesIO(payload),'students.csv')
    }, content_type='multipart/form-data')
    assert response.status_code == 200
    with app.get_db() as conn:
        rows = conn.execute("SELECT student_code,email FROM students WHERE section_id=301 ORDER BY student_code").fetchall()
    assert [(r['student_code'], r['email']) for r in rows] == [('990001','990001@clases.edu'),('990002','990002@clases.edu')]


def test_question_csv_import_supports_multiple_types():
    import io
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute("INSERT OR IGNORE INTO subjects(id,name,description,is_archived,created_at,updated_at) VALUES(302,'CSV Subject','',0,?,?)", (stamp, stamp))
        conn.execute("INSERT OR IGNORE INTO categories(id,subject_id,name,description,sort_order,is_archived,created_at,updated_at) VALUES(302,302,'CSV Category','',0,0,?,?)", (stamp, stamp))
        conn.commit()
    payload = ('Tipo,Pregunta,Opciones,Respuesta,Pares,Orden,Tolerancia,SensibleMayusculas,Script,Audio\n'
               'multiple_choice,"2 + 2?","3|4|5",4,,,,,,\n'
               'numeric,"10 / 2?",,5,,,0,,,\n'
               'matching,"Relaciona",,,"A=>1|B=>2",,,,,\n').encode('utf-8')
    client = _teacher_client()
    response = client.post('/teacher/questions/import', data={
        'csrf_token':'teacher-test-token','default_subject_id':'302','default_category_id':'302',
        'csv_file':(io.BytesIO(payload),'questions.csv')
    }, content_type='multipart/form-data')
    assert response.status_code == 302
    with app.get_db() as conn:
        types = {r['type'] for r in conn.execute("SELECT type FROM question_bank WHERE subject_id=302").fetchall()}
    assert {'multiple_choice','numeric','matching'}.issubset(types)


def test_listening_without_audio_cannot_be_added_to_exam_version():
    import io
    with app.get_db() as conn:
        stamp = app.now_iso()
        teacher_id = conn.execute("SELECT id FROM teachers ORDER BY id LIMIT 1").fetchone()['id']
        conn.execute("INSERT OR IGNORE INTO subjects(id,name,description,is_archived,created_at,updated_at) VALUES(303,'CSV Audio','',0,?,?)", (stamp, stamp))
        conn.execute("INSERT OR IGNORE INTO categories(id,subject_id,name,description,sort_order,is_archived,created_at,updated_at) VALUES(303,303,'Listening','',0,0,?,?)", (stamp, stamp))
        conn.execute("INSERT INTO exams(subject_id,teacher_id,title,description,is_published,is_archived,created_at,updated_at) VALUES(303,?,'Audio Test','',0,0,?,?)", (teacher_id, stamp, stamp))
        exam_id = conn.execute("SELECT id FROM exams WHERE subject_id=303 ORDER BY id DESC LIMIT 1").fetchone()['id']
        conn.execute("INSERT INTO exam_versions(exam_id,name,is_active,created_at,updated_at) VALUES(?,'A',1,?,?)", (exam_id, stamp, stamp))
        version_id = conn.execute("SELECT id FROM exam_versions WHERE exam_id=? ORDER BY id DESC LIMIT 1", (exam_id,)).fetchone()['id']
        conn.commit()
    payload = 'Tipo,Pregunta,Opciones,Respuesta,Script\nlistening,"Listen","Yes|No",Yes,\n'.encode('utf-8')
    client = _teacher_client()
    response = client.post('/teacher/questions/import', data={
        'csrf_token':'teacher-test-token','default_subject_id':'303','default_category_id':'303',
        'csv_file':(io.BytesIO(payload),'listening.csv')
    }, content_type='multipart/form-data')
    assert response.status_code == 302
    with app.get_db() as conn:
        q = conn.execute("SELECT * FROM question_bank WHERE subject_id=303 AND type='listening' ORDER BY created_at DESC LIMIT 1").fetchone()
    assert q is not None and not q['audio']
    response = client.post(f'/teacher/exams/{exam_id}/versions/{version_id}/questions', data={
        'csrf_token':'teacher-test-token','question_ids':q['id']
    })
    assert response.status_code == 302
    with app.get_db() as conn:
        saved = conn.execute("SELECT 1 FROM exam_version_questions WHERE version_id=? AND question_id=?", (version_id,q['id'])).fetchone()
    assert saved is None


def test_clean_question_supports_browser_voice_without_exposing_script():
    q = {
        'id':'tts_q','subject_id':1,'subject_name':'English','category_id':1,'category_name':'Listening',
        'type':'listening','prompt':'Listen','choices':[['a','Yes'],['b','No']],'answer':'a',
        'script':'This should not be embedded in the rendered question payload.',
        'listening_source':'tts','tts_lang':'en-US'
    }
    import random
    cleaned = app.clean_question(q, random.Random(10))
    assert cleaned['listening_source'] == 'tts'
    assert cleaned['tts_lang'] == 'en-US'
    assert 'script' not in cleaned
    assert 'audio' not in cleaned


def test_version_archive_and_restore_routes_exist():
    routes = {rule.rule for rule in app.app.url_map.iter_rules()}
    assert '/teacher/exams/<int:exam_id>/versions/<int:version_id>/archive' in routes
    assert '/teacher/exams/<int:exam_id>/versions/<int:version_id>/restore' in routes
    assert '/api/listening-script/<question_id>' in routes


def test_student_result_page_does_not_offer_history_back_navigation():
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute("INSERT OR IGNORE INTO sections(id,name,description,is_archived,created_at,updated_at) VALUES(401,'Result Test','',0,?,?)", (stamp, stamp))
        conn.execute(
            """INSERT OR REPLACE INTO students
               (id,full_name,section_id,student_code,email,password_hash,must_change_password,password_updated_at,last_login_at,notes,is_active,is_archived,created_at,updated_at)
               VALUES(401,'Result Student',401,'RES-401','result.student@example.com',?,0,?,?,?,?,0,?,?)""",
            (app.generate_password_hash('Student123'), stamp, stamp, None, 1, stamp, stamp),
        )
        conn.execute(
            """INSERT OR REPLACE INTO attempts
               (id,student_id,student_email,student_name,section,started_at,submitted_at,seed,score,total,percentage,grade10,category_scores,type_scores,answers_json,status,questions_json,assessment_title,assessment_subject,category_labels_json,ui_language)
               VALUES('attempt-result',401,'result.student@example.com','Result Student','Result Section',?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                stamp, stamp, 1, 4.0, 4.0, 100.0, 10.0,
                json.dumps({}), json.dumps({}), json.dumps({}), 'submitted',
                json.dumps([]), 'Sample Exam', 'Sample Subject', json.dumps({}), 'es',
            ),
        )
        conn.commit()

    client = app.app.test_client()
    with client.session_transaction() as sess:
        sess['student_authenticated'] = True
        sess['student_id'] = 401
    response = client.get('/result/attempt-result')
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'history.back()' not in html
    assert 'Back to previous page' not in html


def test_teacher_dashboard_stats_are_scoped_to_teacher():
    with app.get_db() as conn:
        stamp = app.now_iso()
        for teacher_id in (810, 811):
            conn.execute(
                """INSERT OR REPLACE INTO teachers
                   (id,full_name,email,password_hash,role,is_active,must_change_password,created_at,updated_at)
                   VALUES(?,?,?,?, 'teacher',1,0,?,?)""",
                (teacher_id, f'Dashboard Teacher {teacher_id}', f'dashboard{teacher_id}@example.com', app.generate_password_hash('Teacher123'), stamp, stamp),
            )
        conn.execute("DELETE FROM attempts WHERE id IN ('dashboard-a','dashboard-b','dashboard-other')")
        attempts = [
            ('dashboard-a', 810, 80.0, 8.0, 0, 0),
            ('dashboard-b', 810, 60.0, 6.0, 2, 1),
            ('dashboard-other', 811, 100.0, 10.0, 0, 0),
        ]
        for attempt_id, teacher_id, percentage, grade, focus, copy in attempts:
            conn.execute(
                """INSERT INTO attempts
                   (id,student_name,section,started_at,submitted_at,seed,score,total,percentage,grade10,status,
                    focus_departures,copy_attempts,teacher_id)
                   VALUES(?, 'Dashboard Student','A',?,?,1,8,10,?,?,'submitted',?,?,?)""",
                (attempt_id, stamp, stamp, percentage, grade, focus, copy, teacher_id),
            )
        conn.commit()
        visible = conn.execute("SELECT * FROM attempts WHERE teacher_id=810 AND status='submitted'").fetchall()
        stats = app.teacher_dashboard_stats(conn, 810, visible)
    assert stats['submitted'] == 2
    assert stats['average_percentage'] == 70.0
    assert stats['average_grade'] == 7.0
    assert stats['clean_count'] == 1
    assert stats['review_count'] == 1
    assert stats['focus_total'] == 2
    assert stats['blocked_total'] == 1


def test_teacher_dashboard_renders_five_item_mobile_navigation():
    response = _teacher_client().get('/teacher')
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'class="app-sidebar"' in html
    assert 'class="mobile-app-bar"' in html
    assert 'data-more-open' in html
    assert 'id="teacherMoreSheet"' in html
    nav = re.search(r'<nav class="mobile-bottom-nav".*?</nav>', html, re.S)
    assert nav is not None
    assert len(re.findall(r'<(?:a|button)\b', nav.group(0))) == 5
    assert 'method="post" action="/teacher/logout"' in html


def test_local_phosphor_sprite_macro_and_license_are_complete():
    root = Path(app.__file__).parent
    sprite = root / 'static/icons/phosphor-sprite.svg'
    macro = root / 'templates/_icons.html'
    license_file = root / 'static/icons/LICENSE'
    assert sprite.is_file() and macro.is_file() and license_file.is_file()

    sprite_source = sprite.read_text(encoding='utf-8')
    macro_source = macro.read_text(encoding='utf-8')
    license_source = license_file.read_text(encoding='utf-8')
    symbol_names = set(re.findall(r'<symbol id="ph-([^"]+)"', sprite_source))
    literal_uses = set()
    for template in (root / 'templates').glob('*.html'):
        literal_uses.update(re.findall(r"icon\(\s*['\"]([^'\"]+)", template.read_text(encoding='utf-8')))

    assert literal_uses <= symbol_names
    assert "url_for('static', filename='icons/phosphor-sprite.svg')" in macro_source
    assert 'aria-hidden="true"' in macro_source and 'focusable="false"' in macro_source
    assert 'MIT License' in license_source
    assert '@phosphor-icons/core 2.1.1' in license_source


def test_templates_have_no_remote_icon_dependency():
    root = Path(app.__file__).parent
    source = '\n'.join(
        path.read_text(encoding='utf-8')
        for folder in ('templates', 'static/css', 'static/js')
        for path in (root / folder).glob('**/*')
        if path.is_file() and path.suffix in {'.html', '.css', '.js'}
    )
    remote_icon_url = re.compile(r'https?://[^\s"\']*(?:phosphor|unpkg|cdnjs)[^\s"\']*', re.I)
    assert not remote_icon_url.search(source)


def test_primary_teacher_navigation_renders_local_accessible_svg_icons():
    html = _teacher_client().get('/teacher').get_data(as_text=True)
    sidebar = re.search(r'<aside class="app-sidebar".*?</aside>', html, re.S)
    assert sidebar is not None
    markup = sidebar.group(0)
    for icon_name in ('chart-bar', 'file-text', 'users-three', 'question', 'folders'):
        assert f'/static/icons/phosphor-sprite.svg#ph-{icon_name}' in markup
    assert markup.count('aria-hidden="true"') >= 5
    assert markup.count('focusable="false"') >= 5
    assert ('Results' in markup or 'Resultados' in markup)
    assert ('Exams' in markup or 'Exámenes' in markup)
    assert not set('⌂☰⋯📊📄👥❓📁').intersection(markup)


def test_teacher_roster_renders_semantic_tables_and_preserves_row_actions():
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute(
            """INSERT OR REPLACE INTO sections(id,name,description,is_archived,created_at,updated_at)
               VALUES(913,'Responsive Roster','Morning group',0,?,?)""",
            (stamp, stamp),
        )
        conn.execute(
            """INSERT OR REPLACE INTO students
               (id,full_name,section_id,student_code,email,password_hash,must_change_password,notes,is_active,is_archived,created_at,updated_at)
               VALUES(913,'Roster Student',913,'RST-913','roster.student@example.com',?,1,'Needs front-row seating',1,0,?,?)""",
            (app.generate_password_hash('Student123'), stamp, stamp),
        )
        conn.commit()

    response = _teacher_client().get('/teacher/students')
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert html.count('data-roster-table="sections"') == 1
    assert html.count('data-roster-table="students"') == 1
    for dataset in ('sections', 'students'):
        table = re.search(rf'<table[^>]+data-roster-table="{dataset}".*?</table>', html, re.S)
        assert table is not None
        assert '<thead>' in table.group(0) and '<tbody>' in table.group(0)
        assert 'scope="col"' in table.group(0) and 'scope="row"' in table.group(0)
        assert 'data-label=' in table.group(0)

    section_row = re.search(r'<tr data-section-row="913".*?</tr>', html, re.S)
    student_row = re.search(r'<tr[^>]+data-student-row="913".*?</tr>', html, re.S)
    assert section_row is not None and student_row is not None
    assert 'data-url="/teacher/sections/913/edit"' in section_row.group(0)
    assert 'action="/teacher/sections/913/archive"' in section_row.group(0)
    assert 'data-confirm' in section_row.group(0) and 'name="csrf_token"' in section_row.group(0)
    assert 'data-url="/teacher/students/913/edit"' in student_row.group(0)
    assert 'data-url="/teacher/students/913/password"' in student_row.group(0)
    assert 'action="/teacher/students/913/archive"' in student_row.group(0)
    assert 'data-section-id="913"' in student_row.group(0)
    assert 'data-active="1"' in student_row.group(0)
    assert 'data-confirm' in student_row.group(0) and 'name="csrf_token"' in student_row.group(0)


def test_teacher_roster_keeps_password_reset_disabled_without_email():
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute(
            """INSERT OR REPLACE INTO sections(id,name,description,is_archived,created_at,updated_at)
               VALUES(914,'Pending Accounts','',0,?,?)""",
            (stamp, stamp),
        )
        conn.execute(
            """INSERT OR REPLACE INTO students
               (id,full_name,section_id,student_code,email,password_hash,must_change_password,notes,is_active,is_archived,created_at,updated_at)
               VALUES(914,'Pending Student',914,'RST-914',NULL,NULL,0,NULL,1,0,?,?)""",
            (stamp, stamp),
        )
        conn.commit()

    html = _teacher_client().get('/teacher/students').get_data(as_text=True)
    student_row = re.search(r'<tr[^>]+data-student-row="914".*?</tr>', html, re.S)
    assert student_row is not None
    reset = re.search(r'<button[^>]+data-open-password-modal[^>]*>', student_row.group(0))
    assert reset is not None and 'disabled' in reset.group(0)


def test_student_dashboard_route_renders_backend_statistics(monkeypatch):
    student = {'id': 912, 'full_name': 'Dashboard Student', 'email': 'student@example.com', 'section_name': 'B'}
    exams = [
        {'id': 1, 'title': 'Resume Exam', 'description': '', 'subject_name': 'Math', 'attempt_status': 'in_progress', 'grade10': None, 'version_count': 2, 'version_mode': 'random', 'attempt_id': 'one'},
        {'id': 2, 'title': 'Available Exam', 'description': '', 'subject_name': 'Science', 'attempt_status': None, 'grade10': None, 'version_count': 1, 'version_mode': 'fixed', 'attempt_id': None},
        {'id': 3, 'title': 'Completed Exam', 'description': '', 'subject_name': 'History', 'attempt_status': 'submitted', 'grade10': 8.4, 'version_count': 1, 'version_mode': 'fixed', 'attempt_id': 'three'},
    ]
    monkeypatch.setattr(app, 'student_exam_rows', lambda _conn, _student_id: (student, exams))
    client = app.app.test_client()
    with client.session_transaction() as sess:
        sess['student_authenticated'] = True
        sess['student_id'] = student['id']
        sess['csrf_token'] = 'student-dashboard-token'
    response = client.get('/student')
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'data-stat="total-assigned" data-value="3"' in html
    assert 'data-stat="available" data-value="1"' in html
    assert 'data-stat="in-progress" data-value="1"' in html
    assert 'data-stat="completed" data-value="1"' in html
    assert 'data-stat="average-grade" data-value="8.4"' in html
    assert 'Resume Exam' in html


def _literal_translation_keys():
    keys = set()
    tree = ast.parse(Path(app.__file__).read_text(encoding='utf-8'))
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call) or not node.args or not isinstance(node.args[0], ast.Constant):
            continue
        name = getattr(node.func, 'id', getattr(node.func, 'attr', ''))
        if name in {'tr', 'rt', 'flash_ui'} and isinstance(node.args[0].value, str):
            keys.add(node.args[0].value)
    pattern = re.compile(r"(?<![A-Za-z_])t\(\s*['\"]([^'\"]+)['\"]")
    for template in (Path(app.__file__).parent / 'templates').glob('*.html'):
        keys.update(pattern.findall(template.read_text(encoding='utf-8')))
    return keys


def test_every_literal_ui_key_has_a_spanish_translation():
    allowlist = {'Assessment Studio', 'CSV', 'Excel', 'English', 'NIE'}
    missing = _literal_translation_keys() - set(app.TRANSLATIONS_ES) - allowlist
    assert not missing, f'Missing Spanish translations: {sorted(missing)}'


def test_global_language_renders_spanish_and_english_login():
    client = app.app.test_client()
    with app.get_db() as conn:
        original = app.setting(conn, 'ui_language')
        conn.execute("UPDATE app_settings SET value='es' WHERE key='ui_language'")
        conn.commit()
    assert 'Portal del estudiante' in client.get('/').get_data(as_text=True)
    with app.get_db() as conn:
        conn.execute("UPDATE app_settings SET value='en' WHERE key='ui_language'")
        conn.commit()
    assert 'Student portal' in client.get('/').get_data(as_text=True)
    with app.get_db() as conn:
        conn.execute("UPDATE app_settings SET value=? WHERE key='ui_language'", (original,))
        conn.commit()


def test_exam_js_dictionary_has_matching_spanish_and_english_surfaces():
    spanish = app.exam_js_strings('es')
    english = app.exam_js_strings('en')
    assert spanish.keys() == english.keys()
    assert spanish['invalidCurrent'] == 'Completa correctamente la pregunta actual antes de continuar.'
    assert english['invalidCurrent'] == 'Complete the current question before continuing.'
    assert spanish['allAnswered'] != english['allAnswered']


def test_public_login_is_basic_without_marketing_sidebar():
    html = app.app.test_client().get('/').get_data(as_text=True)
    assert 'login-card-basic' in html
    assert 'login-aside' not in html and 'feature-stack' not in html and 'account-kicker' not in html


def test_setup_is_admin_only():
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute("""INSERT OR REPLACE INTO teachers(id,full_name,email,password_hash,role,is_active,must_change_password,created_at,updated_at)
                        VALUES(880,'Ordinary Teacher','ordinary@example.com',?,'teacher',1,0,?,?)""",
                     (app.generate_password_hash('Teacher123'), stamp, stamp))
        conn.commit()
    client = app.app.test_client()
    with client.session_transaction() as sess:
        sess.update(teacher_authenticated=True, teacher_id=880, csrf_token='teacher-only')
    assert client.get('/teacher/setup').status_code == 403
    assert _teacher_client().get('/teacher/setup').status_code == 200


def test_setup_policy_text_change_auto_bumps_version_and_invalidates_session_acknowledgment():
    keys = ('institution_name', 'assessment_subtitle', 'listening_max_plays', 'ui_language',
            'student_rules_es', 'student_rules_en', 'rules_version', 'require_rules_acknowledgment')
    with app.get_db() as conn:
        original = {key: app.setting(conn, key) for key in keys}
    expected_version = app.bump_rules_version(original['rules_version'])
    client = _admin_client()
    with client.session_transaction() as sess:
        sess['rules_acknowledged_version'] = original['rules_version']
        sess['rules_acknowledged_at'] = app.now_iso()
    response = client.post('/teacher/setup', data={
        'csrf_token': 'admin-test-token',
        'institution_name': original['institution_name'],
        'assessment_subtitle': original['assessment_subtitle'],
        'listening_max_plays': original['listening_max_plays'],
        'ui_language': original['ui_language'],
        'student_rules_es': f"{original['student_rules_es']} Cambio de revision.",
        'student_rules_en': original['student_rules_en'],
        'rules_version': original['rules_version'],
        'require_rules_acknowledgment': 'on',
    })
    assert response.status_code == 302
    with app.get_db() as conn:
        assert app.setting(conn, 'rules_version') == expected_version
        with app.app.test_request_context('/'):
            app.session['rules_acknowledged_version'] = original['rules_version']
            app.session['rules_acknowledged_at'] = app.now_iso()
            assert not app.rules_acknowledged(conn)
        for key, value in original.items():
            conn.execute("UPDATE app_settings SET value=? WHERE key=?", (value, key))
        conn.commit()


def test_app_and_seed_share_canonical_policy_defaults():
    assert app.DEFAULT_SETTINGS['student_rules_es'] == policy_defaults.DEFAULT_STUDENT_RULES_ES
    assert app.DEFAULT_SETTINGS['student_rules_en'] == policy_defaults.DEFAULT_STUDENT_RULES_EN
    assert app.DEFAULT_SETTINGS['rules_version'] == policy_defaults.DEFAULT_RULES_VERSION
    assert seed.DEFAULT_STUDENT_RULES_ES == policy_defaults.DEFAULT_STUDENT_RULES_ES
    assert seed.DEFAULT_STUDENT_RULES_EN == policy_defaults.DEFAULT_STUDENT_RULES_EN
    assert seed.DEFAULT_RULES_VERSION == policy_defaults.DEFAULT_RULES_VERSION


@pytest.mark.parametrize(('question', 'data'), [
    ({'id':'mc','type':'multiple_choice','choices':[['a','A']], 'answer':'a'}, {'q_mc':'bad'}),
    ({'id':'listen','type':'listening','choices':[['a','A']], 'answer':'a'}, {'q_listen':''}),
    ({'id':'tf','type':'true_false','choices':[['true','True'],['false','False']], 'answer':'true'}, {'q_tf':'maybe'}),
    ({'id':'short','type':'short_answer','answer':['ok']}, {'q_short':'   '}),
    ({'id':'num','type':'numeric','answer':{'value':1,'tolerance':0}}, {'q_num':'NaN'}),
    ({'id':'order','type':'order','items':[['a','A'],['b','B']], 'answer':['a','b']}, {'q_order':'["a","a"]','q_order__touched':'1'}),
    ({'id':'match','type':'matching','left':[['l1','L1'],['l2','L2']], 'right':[['r1','R1'],['r2','R2']], 'answer':{'l1':'r1','l2':'r2'}}, {'q_match__l1':'r1','q_match__l2':'r1'}),
])
def test_required_answer_validation_rejects_missing_or_malformed_by_type(question, data):
    invalid, _ = app.validate_required_answers(data, [dict(question, category_id=1)])
    assert invalid == [question['id']]


def _submission_attempt(attempt_id='submit-test', student_id=990, teacher_id=None):
    questions = [
        {'id':'mc','category_id':1,'type':'multiple_choice','choices':[['a','A'],['b','B']], 'answer':'a'},
        {'id':'short','category_id':1,'type':'short_answer','answer':['yes'],'case_sensitive':False},
        {'id':'num','category_id':1,'type':'numeric','answer':{'value':2.5,'tolerance':0}},
        {'id':'order','category_id':1,'type':'order','items':[['a','A'],['b','B']], 'answer':['a','b']},
        {'id':'match','category_id':1,'type':'matching','left':[['l1','L1'],['l2','L2']], 'right':[['r1','R1'],['r2','R2']], 'answer':{'l1':'r1','l2':'r2'}},
    ]
    questions = [dict(question, subject_id=1, subject_name='General', category_name='General', prompt=question['id']) for question in questions]
    with app.get_db() as conn:
        if teacher_id is None:
            teacher_id = conn.execute('SELECT id FROM teachers ORDER BY id LIMIT 1').fetchone()['id']
        stamp = app.now_iso()
        conn.execute("DELETE FROM attempts WHERE id=?", (attempt_id,))
        conn.execute("""INSERT INTO attempts(id,teacher_id,student_id,student_name,section,started_at,seed,status,questions_json,assessment_title,ui_language)
                        VALUES(?,?,?,?,?,?,1,'in_progress',?,?, 'en')""",
                     (attempt_id, teacher_id, student_id, 'Submit Student', 'T', stamp, json.dumps(questions), 'Submit Exam'))
        conn.commit()
    return questions, teacher_id


def _student_submit_client(attempt_id='submit-test', student_id=990):
    client = app.app.test_client()
    with client.session_transaction() as sess:
        sess.update(student_authenticated=True, student_id=student_id, attempt_id=attempt_id, csrf_token='submit-token')
    return client


def _valid_submission_data():
    return {'csrf_token':'submit-token','q_mc':'a','q_short':'yes','q_num':'2.5','q_order':'["a","b"]',
            'q_order__touched':'1','q_match__l1':'r1','q_match__l2':'r2'}


def test_submit_requires_csrf_and_student_ownership():
    _submission_attempt()
    client = _student_submit_client()
    assert client.post('/submit', data={}).status_code == 403
    wrong = _student_submit_client(student_id=991)
    assert wrong.post('/submit', data=_valid_submission_data()).status_code == 403


def test_incomplete_submission_stays_in_progress_and_valid_submission_succeeds():
    _submission_attempt()
    client = _student_submit_client()
    data = _valid_submission_data()
    data['q_num'] = 'Infinity'
    response = client.post('/submit', data=data)
    assert response.status_code == 302 and response.headers['Location'].endswith('/exam')
    with app.get_db() as conn:
        attempt = conn.execute("SELECT status,answers_json FROM attempts WHERE id='submit-test'").fetchone()
    assert attempt['status'] == 'in_progress' and json.loads(attempt['answers_json'])['num'] == 'Infinity'
    response = client.post('/submit', data=_valid_submission_data())
    assert response.status_code == 302 and '/result/submit-test' in response.headers['Location']
    with app.get_db() as conn:
        assert conn.execute("SELECT status FROM attempts WHERE id='submit-test'").fetchone()['status'] == 'submitted'


def test_rules_acknowledgment_lifecycle_and_attempt_language_snapshot():
    client = app.app.test_client()
    with app.get_db() as conn:
        stamp = app.now_iso()
        teacher_id = conn.execute("SELECT id FROM teachers ORDER BY id LIMIT 1").fetchone()['id']
        conn.execute("INSERT OR REPLACE INTO sections(id,name,description,is_archived,created_at,updated_at) VALUES(591,'Rules Test','',0,?,?)", (stamp, stamp))
        conn.execute("""INSERT OR REPLACE INTO students
                        (id,full_name,section_id,student_code,email,password_hash,must_change_password,is_active,is_archived,created_at,updated_at)
                        VALUES(591,'Rules Student',591,'RULES-591','rules.student@example.com',?,0,1,0,?,?)""",
                     (app.generate_password_hash('Student123'), stamp, stamp))
        conn.execute("INSERT OR REPLACE INTO subjects(id,name,description,is_archived,created_at,updated_at) VALUES(591,'Rules Subject','',0,?,?)", (stamp, stamp))
        conn.execute("INSERT OR REPLACE INTO categories(id,subject_id,name,description,sort_order,is_archived,created_at,updated_at) VALUES(591,591,'Rules Category','',0,0,?,?)", (stamp, stamp))
        conn.execute("""INSERT OR REPLACE INTO question_bank
                        (id,teacher_id,subject_id,category_id,type,prompt,data_json,answer_json,is_active,is_archived,created_at,updated_at)
                        VALUES('rules_q1',?,591,591,'multiple_choice','Rules question','{"choices":[["a","A"],["b","B"]]}','"a"',1,0,?,?)""",
                     (teacher_id, stamp, stamp))
        exam_id = conn.execute("INSERT INTO exams(teacher_id,subject_id,title,description,is_published,is_archived,created_at,updated_at) VALUES(?,591,'Rules Exam','',1,0,?,?)", (teacher_id, stamp, stamp)).lastrowid
        version_id = conn.execute("INSERT INTO exam_versions(exam_id,name,is_active,created_at,updated_at) VALUES(?,'A',1,?,?)", (exam_id, stamp, stamp)).lastrowid
        conn.execute("INSERT INTO exam_version_questions(version_id,question_id,position) VALUES(?,'rules_q1',0)", (version_id,))
        assignment_id = conn.execute("INSERT INTO exam_assignments(exam_id,section_id,version_mode,fixed_version_id,is_active,created_at,updated_at) VALUES(?,591,'fixed',?,1,?,?)", (exam_id, version_id, stamp, stamp)).lastrowid
        original_language = app.setting(conn, 'ui_language')
        conn.execute("UPDATE app_settings SET value='en' WHERE key='ui_language'")
        version = app.setting(conn, 'rules_version')
        conn.commit()
    with client.session_transaction() as sess:
        sess.update(student_authenticated=True, student_id=591, csrf_token='rules-token')
    html = client.get('/student').get_data(as_text=True)
    assert 'studentRulesModal' in html and 'data-static-modal' in html
    bypass = client.post(f"/student/exams/{assignment_id}/start", data={'csrf_token':'rules-token'})
    assert bypass.status_code == 302 and bypass.headers['Location'].endswith('/student')
    assert client.post('/student/rules/acknowledge', data={'csrf_token':'rules-token','rules_version':version}).status_code == 302
    client.post('/student/rules/acknowledge', data={'csrf_token':'rules-token','rules_version':version,'accept_rules':'1'})
    with client.session_transaction() as sess:
        assert sess['rules_acknowledged_version'] == version
        sess['attempt_id'] = 'language-snapshot'
    _submission_attempt('language-snapshot', student_id=591)
    with app.get_db() as conn:
        conn.execute("UPDATE attempts SET ui_language='en',policy_version=?,policy_accepted_at=?,policy_text='Rules' WHERE id='language-snapshot'", (version, app.now_iso()))
        conn.execute("UPDATE app_settings SET value='es' WHERE key='ui_language'")
        conn.commit()
    response = client.get('/exam')
    assert response.status_code == 200 and 'Previous' in response.get_data(as_text=True)
    client.post('/student/logout', data={'csrf_token':'rules-token'})
    with client.session_transaction() as sess:
        assert 'rules_acknowledged_version' not in sess
    with app.get_db() as conn:
        conn.execute("UPDATE app_settings SET value=? WHERE key='ui_language'", (original_language,))
        conn.commit()


def test_penalties_are_tenant_owned_auditable_revocable_and_raw_grade_is_unchanged():
    _, owner_id = _submission_attempt('penalty-attempt', student_id=995)
    with app.get_db() as conn:
        conn.execute("UPDATE attempts SET status='submitted',submitted_at=?,score=8,total=10,percentage=80,grade10=8 WHERE id='penalty-attempt'", (app.now_iso(),))
        conn.commit()
    client = app.app.test_client()
    with client.session_transaction() as sess:
        sess.update(teacher_authenticated=True, teacher_id=owner_id, csrf_token='penalty-token')
    response = client.post('/teacher/results/penalty-attempt/penalties', data={'csrf_token':'penalty-token','points':'1.5','reason':'Reviewed focus evidence'})
    assert response.status_code == 302
    with app.get_db() as conn:
        attempt = conn.execute("SELECT grade10 FROM attempts WHERE id='penalty-attempt'").fetchone()
        penalty = conn.execute("SELECT * FROM attempt_penalties WHERE attempt_id='penalty-attempt'").fetchone()
        assert attempt['grade10'] == 8 and app.active_penalty_total(conn, 'penalty-attempt') == 1.5
    html = client.get('/result/penalty-attempt').get_data(as_text=True)
    assert '6.5' in html and 'Reviewed focus evidence' in html
    pdf = client.get('/report/penalty-attempt.pdf')
    assert pdf.status_code == 200
    assert pdf.mimetype == 'application/pdf'
    assert pdf.data.startswith(b'%PDF')
    client.post(f"/teacher/results/penalty-attempt/penalties/{penalty['id']}/revoke", data={'csrf_token':'penalty-token'})
    with app.get_db() as conn:
        assert app.active_penalty_total(conn, 'penalty-attempt') == 0
        assert conn.execute("SELECT is_active,revoked_at FROM attempt_penalties WHERE id=?", (penalty['id'],)).fetchone()['revoked_at']


def test_penalty_cap_cross_teacher_and_export_values():
    _, owner_id = _submission_attempt('penalty-cap', student_id=996)
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute("""INSERT OR REPLACE INTO teachers(id,full_name,email,password_hash,role,is_active,must_change_password,created_at,updated_at)
                        VALUES(881,'Other Export Teacher','other-export@example.com',?,'teacher',1,0,?,?)""",
                     (app.generate_password_hash('Teacher123'), stamp, stamp))
        conn.execute("UPDATE attempts SET status='submitted',submitted_at=?,score=5,total=10,percentage=50,grade10=5 WHERE id='penalty-cap'", (app.now_iso(),))
        conn.commit()
    owner = app.app.test_client()
    with owner.session_transaction() as sess:
        sess.update(teacher_authenticated=True, teacher_id=owner_id, csrf_token='owner-token')
    owner.post('/teacher/results/penalty-cap/penalties', data={'csrf_token':'owner-token','points':'6','reason':'Too much'})
    with app.get_db() as conn:
        assert app.active_penalty_total(conn, 'penalty-cap') == 0
    other = app.app.test_client()
    with other.session_transaction() as sess:
        sess.update(teacher_authenticated=True, teacher_id=881, csrf_token='other-token')
    assert other.post('/teacher/results/penalty-cap/penalties', data={'csrf_token':'other-token','points':'1','reason':'Forged'}).status_code == 404
    owner.post('/teacher/results/penalty-cap/penalties', data={'csrf_token':'owner-token','points':'1.25','reason':'Reviewed evidence'})
    exported = list(csv.reader(io.StringIO(owner.get('/teacher/export.csv').get_data(as_text=True))))
    csv_row = next(row for row in exported[1:] if row[0] == 'penalty-cap')
    assert list(map(float, csv_row[12:15])) == [5.0, 1.25, 3.75]

    xlsx = owner.get('/teacher/export.xlsx')
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.data), read_only=True, data_only=True)
    result_rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    xlsx_row = next(row for row in result_rows[1:] if row[0] == 'penalty-cap')
    assert list(map(float, xlsx_row[12:15])) == [5.0, 1.25, 3.75]


def test_integrity_event_presentation_pairs_away_periods_and_builds_summary():
    attempt = {
        'seed': 7,
        'questions_json': json.dumps([
            {'id':'q-one','subject_id':1,'category_id':1,'type':'multiple_choice','prompt':'First prompt','choices':[['a','A']]},
            {'id':'q-two','subject_id':1,'category_id':1,'type':'short_answer','prompt':'Second prompt for the teacher'},
        ]),
    }
    events = [
        {'event_type':'visibility_hidden','occurred_at':'2026-08-15T14:05:00','detail_json':json.dumps({'question':'q-two'})},
        {'event_type':'focus_return','occurred_at':'2026-08-15T14:05:06','detail_json':json.dumps({'question':'q-two','duration_ms':6300})},
        {'event_type':'copy','occurred_at':'2026-08-15T14:05:07','detail_json':json.dumps({'question':'q-one'})},
        {'event_type':'shortcut','occurred_at':'2026-08-15T14:05:08','detail_json':json.dumps({'question':'q-one','shortcut':'Ctrl+V'})},
        {'event_type':'fullscreen_exit','occurred_at':'2026-08-15T14:05:09','detail_json':'{'},
    ]

    review = app.integrity_event_presentation(events, attempt, 'en')
    assert len(review['items']) == 4
    assert review['items'][0]['title'] == 'Left the exam for 6.3 seconds'
    assert review['items'][0]['duration'] == '6.3 seconds'
    assert review['items'][0]['question_label'] == 'Question 2'
    assert review['items'][0]['question_excerpt'] == 'Second prompt for the teacher'
    assert review['items'][2]['title'] == 'Blocked attempt to paste'
    assert [item['icon'] for item in review['items']] == ['eye', 'copy', 'keyboard', 'arrows-out']
    assert review['items'][2]['shortcut'] == 'Ctrl+V'
    assert review['items'][3]['technical_details'] == '{'
    assert review['summary'] == {
        'away_periods': 1,
        'total_away_ms': 6300.0,
        'blocked_actions': 2,
        'fullscreen_exits': 1,
        'total_away_time': '6.3 seconds',
    }


@pytest.mark.parametrize(('event_type', 'english_title', 'spanish_title'), [
    ('visibility_hidden', 'Left or switched away from the exam tab', 'Salio o cambio de la pestana del examen'),
    ('focus_return', 'Returned to the exam', 'Regreso al examen'),
    ('window_blur', 'The exam window lost focus', 'La ventana del examen perdio el foco'),
    ('pagehide', 'Left or reloaded the exam page', 'Salio o recargo la pagina del examen'),
    ('contextmenu', 'Blocked context menu attempt', 'Intento bloqueado de abrir el menu contextual'),
    ('copy', 'Blocked copy attempt', 'Intento bloqueado de copiar'),
    ('cut', 'Blocked cut attempt', 'Intento bloqueado de cortar'),
    ('paste', 'Blocked paste attempt', 'Intento bloqueado de pegar'),
    ('fullscreen_exit', 'Left fullscreen', 'Salio de pantalla completa'),
])
def test_integrity_event_titles_are_mapped_bilingually(event_type, english_title, spanish_title):
    event = {'event_type':event_type, 'occurred_at':'2026-08-15T14:05:00', 'detail_json':'{}'}
    attempt = {'seed':1, 'questions_json':'[]'}
    assert app.integrity_event_presentation([event], attempt, 'en')['items'][0]['title'] == english_title
    assert app.integrity_event_presentation([event], attempt, 'es')['items'][0]['title'] == spanish_title


def test_integrity_duration_and_timestamp_formatting_are_human_readable():
    assert app.format_integrity_duration(65000, 'en') == '1 minute 5 seconds'
    assert app.format_integrity_duration(65000, 'es') == '1 minuto 5 segundos'
    assert app.format_integrity_duration('bad', 'en') is None
    assert app.format_integrity_timestamp('2026-08-15T14:05:06', 'en') == 'Aug 15, 2026, 2:05:06 PM'
    assert app.format_integrity_timestamp('2026-08-15T14:05:06', 'es') == '15/08/2026, 14:05:06'
    assert app.format_integrity_timestamp('bad', 'en') == 'Time unavailable'


def test_integrity_event_details_render_for_owner_teacher_only():
    _, owner_id = _submission_attempt('integrity-render', student_id=997)
    with app.get_db() as conn:
        stamp = app.now_iso()
        conn.execute("UPDATE attempts SET status='submitted',submitted_at=?,score=1,total=1,percentage=100,grade10=10 WHERE id='integrity-render'", (stamp,))
        conn.execute(
            "INSERT INTO integrity_events(attempt_id,event_type,occurred_at,detail_json) VALUES('integrity-render','copy',?,?)",
            (stamp, json.dumps({'question':'mc','secret':'teacher-only-raw'})),
        )
        conn.commit()

    teacher = app.app.test_client()
    with teacher.session_transaction() as sess:
        sess.update(teacher_authenticated=True, teacher_id=owner_id, csrf_token='integrity-token')
    teacher_html = teacher.get('/result/integrity-render').get_data(as_text=True)
    assert 'data-integrity-review' in teacher_html
    assert 'teacher-only-raw' in teacher_html
    assert '<details class="integrity-technical">' in teacher_html
    assert 'event_type: copy' in teacher_html
    assert '/static/icons/phosphor-sprite.svg#ph-copy' in teacher_html
    assert '<div class="integrity-event-icon"><svg class="ui-icon" aria-hidden="true" focusable="false">' in teacher_html

    student = app.app.test_client()
    with student.session_transaction() as sess:
        sess.update(student_authenticated=True, student_id=997, attempt_id='integrity-render')
    student_html = student.get('/result/integrity-render').get_data(as_text=True)
    assert 'data-integrity-review' not in student_html
    assert 'teacher-only-raw' not in student_html
    assert 'event_type: copy' not in student_html


def test_exam_js_has_central_no_skip_guard_and_order_touch_marker():
    source = (Path(app.__file__).parent / 'static/js/exam.js').read_text(encoding='utf-8')
    template = (Path(app.__file__).parent / 'templates/exam.html').read_text(encoding='utf-8')
    assert 'function navigateTo(index)' in source
    assert "nextBtn.addEventListener('click', () => navigateTo(current + 1))" in source
    assert "dot.addEventListener('click', () => navigateTo(i))" in source
    assert 'for (let i = current; i < target; i += 1)' in source
    assert 'if (i !== current) showCard(i);' in source
    assert 'showQuestionError(cards[i], validity.message);' in source
    assert 'Number.isFinite(value)' in source
    assert '__touched' in template and 'csrf_token' in template
